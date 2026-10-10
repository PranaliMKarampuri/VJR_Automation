import csv
import tempfile
import unittest
from pathlib import Path

import openpyxl

import phi_rectangular_adapter as phi
from stage1_app_backend import FIXED_TEMPLATE, generate_vjr
from stage1_app_backend import preview_input


ROOT = Path(__file__).resolve().parents[1]
HEADERS = ["Subject", "Page Index", "Label", "Count", "DUCT SIZE", "ACCESSORIES1",
           "ACCESSORIES2", "REMARK", "UNITS", "Length", "Length Unit", "DAMPER TYPE"]


def make_phi(path):
    book = openpyxl.Workbook()
    data = book.active
    data.title = "DATA-DUCT"
    data.append(HEADERS)
    data.append(["SPIRAL DUCT", 1, "SPIRAL DUCT", 0, '10"', ".", ".", ".", "SYS-X", 10, "ft", "NO"])
    data.append(["RECT. DUCT ft.", 2, "RECT. DUCT", 0, "8X8", 'w/ 1\" ACOUSTIC INSULATION', ".", "LOW PRESSURE", "SYS-X", 7.5, "ft", "NO"])
    data.append(["RECT. DUCT ft.", 3, "RECT. DUCT", 0, "10 x 6", "18G", "ACCESSORY-X", ".", "SYS-Y", 2, "ft", "NO"])
    other = book.create_sheet("DATA")
    other.append(HEADERS)
    other.append(["RECT. DUCT ft.", 4, "RECT. DUCT", 0, "12X8", ".", ".", ".", "SHOULD-NOT-READ", 99, "ft", "NO"])
    book.save(path)
    book.close()


class PhiRectangularAdapterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workbook = self.root / "phi.xlsx"
        make_phi(self.workbook)

    def tearDown(self):
        self.temp.cleanup()

    def test_phi_workbook_detection(self):
        self.assertTrue(phi.detect_phi_workbook(self.workbook))

    def test_shared_headers_without_phi_subject_route_to_universal_input(self):
        path = self.root / "non_phi_manual.xlsx"
        book = openpyxl.Workbook()
        data_duct = book.active
        data_duct.title = "DATA-DUCT"
        data_duct.append(HEADERS)
        data_duct.append(["RECT. DUCT m", 2, "RECT. DUCT", 0, "200X150", ".", ".", "LOW PRESSURE", "SYS-X", 7.5, "m", "NO"])
        data = book.create_sheet("DATA")
        data.append(HEADERS)
        data.append(["RECT. DUCT m", 2, "RECT. DUCT", 0, "200X150", ".", ".", "LOW PRESSURE", "SYS-X", 7.5, "m", "NO"])
        book.save(path)
        book.close()

        self.assertFalse(phi.detect_phi_workbook(path))
        result = preview_input(path)
        self.assertEqual(result["format"], "Excel workbook")
        self.assertEqual(result["category_counts"], {"rectangular": 1})
        self.assertEqual(result["records"][0]["source_sheet"], "DATA")
        self.assertEqual(result["records"][0]["source_row"], 2)

    def test_data_duct_is_selected_not_data(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual({row["source_sheet"] for row in records}, {"DATA-DUCT"})
        self.assertNotIn("SHOULD-NOT-READ", {row["system_unit"] for row in records})

    def test_only_exact_rectangular_subject_is_kept(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual(len(records), 2)
        self.assertTrue(all(row["source_subject"] == "RECT. DUCT ft." for row in records))

    def test_rectangular_size_parsing_uses_ph_i_inches(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual((records[0]["width"], records[0]["height"]), (8, 8))
        self.assertEqual((records[1]["width"], records[1]["height"]), (10, 6))
        self.assertEqual(records[0]["raw_duct_size"], "8X8")

    def test_length_comes_from_length_column(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual(records[0]["length"], 7.5)
        self.assertEqual(records[0]["count"], 0)

    def test_length_unit_is_preserved(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual(records[0]["length_unit"], "ft")

    def test_units_maps_to_system_unit(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual(records[0]["system_unit"], "SYS-X")

    def test_accessories_are_preserved_and_dot_is_missing(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual(records[0]["accessories"], ['w/ 1\" ACOUSTIC INSULATION', None])
        self.assertEqual(records[1]["accessories"], ["18G", "ACCESSORY-X"])

    def test_source_sheet_and_physical_rows_are_preserved(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertEqual([(row["source_sheet"], row["source_row"]) for row in records],
                         [("DATA-DUCT", 3), ("DATA-DUCT", 4)])

    def test_absent_type_does_not_reject_phi_and_is_not_inferred(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        self.assertIsNone(records[0]["type"])
        self.assertEqual(records[0]["pressure_material"], "LOW PRESSURE")

    def test_mapper_input_audit_keeps_source_fields(self):
        records, _ = phi.normalize_phi_workbook(self.workbook)
        internal = phi.write_mapper_input(records, self.root / "normalized.csv")
        with internal.open(encoding="utf-8-sig", newline="") as stream:
            row = next(csv.DictReader(stream))
        self.assertEqual(row["type"], "")
        self.assertEqual(row["source_sheet"], "DATA-DUCT")
        self.assertEqual(row["source_row"], "3")
        self.assertEqual(row["raw_duct_size"], "8X8")

    def test_existing_kativik_csv_upload_flow_remains_compatible(self):
        with tempfile.TemporaryDirectory() as temp_name:
            sample = ROOT / "test_input" / "stage1_second_project_sample.csv"
            upload = Path(temp_name) / "kativik_schema.csv"
            upload.write_bytes(sample.read_bytes())
            result = generate_vjr(upload, FIXED_TEMPLATE, Path(temp_name) / "jobs")
            self.assertEqual(result["records_processed"], 3)
            self.assertTrue(result["structure_preserved"])

    def test_phi_workbook_routes_through_upload_backend(self):
        result = generate_vjr(self.workbook, FIXED_TEMPLATE, self.root / "jobs")
        self.assertEqual(result["records_processed"], 2)
        self.assertEqual(result["records_with_writes"], 2)
        with Path(result["audit"]).open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row["source_sheet"] for row in rows], ["DATA-DUCT", "DATA-DUCT"])
        self.assertEqual([row["source_row"] for row in rows], ["3", "4"])
        self.assertIn("PHI CONTEMPORAIN", Path(result["report"]).read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
