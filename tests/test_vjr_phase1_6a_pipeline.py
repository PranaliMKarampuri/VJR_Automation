import hashlib
import tempfile
import unittest
import warnings
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

import spiral_damper_takeoff as damper
import vjr_phase1_5_pipeline as phase1_5
from vjr_phase1_6a_pipeline import (
    DEFAULT_PHASE6A_INPUT,
    DEFAULT_TEMPLATE,
    PipelineValidationError,
    generate,
)


class Phase1To6AIntegrationTests(unittest.TestCase):
    def test_phase6a_repeated_and_metric_diameters_aggregate_exactly(self):
        records, errors = damper.read_records(DEFAULT_PHASE6A_INPUT)
        self.assertEqual(errors, [])
        values, errors = damper.build_values(records, DEFAULT_TEMPLATE)
        self.assertEqual(errors, [])
        self.assertEqual(values, {"M16": 3, "M20": 3})

    def test_phase6a_rejects_nonmatching_diameter_without_rounding(self):
        values, errors = damper.build_values(
            [{"Diameter": "12.01", "DiameterUnit": "IN", "DamperCount": "1"}],
            DEFAULT_TEMPLATE,
        )
        self.assertEqual(values, {})
        self.assertTrue(any("does not exactly match" in error for error in errors))

    def test_phase6a_rejects_missing_damper_count(self):
        with tempfile.TemporaryDirectory(prefix="vjr-phase6g-input-") as temp_dir:
            csv_path = Path(temp_dir) / "missing_count.csv"
            csv_path.write_text(
                "Diameter,DiameterUnit,DamperCount\n12,IN,\n", encoding="utf-8"
            )
            _, errors = damper.read_records(csv_path)
        self.assertTrue(any("missing required field DamperCount" in error for error in errors))

    def test_end_to_end_phase1_5_inputs_and_phase6a_cells_are_preserved(self):
        source_hash = hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest()
        p1_values, _, _ = phase1_5.build_input_maps(
            DEFAULT_TEMPLATE,
            phase1_5.DEFAULT_PHASE2_INPUT,
            phase1_5.DEFAULT_PHASE3_INPUT,
            phase1_5.DEFAULT_PHASE4_INPUT,
            phase1_5.DEFAULT_PHASE5_INPUT,
        )
        with tempfile.TemporaryDirectory(prefix="vjr-phase1-6a-test-") as temp_dir:
            output = Path(temp_dir) / "integrated.xlsx"
            report = generate(
                DEFAULT_TEMPLATE,
                phase1_5.DEFAULT_PHASE2_INPUT,
                phase1_5.DEFAULT_PHASE3_INPUT,
                phase1_5.DEFAULT_PHASE4_INPUT,
                phase1_5.DEFAULT_PHASE5_INPUT,
                DEFAULT_PHASE6A_INPUT,
                output,
            )
            self.assertTrue(output.is_file())
            self.assertEqual(hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest(), source_hash)
            self.assertTrue(any("PHASE 6A INPUT CELLS ADDED: SPIRALE: M16=3, M20=3." in line for line in report))
            self.assertTrue(any("K:AX calculation formulas exactly" in line for line in report))
            self.assertTrue(any("NOT CHECKED: formula recalculation" in line for line in report))

            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
                source = load_workbook(DEFAULT_TEMPLATE, data_only=False)
                result = load_workbook(output, data_only=False)
            self.assertEqual(source.sheetnames, result.sheetnames)
            for sheet_name, values in p1_values.items():
                for coordinate, expected in values.items():
                    actual = result[sheet_name][coordinate].value
                    if expected is None:
                        self.assertIsNone(actual, f"{sheet_name}!{coordinate}")
                    elif isinstance(expected, Decimal):
                        self.assertEqual(float(actual), float(expected), f"{sheet_name}!{coordinate}")
                    else:
                        self.assertEqual(actual, expected, f"{sheet_name}!{coordinate}")
            spiral_sheet = result["SPIRALE"]
            self.assertEqual(spiral_sheet["M16"].value, 3)
            self.assertEqual(spiral_sheet["M20"].value, 3)
            self.assertEqual(spiral_sheet["B20"].value, 25)
            self.assertEqual(spiral_sheet["C20"].value, 1)
            self.assertEqual(spiral_sheet["D20"].value, 3)
            self.assertEqual(result["TAKE OFF"]["A20"].value, 400)
            self.assertEqual(result["ROND ET FLEX"]["F4"].value, 30)
            self.assertEqual(result["ROND ET FLEX"]["G2"].value, "Acoustique standard")

            with zipfile.ZipFile(DEFAULT_TEMPLATE) as original, zipfile.ZipFile(output) as generated:
                self.assertEqual(set(original.namelist()), set(generated.namelist()))
                expected_parts = {
                    phase1_5._sheet_part(original, "TAKE OFF"),
                    phase1_5._sheet_part(original, "SPIRALE"),
                    phase1_5._sheet_part(original, "ROND ET FLEX"),
                }
                changed = {name for name in original.namelist()
                           if original.read(name) != generated.read(name)}
                self.assertEqual(changed, expected_parts)
                self.assertIsNone(generated.testzip())

    def test_output_path_cannot_overwrite_template(self):
        with self.assertRaises(PipelineValidationError):
            generate(
                DEFAULT_TEMPLATE,
                phase1_5.DEFAULT_PHASE2_INPUT,
                phase1_5.DEFAULT_PHASE3_INPUT,
                phase1_5.DEFAULT_PHASE4_INPUT,
                phase1_5.DEFAULT_PHASE5_INPUT,
                DEFAULT_PHASE6A_INPUT,
                DEFAULT_TEMPLATE,
            )


if __name__ == "__main__":
    unittest.main()
