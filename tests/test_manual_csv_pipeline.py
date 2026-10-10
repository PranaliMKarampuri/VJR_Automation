import csv
import hashlib
import tempfile
import unittest
from pathlib import Path

import openpyxl

import manual_csv_to_vjr
import manual_to_vjr
import universal_input


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "source_files" / "Ventilation Jean Roy excel.xlsx"


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ManualCsvPipelineTests(unittest.TestCase):
    def _csv(self, path, count):
        headers = [
            "Subject", "Page Index", "Length", "Length Unit", "Count", "Label",
            "BRAND", "MODEL", "DUCT SIZE", "TYPE", "ACCESSORIES1", "ACCESSORIES2",
            "REMARK", "UNITS", "DAMPER TYPE",
        ]
        records = [
            ["RECT. DUCT M", 1, 1.25, "m", 0, "RECT. DUCT", ".", ".", "300X200", "BP", ".", ".", ".", "SYS-B", "NO"],
            ["RECT. DUCT M", 2, 2.75, "m", 0, "RECT. DUCT", ".", ".", "300X200", "BP", ".", ".", ".", "SYS-B", "NO"],
            ["RECT. DUCT M", 3, 4.5, "m", 0, "RECT. DUCT", ".", ".", "400X200", "HP", ".", ".", ".", "SYS-B", "NO"],
        ][:count]
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(records)

    def test_different_record_counts_generate_from_the_supplied_template(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            template_hash = sha256(TEMPLATE)
            result_counts = []
            for count in (2, 3):
                manual = temp / f"manual_{count}.csv"
                output = temp / f"generated_{count}.xlsx"
                report = temp / f"report_{count}.md"
                self._csv(manual, count)
                result = manual_csv_to_vjr.generate(manual, TEMPLATE, output, report)
                result_counts.append(result["records"])
                self.assertEqual(result["records"], count)
                self.assertTrue(output.is_file())
                self.assertTrue(report.is_file())
                book = openpyxl.load_workbook(output, data_only=True, read_only=True)
                template_book = openpyxl.load_workbook(TEMPLATE, read_only=True)
                self.assertEqual(book.sheetnames, template_book.sheetnames)
                template_book.close()
                takeoff = book["TAKE OFF"]
                self.assertEqual(takeoff["N4"].value, "M")
                self.assertEqual(takeoff["C20"].value, 4.0)
                self.assertEqual(takeoff["A21"].value, 400 if count == 3 else None)
                if count == 3:
                    self.assertEqual(takeoff["C21"].value, 4.5)
                book.close()
            self.assertEqual(result_counts, [2, 3])
            self.assertEqual(sha256(TEMPLATE), template_hash)

    def test_turning_vanes_quantity_two_writes_one_code_without_continuation(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            manual = temp / "turning_vanes.csv"
            output = temp / "turning_vanes.xlsx"
            report = temp / "turning_vanes_report.md"
            headers = [
                "Subject", "Page Index", "Length", "Length Unit", "Count", "Label",
                "BRAND", "MODEL", "DUCT SIZE", "TYPE", "ACCESSORIES1", "ACCESSORIES2",
                "REMARK", "UNITS", "DAMPER TYPE",
            ]
            rows = [
                ["RECT. DUCT m", 1, 1, "m", 0, "RECT. DUCT", ".", ".", "550X400", "BP", ".", ".", ".", "SYS-TV", "NO"],
                ["TURNING VANES", 1, 0, ".", 2, "TURNING VANES", ".", ".", "550X400", ".", ".", ".", ".", "SYS-TV", "NO"],
            ]
            with manual.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerows(rows)

            result = manual_csv_to_vjr.generate(manual, TEMPLATE, output, report)
            self.assertEqual(result["records"], 2)
            book = openpyxl.load_workbook(output, data_only=True, read_only=True)
            takeoff = book["TAKE OFF"]
            self.assertEqual([takeoff[f"{col}20"].value for col in "GHIJ"], [7, None, None, None])
            self.assertIsNone(takeoff["A21"].value)
            book.close()
            with result["audit"].open(encoding="utf-8-sig", newline="") as stream:
                audit = list(csv.DictReader(stream))
            turning_vane = next(row for row in audit if row["source_subject"] == "TURNING VANES")
            self.assertEqual(turning_vane["VJR_row"], "20")

    def test_real_53310_mvd_counts_map_to_spiral_dampers_by_system_and_diameter(self):
        source = ROOT / "test_input" / "TAKEOFF_Plans_ingenierie_53310TT-B.xlsx"
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            normalized = temp / "normalized.csv"
            universal_input.normalize_file(source, normalized)
            with normalized.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            spiral_systems = {
                (row["UNITS"].strip().upper(), row["DUCT SIZE"].strip())
                for row in rows if "SPIRAL" in row["Subject"].upper()
            }
            real_mvd = [row for row in rows
                        if "MANUAL VOLUME DAMPER" in row["Subject"].upper()
                        and (row["UNITS"].strip().upper(), row["DUCT SIZE"].strip()) in spiral_systems]
            self.assertEqual(sum(int(row["Count"]) for row in real_mvd), 72)
            mapper_input = temp / "spiral_mbd_only.csv"
            fields = list(rows[0])
            with mapper_input.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                writer.writerows(row for row in rows
                                 if "SPIRAL" in row["Subject"].upper()
                                 or "MANUAL VOLUME DAMPER" in row["Subject"].upper())
            output, report = temp / "spiral_mbd.xlsx", temp / "spiral_mbd.md"
            manual_csv_to_vjr.generate(mapper_input, TEMPLATE, output, report)
            book = openpyxl.load_workbook(output, data_only=True, read_only=True)
            spiral = book["SPIRALE"]
            self.assertEqual({f"M{row}": spiral[f"M{row}"].value for row in (12, 13, 14, 16, 20)},
                             {"M12": 10, "M13": 8, "M14": 3, "M16": 46, "M20": 5})
            self.assertEqual(sum(spiral[f"M{row}"].value or 0 for row in range(11, 64)), 72)
            book.close()

    def test_csv_engine_has_no_manual_row_or_project_name_gate(self):
        source = Path(manual_csv_to_vjr.__file__).read_text(encoding="utf-8")
        self.assertNotRegex(source, r"KATIVIK|KUUJJAAQ")
        self.assertNotRegex(source, r"source_row\s*==\s*\d+")
        self.assertTrue(callable(manual_csv_to_vjr.generate))
        self.assertTrue(callable(manual_to_vjr.resolve_manual))


if __name__ == "__main__":
    unittest.main()
