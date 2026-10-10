import csv
import hashlib
import tempfile
import unittest
from pathlib import Path
import zipfile

import openpyxl

from stage1_app_backend import FIXED_TEMPLATE, Stage1InputError, generate_vjr, preview_input
from manual_to_vjr import audit_shared_formula_package


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "test_input" / "stage1_second_project_sample.csv"


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Stage1BackendTests(unittest.TestCase):
    def test_csv_upload_uses_fixed_template_and_returns_validated_job_artifacts(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            upload = temp / "second_job.csv"
            upload.write_bytes(SAMPLE.read_bytes())
            template_hash = file_hash(FIXED_TEMPLATE)
            result = generate_vjr(upload, output_dir=temp / "jobs")
            self.assertEqual(result["records_processed"], 3)
            self.assertEqual(result["records_with_writes"], 3)
            self.assertGreater(result["cells_written"], 0)
            self.assertEqual(result["formula_changes"], 0)
            self.assertTrue(result["structure_preserved"])
            self.assertEqual(result["run_id"], result["job_id"])
            self.assertEqual(file_hash(result["output"]), result["output_sha256"])
            for key in ("output", "report", "audit", "unresolved"):
                self.assertTrue(Path(result[key]).is_file())
            self.assertTrue(result["output"].parent.is_dir())
            self.assertEqual(file_hash(FIXED_TEMPLATE), template_hash)
            self.assertIn("mapping engine", Path(result["report"]).read_text(encoding="utf-8").lower())
            report_text = Path(result["report"]).read_text(encoding="utf-8")
            self.assertIn(f"Processing run ID: `{result['run_id']}`", report_text)
            self.assertIn(f"Generated VJR SHA-256: `{result['output_sha256']}`", report_text)

    def test_rtl_rollup_bp_reaches_takeoff_type_through_upload_backend(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            source = temp / "rtl_lot1.xlsx"
            workbook = openpyxl.Workbook()
            data = workbook.active
            data.title = "DATA"
            headers = ["Subject", "Length", "Count", "Label", "Length Unit",
                       "DUCT SIZE", "UNITS", "DAMPER TYPE"]
            data.append(headers)
            detail = workbook.create_sheet("DATA-DUCT")
            detail.append(headers)
            rollup = workbook.create_sheet("TAKEOFF-DUCT")
            rollup.append(["UNITS", "Label", "DUCT SIZE", "Length Unit", "Total", "REMARK"])
            for index in range(31):
                size = f"{200 + index}X150"
                system = f"SYS-{index:02d}"
                row = ["RECT. DUCT m", 1.25, 0, "RECT. DUCT", "m", size, system, "NO"]
                data.append(row)
                detail.append(row)
                rollup.append([system, "RECT. DUCT", size, "m", 1.25, "BP"])
            workbook.save(source)
            workbook.close()

            result = generate_vjr(source, output_dir=temp / "fresh_test_outputs")
            self.assertEqual(result["detected_sheets"], ["DATA", "TAKEOFF-DUCT", "DATA-DUCT"])
            self.assertEqual(result["detected_input_records"], 31)
            self.assertEqual(result["records_processed"], 31)
            self.assertEqual(result["run_id"], result["job_id"])
            self.assertEqual(file_hash(result["output"]), result["output_sha256"])

            generated = openpyxl.load_workbook(result["output"], read_only=True, data_only=False)
            try:
                sheet = generated["TAKE OFF"]
                values = [sheet.cell(row, 5).value for row in range(20, 51)]
                self.assertEqual(values, ["BP"] * 31)
            finally:
                generated.close()

            report_text = Path(result["report"]).read_text(encoding="utf-8")
            self.assertIn(f"Processing run ID: `{result['run_id']}`", report_text)
            self.assertIn(f"Generated VJR SHA-256: `{result['output_sha256']}`", report_text)

    def test_kativik_and_53310_workbooks_still_generate_valid_current_jobs(self):
        kativik = next(p for p in (ROOT / "source_files").glob("TAKEOFF_CS KATIVIK*")
                       if p.suffix.casefold() == ".xlsx" and "_VJR" not in p.name)
        plans = ROOT / "test_input" / "TAKEOFF_Plans_ingenierie_53310TT-B.xlsx"
        rtl_normalized = ROOT / "output" / "rtl_lot1_pressure_fix" / "rtl_lot1_normalized.csv"
        with tempfile.TemporaryDirectory() as temp_name:
            output_root = Path(temp_name) / "fresh_jobs"
            for source in (kativik, plans, rtl_normalized):
                result = generate_vjr(source, output_dir=output_root)
                output = Path(result["output"])
                self.assertEqual(result["run_id"], result["job_id"])
                self.assertGreater(result["records_processed"], 0)
                self.assertEqual(file_hash(output), result["output_sha256"])
                with zipfile.ZipFile(output) as package:
                    self.assertIsNone(package.testzip())
                self.assertEqual(audit_shared_formula_package(output), [])
                workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
                try:
                    self.assertTrue({"TAKE OFF", "SPIRALE", "ROND ET FLEX"}.issubset(workbook.sheetnames))
                finally:
                    workbook.close()
                report_text = Path(result["report"]).read_text(encoding="utf-8")
                self.assertIn(f"Processing run ID: `{result['run_id']}`", report_text)
                self.assertIn(f"Generated VJR SHA-256: `{result['output_sha256']}`", report_text)

                if source == plans:
                    detection = preview_input(source)
                    selection = detection["source_selection"]
                    self.assertEqual(selection["grouped_records"], 51)
                    self.assertEqual(selection["pressure_counts"], {"BP": 51})
                    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
                    try:
                        sheet = workbook["TAKE OFF"]
                        rows = list(sheet.iter_rows(min_row=20, max_row=1004,
                                                    max_col=5, values_only=True))
                        groups = [row for row in rows if row[0] is not None and row[1] is not None
                                  and row[2] not in (None, 0)]
                        self.assertEqual(len(groups), 51)
                        self.assertTrue(all(row[4] == "BP" for row in groups))
                    finally:
                        workbook.close()

    def test_xlsx_data_sheet_is_adapted_through_the_same_csv_backend(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "DATA"
            with SAMPLE.open("r", newline="", encoding="utf-8-sig") as stream:
                for row in csv.reader(stream):
                    sheet.append(row)
            manual_xlsx = temp / "manual_excel.xlsx"
            workbook.save(manual_xlsx)
            workbook.close()
            result = generate_vjr(manual_xlsx, FIXED_TEMPLATE, temp / "jobs")
            self.assertEqual(result["records_processed"], 3)
            self.assertGreater(result["cells_written"], 0)
            self.assertTrue(Path(result["output"]).is_file())
            self.assertEqual(audit_shared_formula_package(Path(result["output"])), [])
            with Path(result["audit"]).open("r", newline="", encoding="utf-8-sig") as stream:
                audit = list(csv.DictReader(stream))
            self.assertEqual(len(audit), 3)
            self.assertIn("manual_excel.xlsx", Path(result["report"]).read_text(encoding="utf-8"))

    def test_invalid_uploads_fail_without_creating_job_output(self):
        with tempfile.TemporaryDirectory() as temp_name:
            temp = Path(temp_name)
            old_xls = temp / "old.xls"
            old_xls.write_bytes(b"not an xls file")
            with self.assertRaisesRegex(Stage1InputError, "Legacy .xls"):
                generate_vjr(old_xls, output_dir=temp / "jobs")
            malformed = temp / "malformed.csv"
            malformed.write_text("wrong,headers\n1,2\n", encoding="utf-8")
            with self.assertRaises(Stage1InputError):
                generate_vjr(malformed, output_dir=temp / "jobs")
            self.assertEqual(list((temp / "jobs").iterdir()), [])


if __name__ == "__main__":
    unittest.main()
