import csv
import tempfile
import unittest
from pathlib import Path

import universal_input
import manual_csv_to_vjr


class UniversalInputTests(unittest.TestCase):
    def test_upload_pipeline_maps_directionless_risers_and_reviews_missing_geometry(self):
        import openpyxl
        from stage1_app_backend import generate_vjr

        root = Path(__file__).resolve().parents[1]
        template = root / "source_files" / "Ventilation Jean Roy excel.xlsx"
        with tempfile.TemporaryDirectory() as d:
            temp = Path(d)
            source = temp / "directionless_risers.xlsx"
            book = openpyxl.Workbook()
            data = book.active
            data.title = "DATA"
            data.append(["Subject", "Length", "Length Unit", "Count", "Label", "DUCT SIZE", "UNITS", "REMARK"])
            data.append(["RECTANGULAR RISER", 4.5, "m", 0, "RECTANGULAR RISER", "300X200", "SYS-RI", "."])
            data.append(["RECTANGULAR RISER", 1.25, "m", 0, "RECTANGULAR RISER", "300X200", "SYS-RI", "."])
            data.append(["SPIRAL RISER", 3.25, "m", 0, "SPIRAL RISER", "200", "SYS-RI", "."])
            data.append(["RISER", 2.0, "m", 0, "RISER", "200", "SYS-RI", "."])
            data.append(["RECTANGULAR RISER", None, "m", 0, "RECTANGULAR RISER", "450X250", "SYS-RI", "."])
            data.append(["RISER", 1.0, "m", 0, "RISER", None, "SYS-RI", "."])
            book.save(source)
            book.close()

            result = generate_vjr(source, template, temp / "jobs")
            from manual_to_vjr import audit_shared_formula_package
            self.assertEqual(audit_shared_formula_package(result["output"]), [])
            generated = openpyxl.load_workbook(result["output"], data_only=True, read_only=False)
            try:
                takeoff = generated["TAKE OFF"]
                rect = [(takeoff.cell(row, 3).value, takeoff.cell(row, 6).value)
                        for row in range(20, 1005)
                        if takeoff.cell(row, 1).value == 300 and takeoff.cell(row, 2).value == 200]
                self.assertIn((5.75, "RI"), rect)

                spiral = generated["SPIRALE"]
                diameter_row = next(row for row in range(11, 64) if spiral.cell(row, 1).value == 8)
                self.assertEqual(spiral.cell(diameter_row, 2).value, 5.25)
            finally:
                generated.close()

            with Path(result["audit"]).open(encoding="utf-8-sig", newline="") as stream:
                audit = list(csv.DictReader(stream))
            self.assertEqual(len(audit), 6)
            self.assertEqual([row["category"] for row in audit], ["riser"] * 6)
            mapped = [row for row in audit if row["VJR_sheet"]]
            self.assertEqual(len(mapped), 4)
            self.assertTrue(all(row["status"] == "POPULATED" for row in mapped))
            self.assertCountEqual([row["VJR_sheet"] for row in mapped],
                                  ["TAKE OFF", "TAKE OFF", "SPIRALE", "SPIRALE"])

            with Path(result["unresolved"]).open(encoding="utf-8-sig", newline="") as stream:
                unresolved = list(csv.DictReader(stream))
            self.assertEqual(len(unresolved), 2)
            reasons = {row["reason"] for row in unresolved}
            self.assertTrue(any("project length is missing" in reason for reason in reasons))
            self.assertTrue(any("diameter is missing" in reason for reason in reasons))
            self.assertTrue(all("direction" not in reason.casefold() for reason in reasons))
            report = Path(result["report"]).read_text(encoding="utf-8")
            self.assertIn("Riser records detected: **6**; rectangular mapped: **2**; spiral mapped: **2**; requiring review: **2**.", report)

    def test_rtl_grouped_remark_populates_only_supported_pressure_types(self):
        import openpyxl
        root = Path(__file__).resolve().parents[1]
        template = root / "source_files" / "Ventilation Jean Roy excel.xlsx"
        with tempfile.TemporaryDirectory() as d:
            source, mapper, output, report = [Path(d) / name for name in
                                               ("rtl.xlsx", "normalized.csv", "generated.xlsx", "report.md")]
            book = openpyxl.Workbook()
            data = book.active
            data.title = "DATA"
            data.append(["Subject", "Length", "Count", "Label", "Length Unit", "DUCT SIZE", "UNITS", "DAMPER TYPE"])
            detail = book.create_sheet("DATA-DUCT")
            detail.append(["Subject", "Length", "Count", "Label", "Length Unit", "DUCT SIZE", "UNITS", "DAMPER TYPE"])
            rollup = book.create_sheet("TAKEOFF-DUCT")
            rollup.append(["UNITS", "Label", "DUCT SIZE", "Length Unit", "Total", "REMARK"])
            for index, (pressure, system) in enumerate((("BP", "SYS-BP"), ("MP", "SYS-MP"), ("HP", "SYS-HP"), ("review", "SYS-REVIEW")), 1):
                size = f"{200 + index}X150"
                detail.append(["RECT. DUCT m", 1.25, 0, "RECT. DUCT", "m", size, system, "NO"])
                rollup.append([system, "RECT. DUCT", size, "m", 1.25, pressure])
            book.save(source)
            book.close()

            result = universal_input.normalize_file(source, mapper)
            self.assertEqual(result["source_selection"]["authoritative_sheet"], "TAKEOFF-DUCT")
            self.assertEqual(result["source_selection"]["pressure_counts"], {"BP": 1, "MP": 1, "HP": 1, "": 1})
            with mapper.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual([row["TYPE"] for row in rows], ["BP", "MP", "HP", ""])

            manual_csv_to_vjr.generate(mapper, template, output, report)
            generated = openpyxl.load_workbook(output, data_only=True, read_only=True)
            try:
                takeoff = generated["TAKE OFF"]
                values = [takeoff.cell(row, 5).value for row in range(20, 24)]
                self.assertCountEqual(values[:3], ["BP", "MP", "HP"])
                self.assertIsNone(values[3])
            finally:
                generated.close()

    def test_missing_type_reordered_fields_normalize_rectangular_units(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source, mapper = root / "manual.csv", root / "mapper.csv"
            source.write_text('Length Unit,DUCT SIZE,Subject,Length,ACCESSORIES2,UNITS,ACCESSORIES1\nm,350X250,RECT. DUCT m,1.5,.,SYS-A,.\n', encoding="utf-8")
            detected = universal_input.normalize_file(source, mapper)
            record = detected["records"][0]
            self.assertEqual(record["category"], "rectangular")
            self.assertEqual((record["width"], record["height"]), (350, 250))
            self.assertEqual(record["length"], "1.5")
            self.assertEqual(record["length_unit"], "m")
            with mapper.open(encoding="utf-8-sig", newline="") as stream:
                row = next(csv.DictReader(stream))
            self.assertIn("TYPE", row)
            self.assertEqual(row["TYPE"], "")
            self.assertEqual(row["DUCT SIZE"], "350X250")

    def test_data_sheet_is_preferred_over_pivot_summaries(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            source, mapper = Path(d) / "manual.xlsx", Path(d) / "mapper.csv"
            book = openpyxl.Workbook()
            data = book.active
            data.title = "DATA"
            data.append(["Subject", "DUCT SIZE", "Length", "Length Unit", "Count", "Label"])
            data.append(["SPIRAL DUCT", "200", 1, "m", 0, "SPIRAL DUCT"])
            pivot = book.create_sheet("TO-DUCT")
            pivot.append(["Subject", "DUCT SIZE", "Length", "Length Unit", "Count", "Label"])
            pivot.append(["SPIRAL DUCT", "200", 99, "m", 0, "SPIRAL DUCT"])
            book.save(source)
            book.close()
            result = universal_input.normalize_file(source, mapper)
            self.assertEqual(result["sheets"], ["DATA"])
            self.assertEqual(result["total_records"], 1)

    def test_reconciled_takeoff_ducts_overrides_rectangular_detail_without_double_counting(self):
        source = Path(__file__).resolve().parents[1] / "test_input" / "TAKEOFF_Plans_ingenierie_53310TT-B.xlsx"
        with tempfile.TemporaryDirectory() as d:
            mapper = Path(d) / "mapper.csv"
            result = universal_input.normalize_file(source, mapper)
            selection = result["source_selection"]
            self.assertEqual(selection["authoritative_sheet"], "TAKEOFF-DUCTS")
            self.assertEqual(selection["detail_sheet"], "DATA-DUCTS")
            self.assertEqual(selection["detail_records"], 94)
            self.assertEqual(selection["grouped_records"], 51)
            self.assertAlmostEqual(selection["detail_length_total"], 342.5)
            self.assertAlmostEqual(selection["grouped_length_total"], 342.5)
            self.assertEqual(selection["pressure_counts"], {"BP": 51})
            with mapper.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            rectangular = [row for row in rows if row["Subject"] == "RECT. DUCT m"]
            self.assertEqual(len(rectangular), 51)
            self.assertAlmostEqual(sum(float(row["Length"]) for row in rectangular), 342.5)
            self.assertEqual({row["REMARK"] for row in rectangular}, {"BP"})
            self.assertEqual({row["TYPE"] for row in rectangular}, {"BP"})
            self.assertEqual(sum(item["field"] == "ACCESSORIES1" for item in selection["source_review_items"]), 3)


if __name__ == "__main__":
    unittest.main()
