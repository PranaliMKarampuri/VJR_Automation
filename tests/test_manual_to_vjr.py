import csv
import hashlib
import io
import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

import manual_to_vjr as writer


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


class ManualToVjrWriterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inputs = self.root / "test_input"
        self.outputs = self.root / "output"
        self.inputs.mkdir()
        self.manual = self.inputs / "sample.xlsx"
        self.template = self.inputs / "client_vjr.xlsx"
        self.mapping = self.root / "approved.csv"
        self._manual()
        self._template()

    def tearDown(self):
        self.temp.cleanup()

    def _manual(self, subject="RECT. DUCT M", size="200X300", row_type="BP", label="RECT. DUCT", length_unit="m", accessory1="."):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = "DATA"
        ws.append(["Subject", "Page Index", "Length", "Length Unit", "Count", "Label", "BRAND", "MODEL", "DUCT SIZE", "TYPE", "ACCESSORIES1", "ACCESSORIES2", "REMARK", "UNITS", "DAMPER TYPE"])
        ws.append([subject, 1, 2.0, length_unit, 0, label, ".", ".", size, row_type, accessory1, ".", ".", "UTA-1", "NO"])
        book.save(self.manual)

    def _template(self, formula_at=None, selector="M"):
        book = openpyxl.Workbook()
        takeoff = book.active
        takeoff.title = "TAKE OFF"
        for name in ("SPIRALE", "ROND ET FLEX"):
            book.create_sheet(name)
        takeoff["N4"] = selector
        for col in "ABCDEFGHIJ":
            cell = takeoff[f"{col}20"]
            cell.number_format = "0.00"
        if formula_at:
            takeoff[formula_at] = "=1+1"
        validation = DataValidation(type="list", formula1='"BP,MP,HP"')
        takeoff.add_data_validation(validation)
        validation.add(takeoff["E20"])
        book.save(self.template)

    def _manifest(self, values=None, row=20, source_row=2):
        values = values or {
            f"A{row}": 300, f"B{row}": 200, f"C{row}": 2.0,
            f"D{row}": 26, f"E{row}": "BP", f"F{row}": "HA",
        }
        with self.mapping.open("w", newline="", encoding="utf-8-sig") as stream:
            fields = list(writer.MANIFEST_HEADERS)
            output = csv.DictWriter(stream, fieldnames=fields)
            output.writeheader()
            output.writerow({
                "manifest_scope": "TEST_ONLY",
                "manual_source_row": source_row,
                "estimator_decision": "SAFE_TO_AUTOMATE",
                "estimator_confirmed_by": "Test Estimator",
                "estimator_confirmation_date": "2026-10-03",
                "ppt_rule": "rectangular",
                "target_sheet": "TAKE OFF",
                "target_row": row,
                "dimension_unit": "mm",
                "write_values_json": json.dumps(values),
            })

    def test_input_discovery_and_filename_resolution(self):
        self.assertEqual(writer.discover_inputs(self.inputs), [self.manual])
        self.assertEqual(writer.resolve_manual("sample.xlsx", self.inputs), self.manual.resolve())

    def test_rejects_missing_data_sheet_or_required_columns(self):
        book = openpyxl.Workbook()
        book.save(self.manual)
        with self.assertRaisesRegex(writer.WriterError, "DATA"):
            writer.inspect_manual(self.manual)

    def test_no_manifest_writes_no_cells_and_copies_template_byte_for_byte(self):
        before_input, before_template = file_hash(self.manual), file_hash(self.template)
        result = writer.run(self.manual, self.template, output_dir=self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 0)
        self.assertEqual(result["counts"]["NEEDS_CONFIRMATION"], 1)
        self.assertEqual(result["written_cells"], 0)
        self.assertEqual(file_hash(result["output"]), before_template)
        self.assertEqual(file_hash(self.manual), before_input)

    def test_writes_only_signed_ppt_checked_empty_input_cells_to_new_copy(self):
        self._manifest()
        before_input, before_template = file_hash(self.manual), file_hash(self.template)
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 1)
        self.assertEqual(result["written_cells"], 6)
        output = openpyxl.load_workbook(result["output"], data_only=False)
        self.assertEqual([output["TAKE OFF"][f"{col}20"].value for col in "ABCDEF"], [300, 200, 2, 26, "BP", "HA"])
        self.assertEqual(len(output["TAKE OFF"].data_validations.dataValidation), 1)
        output.close()
        self.assertEqual(file_hash(self.manual), before_input)
        self.assertEqual(file_hash(self.template), before_template)

    def test_template_markup_compatibility_prefixes_remain_declared_after_patch(self):
        compatible_template = self.inputs / "compatible_client_vjr.xlsx"
        output_path = self.outputs / "compatible_output.xlsx"
        with zipfile.ZipFile(self.template) as source, zipfile.ZipFile(compatible_template, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    marker = b"<worksheet "
                    self.assertIn(marker, data)
                    data = data.replace(
                        marker,
                        marker + b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                        b'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" '
                        b'mc:Ignorable="x14ac" ', 1)
                    data = data.replace(b"<sheetFormatPr ", b'<sheetFormatPr x14ac:dyDescent="0.25" ', 1)
                target.writestr(info, data)
        changed = writer._patch_copy(compatible_template, output_path, {"TAKE OFF": {"A20": 300}})
        with zipfile.ZipFile(output_path) as output:
            for part in changed:
                xml = output.read(part)
                root = ET.fromstring(xml)
                ignorable = root.attrib.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "")
                declared = {prefix for _, (prefix, _) in ET.iterparse(io.BytesIO(xml), events=("start-ns",))}
                self.assertEqual(ignorable, "x14ac")
                self.assertIn("x14ac", declared)

    def test_writes_materialize_all_five_overwritten_shared_formula_masters(self):
        output_path = self.outputs / "shared_formula_masters_output.xlsx"
        writes = {"TAKE OFF": {
            "A546": 450, "B546": 300, "D546": None, "E546": "BP", "F546": "HA",
        }}
        changed = writer._patch_copy(writer.DEFAULT_TEMPLATE, output_path, writes)
        self.assertIn("xl/worksheets/sheet2.xml", changed)
        self.assertEqual(writer.audit_shared_formula_package(output_path), [])
        with zipfile.ZipFile(output_path) as package:
            root = ET.fromstring(package.read("xl/worksheets/sheet2.xml"))
            cells = {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}
            for address in ("A546", "B546", "D546", "E546", "F546"):
                self.assertIsNone(cells[address].find("{*}f"))
            for address, formula in {
                "A553": 'IF(A552="","",IF(C553="","",A552))',
                "B553": 'IF(B552="","",IF(C553="","",B552))',
                "D553": 'IF(A553="","",IF(D552="","",D552))',
                "E553": 'IF(A553="","",IF(E552="","",E552))',
                "F553": 'IF(A553="","",IF(F552="","",F552))',
            }.items():
                formula_node = cells[address].find("{*}f")
                self.assertEqual(formula_node.text, formula)
                self.assertNotIn("t", formula_node.attrib)
                self.assertNotIn("si", formula_node.attrib)
            # An unrelated pre-existing shared-formula group remains shared.
            untouched = cells["A100"].find("{*}f")
            self.assertEqual(untouched.attrib.get("t"), "shared")
            self.assertEqual(untouched.attrib.get("ref"), "A100:A157")
            self.assertEqual(cells["A101"].find("{*}f").attrib.get("si"), untouched.attrib.get("si"))

        result = openpyxl.load_workbook(output_path, data_only=False)
        try:
            self.assertEqual([result["TAKE OFF"][cell].value for cell in ("A546", "B546", "D546", "E546", "F546")],
                             [450, 300, None, "BP", "HA"])
            self.assertEqual(result["TAKE OFF"]["A553"].value, '=IF(A552="","",IF(C553="","",A552))')
            self.assertEqual(result["TAKE OFF"]["F553"].value, '=IF(A553="","",IF(F552="","",F552))')
        finally:
            result.close()

    def test_write_to_shared_formula_follower_preserves_master_and_other_followers(self):
        output_path = self.outputs / "shared_formula_follower_output.xlsx"
        writer._patch_copy(writer.DEFAULT_TEMPLATE, output_path, {"TAKE OFF": {"A547": 777}})
        self.assertEqual(writer.audit_shared_formula_package(output_path), [])
        with zipfile.ZipFile(output_path) as package:
            root = ET.fromstring(package.read("xl/worksheets/sheet2.xml"))
            cells = {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}
            self.assertEqual(cells["A547"].findtext("{*}v"), "777")
            self.assertIsNone(cells["A546"].find("{*}f").attrib.get("t"))
            self.assertEqual(cells["A546"].findtext("{*}f"), 'IF(A545="","",IF(C546="","",A545))')
            self.assertEqual(cells["A548"].findtext("{*}f"), 'IF(A547="","",IF(C548="","",A547))')

    def test_new_layout_shared_formula_group_in_arbitrary_column_is_materialized(self):
        template = self.inputs / "new_layout_template.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "TAKE OFF"
        spiral = workbook.create_sheet("SPIRALE")
        workbook.create_sheet("ROND ET FLEX")
        for row in range(20, 23):
            spiral[f"H{row}"] = "=ROW()"
        workbook.save(template)
        workbook.close()
        staged = self.inputs / "new_layout_shared_template.xlsx"
        with zipfile.ZipFile(template) as source, zipfile.ZipFile(staged, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet2.xml":
                    root = ET.fromstring(data)
                    cells = {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}
                    master = cells["H20"].find("{*}f")
                    master.attrib.update({"t": "shared", "ref": "H20:H22", "si": "900"})
                    for address in ("H21", "H22"):
                        formula = cells[address].find("{*}f")
                        formula.attrib.clear()
                        formula.attrib.update({"t": "shared", "si": "900"})
                        formula.text = None
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, data)
        output_path = self.outputs / "new_layout_output.xlsx"
        writer._patch_copy(staged, output_path, {"SPIRALE": {"H21": 99}})
        self.assertEqual(writer.audit_shared_formula_package(output_path), [])
        result = openpyxl.load_workbook(output_path, data_only=False)
        try:
            self.assertEqual(result["SPIRALE"]["H20"].value, "=ROW()")
            self.assertEqual(result["SPIRALE"]["H21"].value, 99)
            self.assertEqual(result["SPIRALE"]["H22"].value, "=ROW()")
        finally:
            result.close()

    def test_formula_target_blocks_entire_record(self):
        self._template(formula_at="A20")
        self._manifest()
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 0)
        self.assertEqual(result["counts"]["NEEDS_CONFIRMATION"], 1)
        self.assertEqual(file_hash(result["output"]), file_hash(self.template))

    def test_test_only_manifest_is_rejected_outside_test_mode(self):
        self._manifest()
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=False)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 0)
        self.assertEqual(file_hash(result["output"]), file_hash(self.template))

    def test_exact_36_inches_stays_needs_confirmation(self):
        self._manual(subject="RECT. DUCT FT", size="36X20", length_unit="ft")
        self._template(selector="A")
        self._manifest({"A20": 36, "B20": 20, "C20": 2.0, "D20": 22, "E20": "BP", "F20": "RE"})
        with self.mapping.open(encoding="utf-8-sig", newline="") as stream:
            values = next(csv.DictReader(stream))
        values["dimension_unit"] = "inch"
        with self.mapping.open("w", newline="", encoding="utf-8-sig") as stream:
            output = csv.DictWriter(stream, fieldnames=list(writer.MANIFEST_HEADERS))
            output.writeheader()
            output.writerow(values)
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 0)
        self.assertEqual(result["counts"]["NEEDS_CONFIRMATION"], 1)

    def test_wrong_accessory_code_is_not_written(self):
        self._manifest({"A20": 300, "B20": 200, "C20": 2.0, "D20": 26, "E20": "BP", "F20": "HA", "G20": 999})
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 0)
        self.assertEqual(file_hash(result["output"]), file_hash(self.template))

    def test_explicit_one_inch_acoustic_insulation_uses_ppt_code_21(self):
        self._manual(accessory1='w/ 1" ACOUSTIC INSULATION')
        self._manifest({"A20": 300, "B20": 200, "C20": 2.0, "D20": 26, "E20": "BP", "F20": "HA", "G20": 21})
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["SAFE_TO_AUTOMATE"], 1)
        book = openpyxl.load_workbook(result["output"], data_only=False)
        self.assertEqual(book["TAKE OFF"]["G20"].value, 21)
        book.close()

    def test_unsupported_source_record_is_never_promoted_or_written(self):
        self._manual(subject="VENT CAP", size="150", row_type=".", label="VENT CAP")
        self._manifest()
        result = writer.run(self.manual, self.template, self.mapping, self.outputs, test_mode=True)
        self.assertEqual(result["counts"]["UNSUPPORTED"], 1)
        self.assertEqual(result["written_cells"], 0)
        self.assertEqual(file_hash(result["output"]), file_hash(self.template))


if __name__ == "__main__":
    unittest.main()
