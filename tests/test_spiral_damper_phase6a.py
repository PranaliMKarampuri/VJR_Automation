import hashlib
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from spiral_damper_takeoff import (
    DEFAULT_TEMPLATE,
    InputValidationError,
    build_values,
    generate,
    read_records,
)


class SpiralDamperPhase6ATests(unittest.TestCase):
    def test_aggregates_repeated_exact_diameters_and_writes_only_m(self):
        records = [
            {"Diameter": "12", "DiameterUnit": "IN", "DamperCount": "2"},
            {"Diameter": "304.8", "DiameterUnit": "MM", "DamperCount": "1"},
            {"Diameter": "8", "DiameterUnit": "IN", "DamperCount": "3"},
        ]
        values, errors = build_values(records, DEFAULT_TEMPLATE)
        self.assertEqual(errors, [])
        self.assertEqual(values, {"M16": 3, "M20": 3})

    def test_missing_required_fields_are_reported(self):
        with tempfile.TemporaryDirectory(prefix="vjr-damper-input-") as temporary:
            path = Path(temporary) / "missing.csv"
            path.write_text("Diameter,DiameterUnit,DamperCount\n12,IN,\n", encoding="utf-8")
            _, errors = read_records(path)
        self.assertTrue(any("missing required field DamperCount" in error for error in errors))

    def test_nonmatching_diameter_is_rejected_without_rounding(self):
        values, errors = build_values(
            [{"Diameter": "12.01", "DiameterUnit": "IN", "DamperCount": "1"}],
            DEFAULT_TEMPLATE,
        )
        self.assertEqual(values, {})
        self.assertTrue(any("does not exactly match" in error for error in errors))

    def test_end_to_end_preserves_template_except_m_input_cells(self):
        template_hash = hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory(prefix="vjr-phase6a-") as temporary:
            output = Path(temporary) / "spiral_damper.xlsx"
            report = generate(DEFAULT_TEMPLATE, Path("data/spiral_damper_phase6a_test_cases.csv"), output)
            self.assertTrue(output.is_file())
            self.assertTrue(any("wrote only SPIRALE!M input cells: M16=3, M20=3" in line for line in report))
            self.assertTrue(any("formula recalculation was not performed" in line for line in report))
            self.assertEqual(hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest(), template_hash)

            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
                source = load_workbook(DEFAULT_TEMPLATE, data_only=False)
                result = load_workbook(output, data_only=False)
            self.assertEqual(source.sheetnames, result.sheetnames)
            source_spiral, result_spiral = source["SPIRALE"], result["SPIRALE"]
            self.assertEqual(result_spiral["M16"].value, 3)
            self.assertEqual(result_spiral["M20"].value, 3)
            self.assertEqual(result_spiral["B20"].value, source_spiral["B20"].value)
            self.assertEqual(result_spiral["L20"].value, source_spiral["L20"].value)
            self.assertEqual(source_spiral["M20"].style_id, result_spiral["M20"].style_id)
            self.assertEqual(set(map(str, source_spiral.merged_cells.ranges)),
                             set(map(str, result_spiral.merged_cells.ranges)))
            self.assertEqual(len(source_spiral._images), len(result_spiral._images))
            with zipfile.ZipFile(DEFAULT_TEMPLATE) as original, zipfile.ZipFile(output) as generated:
                changed = [name for name in original.namelist() if original.read(name) != generated.read(name)]
                self.assertEqual(changed, ["xl/worksheets/sheet3.xml"])
                self.assertIsNone(generated.testzip())


if __name__ == "__main__":
    unittest.main()
