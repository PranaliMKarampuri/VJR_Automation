import hashlib
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from vjr_phase1_5_pipeline import (
    DEFAULT_PHASE2_INPUT,
    DEFAULT_PHASE3_INPUT,
    DEFAULT_PHASE4_INPUT,
    DEFAULT_PHASE5_INPUT,
    DEFAULT_TEMPLATE,
    TAKEOFF,
    SPIRAL,
    ROUND_FLEX,
    generate,
)


class Phase1To5IntegrationTests(unittest.TestCase):
    def test_end_to_end_integrated_workbook_and_preservation(self):
        source_hash = hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory(prefix="vjr-phase1-5-") as temporary:
            output = Path(temporary) / "integrated.xlsx"
            report = generate(
                DEFAULT_TEMPLATE,
                DEFAULT_PHASE2_INPUT,
                DEFAULT_PHASE3_INPUT,
                DEFAULT_PHASE4_INPUT,
                DEFAULT_PHASE5_INPUT,
                output,
            )
            self.assertTrue(output.is_file())
            self.assertEqual(hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest(), source_hash)
            self.assertTrue(any("K:AX calculation formulas" in line for line in report))
            self.assertTrue(any("NOT CHECKED: formula recalculation" in line for line in report))

            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
                workbook = load_workbook(output, data_only=False, read_only=False)
                source = load_workbook(DEFAULT_TEMPLATE, data_only=False, read_only=False)
            self.assertEqual(source.sheetnames, workbook.sheetnames)
            takeoff = workbook[TAKEOFF]
            # Four Phase 2 rows followed by three Phase 5 rectangular riser rows.
            self.assertEqual(takeoff["A20"].value, 400)
            self.assertEqual(takeoff["F20"].value, "HA")
            self.assertEqual(takeoff["C21"].value, 0)
            self.assertEqual(takeoff["F22"].value, "RE")
            self.assertEqual(takeoff["F23"].value, "RI")
            self.assertEqual(takeoff["A24"].value, 300)
            self.assertEqual(takeoff["F24"].value, "RI")
            self.assertEqual(takeoff["C26"].value, 0)
            self.assertEqual(takeoff["N4"].value, "M")
            self.assertEqual([takeoff[f"F{row}"].value for row in range(20, 27)],
                             ["HA", "HA", "RE", "RI", "RI", "RI", "RI"])
            self.assertEqual(takeoff["K20"].value, source[TAKEOFF]["K20"].value)
            self.assertEqual(takeoff["AX26"].value, source[TAKEOFF]["AX26"].value)

            spiral = workbook[SPIRAL]
            self.assertEqual(spiral["B16"].value, 27)
            self.assertEqual(spiral["B20"].value, 25)
            self.assertEqual(spiral["C20"].value, 1)
            self.assertEqual(spiral["D20"].value, 3)
            self.assertAlmostEqual(spiral["B24"].value, 6 / 0.3048)
            self.assertEqual(spiral["AA20"].value, source[SPIRAL]["AA20"].value)

            round_flex = workbook[ROUND_FLEX]
            self.assertEqual(round_flex["A19"].value, 25)
            self.assertEqual(round_flex["A21"].value, 35)
            self.assertEqual(round_flex["F4"].value, 30)
            self.assertEqual(round_flex["G4"].value, 10)
            self.assertEqual(round_flex["J4"].value, 20)
            self.assertEqual(round_flex["K4"].value, 8)
            self.assertEqual(round_flex["H4"].value, source[ROUND_FLEX]["H4"].value)

            with zipfile.ZipFile(DEFAULT_TEMPLATE) as original_zip, zipfile.ZipFile(output) as output_zip:
                self.assertEqual(set(original_zip.namelist()), set(output_zip.namelist()))
                changed = [name for name in original_zip.namelist() if original_zip.read(name) != output_zip.read(name)]
                self.assertEqual(set(changed), {
                    "xl/worksheets/sheet2.xml",
                    "xl/worksheets/sheet3.xml",
                    "xl/worksheets/sheet4.xml",
                })
                self.assertIsNone(output_zip.testzip())

            for sheet_name in (TAKEOFF, SPIRAL, ROUND_FLEX):
                self.assertEqual(
                    set(map(str, source[sheet_name].merged_cells.ranges)),
                    set(map(str, workbook[sheet_name].merged_cells.ranges)),
                )
                self.assertEqual(len(source[sheet_name]._images), len(workbook[sheet_name]._images))
                for row in source[sheet_name].iter_rows():
                    for cell in row:
                        self.assertEqual(cell.style_id, workbook[sheet_name][cell.coordinate].style_id)


if __name__ == "__main__":
    unittest.main()
