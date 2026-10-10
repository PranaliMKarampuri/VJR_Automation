"""Read-only Manual Team to VJR proposal generator (no workbook writer)."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl

import vjr_ppt_rules as rules


ROOT = Path(__file__).resolve().parent
MANUAL_DEFAULT = ROOT / "source_files" / "TAKEOFF_CS KATIVIK - CENTRE D'ÉDUCATION AUX ADULTES.xlsx"
VJR_DEFAULT = ROOT / "source_files" / "TAKEOFF_CS KATIVIK - CENTRE D'ÉDUCATION AUX ADULTES_VJR.xlsx"
MAPPING_DEFAULT = ROOT / "reports" / "manual_to_vjr_mapping.csv"
CROSSWALK_DEFAULT = ROOT / "reports" / "kuujjuaq_rectangular_estimator_crosswalk.csv"
DECISIONS_DEFAULT = ROOT / "reports" / "kuujjuaq_estimator_decision_package.csv"
OUTPUT_CSV_DEFAULT = ROOT / "reports" / "kuujjuaq_manual_to_vjr_dry_run.csv"
OUTPUT_MD_DEFAULT = ROOT / "reports" / "kuujjuaq_manual_to_vjr_dry_run.md"
OUTPUT_VALIDATION_DEFAULT = ROOT / "output" / "kuujjuaq_stage1_validation_report.txt"

FIELDS = [
    "manual_source_sheet", "manual_source_row", "manual_system_unit", "manual_duct_size",
    "manual_length", "manual_unit", "manual_type", "manual_accessories",
    "vjr_candidate_sheet", "vjr_candidate_row", "vjr_destination", "proposed_width",
    "proposed_height", "proposed_length", "proposed_metal", "proposed_type",
    "proposed_installation", "proposed_accessories", "rule_status", "mapping_status",
    "confidence", "source_evidence", "reason", "estimator_confirmation_required",
    "rule_category", "mapping_basis", "record_type", "manual_count", "candidate_row_protection",
]

SAFE_TO_AUTOMATE = "SAFE_TO_AUTOMATE"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _present(value: Any) -> bool:
    return value is not None and str(value).strip() not in {"", ".", "-", "N/A", "NA"}


def _json(values: Any) -> str:
    return json.dumps(values, ensure_ascii=False, default=str, separators=(",", ":"))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() in {"", ".", "-"}:
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def classify(subject: Any, label: Any, record_type: Any = "") -> str:
    value = " ".join(str(x or "") for x in (subject, label, record_type)).upper()
    if "RECT" in value and "DUCT" in value:
        return "rectangular"
    if "SPIRAL" in value:
        return "spiral"
    if "RISER" in value:
        return "riser"
    if "45" in value or "90" in value or "ELBOW" in value:
        return "elbow"
    if "VANE" in value:
        return "T-vane"
    if "DAMPER" in value or any(code in value.split() for code in ("RCF", "RCFF", "VCF", "RM", "BDD", "CD")):
        return "damper"
    if "CANEVAS" in value or "CANVAS" in value:
        return "canvas"
    if "INSULATION" in value or "ACOUSTIC" in value:
        return "insulation"
    if "FLEX" in value:
        return "flex"
    if "MESH" in value:
        return "other/unknown"
    return "other/unknown"


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def _decision_ids(decisions: list[dict[str, str]]) -> dict[int, list[str]]:
    result: dict[int, list[str]] = defaultdict(list)
    for decision in decisions:
        token = decision.get("manual_source_rows", "")
        if re.search(r"all\s+203\s+rectangular\s+manual\s+data\s+rows", token, re.I):
            result[0].append(decision.get("decision_id", ""))
            continue
        for row_number in re.findall(r"\b\d+\b", token):
            result[int(row_number)].append(decision.get("decision_id", ""))
    return result


def _candidate_rows(*texts: str) -> list[int]:
    found: set[int] = set()
    for text in texts:
        value = text or ""
        if "!" in value:
            # Accept exact row ranges such as TAKE OFF!A26:C26; reject
            # intervals such as SPIRALE!D11:D63 as candidates.
            for match in re.finditer(r"(?:TAKE\s*OFF|SPIRALE|ROND\s*ET\s*FLEX)!\$?[A-Z]{1,3}\$?(\d+)(?::\$?[A-Z]{1,3}\$?(\d+))?", value, re.I):
                if match.group(2) is None or match.group(1) == match.group(2):
                    found.add(int(match.group(1)))
        else:
            # Crosswalk candidate-row lists are plain semicolon-separated integers.
            for match in re.finditer(r"(?:^|[; ,])\s*(\d{1,4})\s*(?=$|[; ,])", value):
                found.add(int(match.group(1)))
    return sorted(found)


def _accessory_values(row: tuple[Any, ...]) -> dict[str, Any]:
    return {"manual_accessory_1": row[10], "manual_accessory_2": row[11], "manual_remark": row[12], "manual_damper_type": row[14]}


def _pressure_from_explicit_text(value: Any) -> str | None:
    text = str(value or "").strip().upper()
    return text if text in {"BP", "MP", "HP"} else None


def _rule_result(category: str, row: tuple[Any, ...], size: Any, system: Any, accessory_values: dict[str, Any]) -> tuple[str, str, str]:
    subject, length, unit, count, label, typ, remark = row[0], row[2], row[3], row[4], row[5], row[9], row[12]
    if category == "rectangular":
        match = re.fullmatch(r"\s*([\d.]+)\s*[xX×]\s*([\d.]+)\s*", str(size or ""))
        if not match:
            return rules.NEEDS_CONFIRMATION, "rectangular_record", "Rectangular size could not be parsed into two explicit dimensions."
        material = remark if _present(remark) and re.search(r"GAUGE|STEEL|ALUMIN", str(remark), re.I) else None
        explicit_accessory = accessory_values["manual_accessory_1"]
        accessories = []
        if isinstance(explicit_accessory, str) and "1\"" in explicit_accessory and "ACOUSTIC" in explicit_accessory.upper():
            accessories = [21]
        result = rules.rectangular_record(width=match.group(1), height=match.group(2), length=length,
            unit=_text(unit), metal_material_code=material,
            pressure=_pressure_from_explicit_text(typ) or _pressure_from_explicit_text(remark),
            system=system if _present(system) else None, accessories=accessories)
        return result.status, "rectangular_record", result.reason
    if category == "spiral":
        diameter, parsed_length = _decimal(size), _decimal(length)
        if diameter is None or parsed_length is None:
            return rules.NEEDS_CONFIRMATION, "spiral_length_by_diameter", "Spiral diameter or length is missing or nonnumeric."
        if _text(unit).lower() == "ft":
            diameter_in = diameter / Decimal("25.4")
            result = rules.spiral_length_by_diameter([(diameter_in, parsed_length)], "ft")
        elif _text(unit).lower() == "m":
            # DATA length unit is meters; source spiral diameter is the explicit millimeter dimension.
            diameter_in = diameter / Decimal("25.4")
            result = rules.spiral_length_by_diameter([(diameter_in, parsed_length)], "m")
        else:
            return rules.NEEDS_CONFIRMATION, "spiral_length_by_diameter", "Spiral length unit is not documented on this Manual record."
        return result.status, "spiral_length_by_diameter", result.reason
    if category == "elbow":
        angle = 45 if "45" in str(subject or label) else 90 if "90" in str(subject or label) else None
        diameter = _decimal(size)
        if angle is None or diameter is None or _decimal(count) is None:
            return rules.NEEDS_CONFIRMATION, "spiral_elbows", "Elbow angle, diameter, or count is missing."
        # Source DATA duct sizes are in mm; preserve exact conversion, with no nominal rounding.
        return rules.spiral_elbows(angle, int(_decimal(count)), diameter / Decimal("25.4")).status, "spiral_elbows", "Apply the PPT elbow-count rule; exact diameter correspondence to a VJR row still requires confirmation."
    if category == "flex":
        qty = _decimal(count)
        if qty is None:
            return rules.NEEDS_CONFIRMATION, "flex_length", "Flex count is missing."
        result = rules.flex_length(int(qty))
        return result.status, "flex_length", result.reason
    if category == "canvas":
        result = rules.canvas_eligibility(direct_to_fan_ahu_rtu=None, kitchen_hood_or_dryer=None)
        return result.status, "canvas_eligibility", result.reason
    if category == "damper":
        return rules.NEEDS_CONFIRMATION, "mbd", "Damper type/destination needs explicit mapping; do not infer from nearby duct topology."
    if category == "riser":
        return rules.NEEDS_CONFIRMATION, "rectangular_riser/spiral_riser_drop", "Riser orientation, geometry, length unit, and duplicate status require confirmation from row-specific source evidence."
    if category == "T-vane":
        return rules.NEEDS_CONFIRMATION, "accessory", "T-vane association must be explicit for this duct; no size-only association is applied."
    return rules.UNSUPPORTED, "special_item", f"No supported automatic destination for source item {subject!r} / {label!r}."


def _occupied_vjr_rows(vjr_path: Path) -> dict[tuple[str, int], bool]:
    book = openpyxl.load_workbook(vjr_path, read_only=True, data_only=False)
    out: dict[tuple[str, int], bool] = {}
    for sheet_name in ("TAKE OFF", "SPIRALE", "ROND ET FLEX"):
        if sheet_name not in book.sheetnames:
            continue
        sheet = book[sheet_name]
        for r, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            out[(sheet_name, r)] = any(v is not None and v != "" for v in values)
    book.close()
    return out


def _protection_status(rows: list[int], sheet: str, occupied: dict[tuple[str, int], bool]) -> str:
    if not rows:
        return "No exact candidate row identified; no insertion row is proposed."
    targets = sheet.split(" / ") if sheet else []
    descriptions = []
    for row in rows:
        states = [occupied.get((target.strip(), row), False) for target in targets]
        descriptions.append(f"row {row}: existing data/formula present" if any(states) else f"row {row}: no data/formula detected")
    return "; ".join(descriptions) + ". Candidate rows are references only and will never be overwritten by this dry-run."


def run_dry_run(manual_path: Path = MANUAL_DEFAULT, vjr_path: Path = VJR_DEFAULT,
                mapping_path: Path = MAPPING_DEFAULT, crosswalk_path: Path = CROSSWALK_DEFAULT,
                decisions_path: Path = DECISIONS_DEFAULT, output_csv: Path = OUTPUT_CSV_DEFAULT,
                output_md: Path = OUTPUT_MD_DEFAULT) -> dict[str, Any]:
    protected_paths = {Path(p).resolve() for p in (manual_path, vjr_path)}
    if Path(output_csv).resolve() in protected_paths or Path(output_md).resolve() in protected_paths:
        raise ValueError("Dry-run outputs cannot target either source workbook.")

    old_mapping = {int(x["Source Row"]): x for x in _read_csv(mapping_path)}
    crosswalk = {int(x["manual_source_row"]): x for x in _read_csv(crosswalk_path)}
    decisions = _decision_ids(_read_csv(decisions_path))
    occupied = _occupied_vjr_rows(Path(vjr_path))

    manual = openpyxl.load_workbook(manual_path, read_only=True, data_only=True)
    if "DATA" not in manual.sheetnames:
        raise ValueError("Manual workbook has no DATA sheet.")
    sheet = manual["DATA"]
    records: list[dict[str, Any]] = []
    raw_length_total = Decimal(0)
    length_by_unit: dict[str, Decimal] = defaultdict(Decimal)
    counts_by_category: Counter[str] = Counter()
    lengths_by_category: dict[str, Decimal] = defaultdict(Decimal)
    counts_by_system: Counter[str] = Counter()
    lengths_by_system: dict[str, Decimal] = defaultdict(Decimal)
    source_sheet_counts: Counter[str] = Counter()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None for v in row):
            continue
        subject, page, length, unit, count, label, brand, model, size, typ, acc1, acc2, remark, system, damper_type = row[:15]
        legacy = old_mapping.get(row_number, {})
        cw = crosswalk.get(row_number, {})
        category = classify(subject, label, legacy.get("Record Type", ""))
        length_num = _decimal(length)
        if length_num is not None:
            raw_length_total += length_num
            length_by_unit[_text(unit) if _present(unit) else "(blank/placeholder)"] += length_num
        counts_by_category[category] += 1
        lengths_by_category[category] += length_num or Decimal(0)
        source_sheet_counts["DATA"] += 1
        if _present(system):
            counts_by_system[str(system)] += 1
            lengths_by_system[str(system)] += length_num or Decimal(0)

        rule_status, rule_name, rule_reason = _rule_result(category, row, size, system, _accessory_values(row))
        # A supported PPT rule alone does not authorize a workbook write.
        if rule_status == rules.SAFE:
            rule_status = "SUPPORTED_BY_PPT"
        crosswalk_basis = cw.get("mapping_basis", "")
        legacy_match = legacy.get("Match Status", "")
        basis = crosswalk_basis or ("GROUP_ROLLUP" if legacy_match == "MATCHED_GROUP_ROLLUP" else "CANDIDATE_MATCH" if legacy_match.startswith("CANDIDATE") else "UNRESOLVED")
        old_candidate_text = legacy.get("Matching VJR Record/Cell if found", "")
        cw_candidate_text = cw.get("vjr_candidate_row", "")
        candidate_text = cw_candidate_text or old_candidate_text
        candidate_rows = _candidate_rows(candidate_text, legacy.get("VJR Cell/Column", ""))
        candidate_sheet = cw.get("vjr_candidate_sheet", "") or legacy.get("VJR Sheet", "")
        destination = ""
        if candidate_sheet:
            destination = candidate_sheet
            dest_field = legacy.get("VJR Input Field", "")
            if dest_field:
                destination += " — " + dest_field
        source_evidence = {
            "legacy_mapping_status": legacy.get("Status", ""),
            "legacy_match_status": legacy_match,
            "legacy_rule_logic": legacy.get("Rule/Logic", ""),
            "crosswalk_basis": crosswalk_basis,
            "crosswalk_evidence": cw.get("mapping_evidence", ""),
            "decision_ids": decisions.get(row_number, []) + (decisions.get(0, []) if category == "rectangular" else []),
            "vjr_candidate_reference": candidate_text,
            "ppt_rule": rule_name,
        }

        # No estimator-approved mapping exists in the supplied crosswalk/decision package.
        unsupported = rule_status == rules.UNSUPPORTED or legacy.get("Status") == "DO_NOT_AUTOMATE_YET"
        if unsupported:
            mapping_status = rules.UNSUPPORTED
        else:
            mapping_status = rules.NEEDS_CONFIRMATION
        reasons = [rule_reason]
        if basis == "GROUP_ROLLUP":
            reasons.append("GROUP_ROLLUP evidence retained as a grouped proposal; it is not an exact row mapping and cannot be safely applied row by row.")
        if candidate_rows:
            reasons.append("Existing-row candidates are comparison references only; row reuse is not authorized and no row will be overwritten.")
        else:
            reasons.append("No confirmed destination row exists; append row 68 and other insertion locations remain unapproved.")
        if row_number in crosswalk or category == "rectangular":
            reasons.append("Rectangular estimator crosswalk has no approved SAFE rows.")
        if not _present(system):
            reasons.append("Manual system/unit is blank or placeholder; none was inferred.")
        if category in {"rectangular", "T-vane", "canvas", "damper"} and not _present(acc1) and not _present(acc2):
            reasons.append("No accessory association is recorded in this row; none was inferred from duct size or adjacency.")
        # A proposed destination value is shown only when the crosswalk documents it.
        proposed_width = cw.get("vjr_width", "")
        proposed_height = cw.get("vjr_height", "")
        proposed_length = cw.get("vjr_length", "")
        proposed_metal = cw.get("vjr_metal", "")
        proposed_type = cw.get("vjr_type", "")
        proposed_install = cw.get("vjr_installation", "")
        proposed_accessories = cw.get("vjr_accessories", "")
        if not cw:
            proposed_width = proposed_height = proposed_length = proposed_metal = proposed_type = proposed_install = proposed_accessories = ""
        confirm = cw.get("estimator_confirmation_needed", "") or legacy.get("Notes", "")
        if not confirm and mapping_status == rules.NEEDS_CONFIRMATION:
            confirm = "Confirm exact destination/group, all required VJR fields, and row insertion/reuse policy before writing."
        confidence = cw.get("confidence", "") or legacy.get("Confidence", "") or "Not established"
        records.append({
            "manual_source_sheet": "DATA", "manual_source_row": row_number,
            "manual_system_unit": _text(system), "manual_duct_size": _text(size),
            "manual_length": _text(length), "manual_unit": _text(unit),
            "manual_type": _text(typ), "manual_accessories": _json(_accessory_values(row)),
            "vjr_candidate_sheet": candidate_sheet,
            "vjr_candidate_row": "; ".join(map(str, candidate_rows)),
            "vjr_destination": destination,
            "proposed_width": proposed_width, "proposed_height": proposed_height,
            "proposed_length": proposed_length, "proposed_metal": proposed_metal,
            "proposed_type": proposed_type, "proposed_installation": proposed_install,
            "proposed_accessories": proposed_accessories,
            "rule_status": rule_status, "mapping_status": mapping_status,
            "confidence": confidence, "source_evidence": _json(source_evidence),
            "reason": " ".join(reasons), "estimator_confirmation_required": confirm,
            "rule_category": category, "mapping_basis": basis,
            "record_type": legacy.get("Record Type", ""),
            "manual_count": _text(count),
            "candidate_row_protection": _protection_status(candidate_rows, candidate_sheet, occupied),
        })
    manual.close()

    with Path(output_csv).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)

    status_counts = Counter(x["mapping_status"] for x in records)
    grouped_count = sum(x["mapping_basis"] == "GROUP_ROLLUP" for x in records)
    candidate_count = sum(x["mapping_basis"] == "CANDIDATE_MATCH" for x in records)
    unresolved_count = sum(x["mapping_status"] != SAFE_TO_AUTOMATE for x in records)
    summary = {
        "total_manual_records": len(records),
        "total_manual_length": str(raw_length_total),
        "length_by_unit": {k: str(v) for k, v in sorted(length_by_unit.items())},
        "safe_count": status_counts[SAFE_TO_AUTOMATE],
        "safe_to_automate_count": status_counts[SAFE_TO_AUTOMATE],
        "needs_confirmation_count": status_counts[rules.NEEDS_CONFIRMATION],
        "unsupported_count": status_counts[rules.UNSUPPORTED],
        "grouped_rollup_count": grouped_count,
        "candidate_mapping_count": candidate_count,
        "unresolved_count": unresolved_count,
        "totals_by_category": {k: {"count": counts_by_category[k], "raw_length": str(lengths_by_category[k])} for k in sorted(counts_by_category)},
        "totals_by_source_sheet": dict(source_sheet_counts),
        "totals_by_system_unit": {k: {"count": counts_by_system[k], "raw_length": str(lengths_by_system[k])} for k in sorted(counts_by_system)},
    }
    _write_report(Path(output_md), summary, records)
    return {"summary": summary, "records": records}


def _write_report(path: Path, summary: dict[str, Any], records: list[dict[str, Any]]) -> None:
    reasons = Counter()
    for record in records:
        for clause in record["reason"].split(". "):
            if clause.strip():
                reasons[clause.strip().rstrip(".")] += 1
    lines = [
        "# Kuujjuaq Manual Team → VJR dry-run proposal",
        "",
        "## Safety boundary",
        "",
        "This report was generated by reading the Manual DATA sheet, VJR workbook, existing 661-row mapping, rectangular estimator crosswalk, and estimator decision package. The VJR workbook was opened read-only. No VJR writer is present in this pipeline. Candidate rows are comparison references only; no insertion location or row reuse is approved here.",
        "",
        "## Counts",
        "",
        f"- Manual records: **{summary['total_manual_records']}**",
        f"- Sum of numeric Manual DATA Length cells (raw values across units): **{summary['total_manual_length']}**",
        f"- Disposition: SAFE_TO_AUTOMATE **{summary['safe_to_automate_count']}**, NEEDS_CONFIRMATION **{summary['needs_confirmation_count']}**, UNSUPPORTED **{summary['unsupported_count']}**",
        f"- GROUP_ROLLUP proposals: **{summary['grouped_rollup_count']}**",
        f"- Candidate mapping proposals (excluding GROUP_ROLLUP): **{summary['candidate_mapping_count']}**",
        f"- Unresolved proposals: **{summary['unresolved_count']}**",
        "",
        "The raw length sum combines the units as recorded and is not a converted physical total. Unit subtotals follow.",
        "",
        "## Manual length subtotals by recorded unit",
        "",
        "| Unit | Raw sum |",
        "|---|---:|",
    ]
    lines.extend(f"| {unit} | {amount} |" for unit, amount in summary["length_by_unit"].items())
    lines += ["", "## Totals by rule category", "", "| Category | Records | Raw length sum |", "|---|---:|---:|"]
    lines.extend(f"| {category} | {data['count']} | {data['raw_length']} |" for category, data in summary["totals_by_category"].items())
    lines += ["", "## Totals by source sheet", "", "| Source sheet | Records |", "|---|---:|"]
    lines.extend(f"| {sheet} | {count} |" for sheet, count in summary["totals_by_source_sheet"].items())
    lines += ["", "## Totals by explicitly recorded system/unit", "", "| Manual system/unit | Records | Raw length sum |", "|---|---:|---:|"]
    lines.extend(f"| {system} | {data['count']} | {data['raw_length']} |" for system, data in summary["totals_by_system_unit"].items())
    lines += ["", "## Most frequent unresolved reason statements", ""]
    lines.extend(f"- {count} records: {reason}" for reason, count in reasons.most_common(10))
    lines += [
        "", "## Estimator/client confirmations required before VJR writing",
        "",
        "1. Approve the row insertion/reuse policy, including whether row 68 is an approved frontier; this dry-run assumes no frontier.",
        "2. Resolve the installation rule conflict for Manual rows 300/304 and VJR row 49. No HA/RE override is applied.",
        "3. Confirm the Manual rectangular dimension unit and exact system/unit associations, destinations, and group boundaries listed in the crosswalk/decision package.",
        "4. Confirm type/pressure/material, accessory association, and any unresolved classification for each row/group; blank source/VJR cells remain blank/unknown, never zero.",
        "5. Approve exact output fields and destinations for records not covered by an approved mapping, including existing client row treatment. Existing VJR rows are not reused or overwritten by this dry-run.",
        "",
        "All row-level proposals and evidence are in the companion CSV. SAFE_TO_AUTOMATE requires an approved exact destination and input-cell write manifest; the current checklist and crosswalk contain none. PPT support for an isolated rule is reported separately and never promotes a row by itself.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_validation_report(summary: dict[str, Any], records: list[dict[str, Any]],
                            source_hashes_before: dict[str, str],
                            source_hashes_after: dict[str, str], path: Path) -> None:
    row_keys = [int(record["manual_source_row"]) for record in records]
    statuses = Counter(record["mapping_status"] for record in records)
    unique_rows = len(set(row_keys)) == len(row_keys)
    coverage_ok = len(records) == 661 and unique_rows and set(row_keys) == set(range(2, 663))
    statuses_ok = set(statuses).issubset({SAFE_TO_AUTOMATE, rules.NEEDS_CONFIRMATION, rules.UNSUPPORTED})
    counts_ok = sum(statuses.values()) == summary["total_manual_records"]
    hashes_ok = source_hashes_before == source_hashes_after and len(source_hashes_before) == 2
    no_writes = summary["safe_to_automate_count"] == 0
    checks = {
        "Manual source row coverage and unique keys": coverage_ok,
        "Disposition vocabulary and totals": statuses_ok and counts_ok,
        "No VJR write set without approved input cells": no_writes,
        "Manual and VJR source workbook SHA-256 unchanged": hashes_ok,
    }
    lines = [
        "Kuujjuaq Stage 1 validation",
        "",
        "Result: " + ("PASS" if all(checks.values()) else "FAIL"),
        f"Manual records: {len(records)}; unique source rows: {len(set(row_keys))}",
        "Disposition counts: " + ", ".join(f"{name}={statuses.get(name, 0)}" for name in (SAFE_TO_AUTOMATE, rules.NEEDS_CONFIRMATION, rules.UNSUPPORTED)),
        "VJR workbook writes: 0 (no signed destination/input-cell manifest exists)",
        "",
        "Checks:",
    ]
    lines.extend(f"{'PASS' if passed else 'FAIL'}  {label}" for label, passed in checks.items())
    lines += ["", "Source workbook hashes (before → after):"]
    for name, digest in source_hashes_before.items():
        lines.append(f"{name}: {digest} → {source_hashes_after.get(name, 'MISSING')}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if not all(checks.values()):
        raise RuntimeError("Stage 1 validation failed; see " + str(path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manual", type=Path, default=MANUAL_DEFAULT)
    parser.add_argument("--vjr", type=Path, default=VJR_DEFAULT)
    parser.add_argument("--mapping", type=Path, default=MAPPING_DEFAULT)
    parser.add_argument("--crosswalk", type=Path, default=CROSSWALK_DEFAULT)
    parser.add_argument("--decisions", type=Path, default=DECISIONS_DEFAULT)
    parser.add_argument("--csv", type=Path, default=OUTPUT_CSV_DEFAULT)
    parser.add_argument("--report", type=Path, default=OUTPUT_MD_DEFAULT)
    parser.add_argument("--validation-report", type=Path, default=OUTPUT_VALIDATION_DEFAULT)
    args = parser.parse_args()
    source_hashes_before = {"Manual": _sha256(args.manual), "VJR": _sha256(args.vjr)}
    result = run_dry_run(args.manual, args.vjr, args.mapping, args.crosswalk, args.decisions, args.csv, args.report)
    source_hashes_after = {"Manual": _sha256(args.manual), "VJR": _sha256(args.vjr)}
    write_validation_report(result["summary"], result["records"], source_hashes_before,
                            source_hashes_after, args.validation_report)
    print(json.dumps(result["summary"], indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
