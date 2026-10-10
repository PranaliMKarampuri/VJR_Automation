#!/usr/bin/env python3
"""Populate documented rectangular and spiral riser inputs in a VJR workbook.

The Phase 5 wrapper reuses Phase 2 rectangular transformations and Phase 3
exact spiral diameter matching. It edits only the two relevant worksheet XML
parts of a copied XLSX package so workbook extensions and other parts survive.
"""

from __future__ import annotations

import argparse
import csv
import sys
import warnings
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

import rectangular_duct_takeoff as rectangular
import spiral_duct_takeoff as spiral


TAKEOFF_SHEET = "TAKE OFF"
SPIRAL_SHEET = "SPIRALE"
REQUIRED_HEADERS = (
    "RiserID", "Floor", "System", "Width", "Height", "Length", "LengthUnit",
    "MetalCode", "Type", "Units", "Diameter", "DiameterUnit",
)
NS_MAIN = rectangular.NS_MAIN
NS_REL_DOC = rectangular.NS_REL_DOC
NS_REL_PKG = rectangular.NS_REL_PKG


class InputValidationError(Exception):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("Input validation failed")


def read_records(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    records: list[dict[str, Any]] = []
    errors: list[str] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            missing = [name for name in REQUIRED_HEADERS if name not in headers]
            if missing:
                return [], ["CSV is missing required columns: " + ", ".join(missing)]
            accessory_fields = []
            for header in headers:
                if header.startswith("Accessory"):
                    suffix = header.removeprefix("Accessory")
                    if suffix.isdigit():
                        accessory_fields.append((int(suffix), header))
            accessory_fields.sort()
            if [n for n, _ in accessory_fields] != list(range(1, len(accessory_fields) + 1)):
                return [], ["Accessory columns must be consecutive, starting with Accessory1"]
            for line, raw in enumerate(reader, start=2):
                record = {key: (raw.get(key) or "").strip() for key in REQUIRED_HEADERS}
                for numeric_field in ("Width", "Height", "Length", "MetalCode", "Diameter"):
                    if record[numeric_field]:
                        record[numeric_field] = rectangular.parse_value(record[numeric_field])
                record["Accessories"] = [rectangular.parse_value(raw.get(col) or "") for _, col in accessory_fields]
                if not any(record[k] for k in REQUIRED_HEADERS) and not any(record["Accessories"]):
                    continue
                if not record["RiserID"]:
                    errors.append(f"CSV row {line}: missing required field RiserID")
                if not record["Floor"]:
                    errors.append(f"CSV row {line}: missing required field Floor")
                system = record["System"].upper()
                if system not in {"RECTANGULAR", "SPIRAL"}:
                    errors.append(f"CSV row {line}: System must be RECTANGULAR or SPIRAL")
                record["System"] = system
                if system == "RECTANGULAR":
                    for field in ("Width", "Height", "Length", "LengthUnit", "MetalCode", "Type", "Units"):
                        if not record[field]:
                            errors.append(f"CSV row {line}: missing required field {field} for rectangular riser")
                    if any(record[k] for k in ("Diameter", "DiameterUnit")):
                        errors.append(f"CSV row {line}: spiral diameter fields must be empty for rectangular risers")
                    if record["LengthUnit"].upper() != "M":
                        errors.append(f"CSV row {line}: rectangular riser LengthUnit must be M; English-unit riser conversion is deferred")
                    if record["Units"].upper() != "M":
                        errors.append(f"CSV row {line}: rectangular riser Units must be M for the documented meter length")
                    record["Installation"] = "RI"
                elif system == "SPIRAL":
                    for field in ("Diameter", "DiameterUnit", "Length", "LengthUnit"):
                        if not record[field]:
                            errors.append(f"CSV row {line}: missing required field {field} for spiral riser")
                    for field in ("Width", "Height", "MetalCode", "Type", "Units"):
                        if record[field]:
                            errors.append(f"CSV row {line}: {field} is not used for a spiral riser")
                    if record["LengthUnit"].upper() != "FT":
                        errors.append(f"CSV row {line}: spiral riser LengthUnit must be FT")
                    if any(record["Accessories"]):
                        errors.append(f"CSV row {line}: no spiral riser accessory rule is documented; accessory values are rejected")
                records.append(record)
    except OSError as exc:
        errors.append(f"Cannot read input CSV {path}: {exc}")
    if not records and not errors:
        errors.append("CSV contains no riser records")
    return records, errors


def _signature(record: dict[str, Any]) -> tuple[Any, ...]:
    ignored = {"Floor", "RiserID"}
    result: list[Any] = []
    for key in REQUIRED_HEADERS:
        if key in ignored:
            continue
        value = str(record.get(key, "")).strip()
        if key in {"System", "LengthUnit", "DiameterUnit", "Units"}:
            value = value.upper()
        result.append(value)
    result.append(tuple(str(v).strip() for v in record.get("Accessories", [])))
    return tuple(result)


def deduplicate_floor_records(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int, list[str]]:
    """Keep a riser ID once when the same documented item appears on floors."""
    errors: list[str] = []
    by_id: dict[str, dict[str, Any]] = {}
    floors_by_id: dict[str, set[str]] = defaultdict(set)
    for line, record in enumerate(records, start=1):
        key = str(record.get("RiserID", "")).strip().casefold()
        if not key:
            errors.append(f"Record {line}: RiserID is required to enforce the one-count rule")
            continue
        floor = str(record.get("Floor", "")).strip().casefold()
        prior = by_id.get(key)
        if prior is None:
            by_id[key] = record
            floors_by_id[key].add(floor)
            continue
        if _signature(prior) != _signature(record):
            errors.append(f"Record {line}: RiserID {record['RiserID']!r} appears with conflicting riser data; cannot safely count it once")
            continue
        if floor in floors_by_id[key]:
            errors.append(f"Record {line}: RiserID {record['RiserID']!r} is repeated on the same floor; confirm whether it is a duplicate")
            continue
        floors_by_id[key].add(floor)
    if errors:
        return [], 0, errors
    kept = list(by_id.values())
    return kept, len(records) - len(kept), []


def _decimal(value: Any, label: str, record: str, errors: list[str]) -> Decimal | None:
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError, AttributeError):
        errors.append(f"Riser {record}: {label} must be numeric; got {value!r}")
        return None
    if not result.is_finite():
        errors.append(f"Riser {record}: {label} must be finite")
        return None
    return result


def diameter_rows(template: Path) -> dict[Decimal, int]:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        workbook = load_workbook(template, data_only=False, read_only=False)
    if SPIRAL_SHEET not in workbook.sheetnames:
        raise InputValidationError([f"Template does not contain {SPIRAL_SHEET!r}"])
    result: dict[Decimal, int] = {}
    sheet = workbook[SPIRAL_SHEET]
    for row in range(spiral.DIAMETER_FIRST_ROW, spiral.DIAMETER_LAST_ROW + 1):
        value = sheet[f"A{row}"].value
        if value is None:
            continue
        try:
            key = Decimal(str(value))
        except InvalidOperation:
            continue
        if key.is_finite():
            result[key] = row
    return result


def build_values(records: list[dict[str, Any]], template: Path) -> tuple[dict[str, dict[str, Any]], int, list[str]]:
    deduped, suppressed, errors = deduplicate_floor_records(records)
    if errors:
        return {}, suppressed, errors
    rectangular_records: list[dict[str, Any]] = []
    spiral_records: list[dict[str, Any]] = []
    for record in deduped:
        if record["System"] == "RECTANGULAR":
            row_record = {key: record[key] for key in rectangular.REQUIRED_COLUMNS}
            row_record["Installation"] = "RI"
            row_record["Units"] = "M"
            row_record["Accessories"] = record["Accessories"]
            rectangular_records.append(row_record)
        else:
            spiral_records.append(record)

    values: dict[str, dict[str, Any]] = {TAKEOFF_SHEET: {}, SPIRAL_SHEET: {}}
    if rectangular_records:
        rect_rows, rect_errors, units = rectangular.transform_records(rectangular_records)
        errors.extend(rect_errors)
        if units != "M":
            errors.append("Rectangular riser lengths require metric project units (M)")
        for index, duct in enumerate(rect_rows):
            row = rectangular.FIRST_DATA_ROW + index
            cells = {
                "A": duct["Width"], "B": duct["Height"], "C": duct["Length"],
                "D": duct["MetalCode"], "E": duct["Type"], "F": "RI",
            }
            cells.update({column: duct["Accessories"][slot] if slot < len(duct["Accessories"]) else None
                          for slot, column in enumerate(rectangular.ACCESSORY_SLOT_COLUMNS)})
            values[TAKEOFF_SHEET].update({f"{column}{row}": val for column, val in cells.items()})
        values[TAKEOFF_SHEET]["N4"] = "M"

    if spiral_records:
        diameter_to_row = diameter_rows(template)
        totals: dict[int, Decimal] = defaultdict(Decimal)
        for record in spiral_records:
            name = record["RiserID"]
            diameter = _decimal(record["Diameter"], "Diameter", name, errors)
            length = _decimal(record["Length"], "Length", name, errors)
            if diameter is None or length is None:
                continue
            if diameter <= 0 or length < 0:
                errors.append(f"Riser {name}: Diameter must be positive and Length cannot be negative")
                continue
            unit = record["DiameterUnit"].upper()
            if unit not in {"IN", "MM"}:
                errors.append(f"Riser {name}: DiameterUnit must be IN or MM")
                continue
            diameter_in = diameter if unit == "IN" else diameter / Decimal("25.4")
            row = diameter_to_row.get(diameter_in)
            if row is None:
                errors.append(f"Riser {name}: diameter {diameter} {unit} converts to {diameter_in} inches, which does not exactly match a SPIRALE diameter key; rejected without rounding")
                continue
            totals[row] += length
        values[SPIRAL_SHEET].update({f"B{row}": total for row, total in totals.items()})
    if errors:
        return {}, suppressed, errors
    return values, suppressed, []


def _sheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels.findall(f"{{{NS_REL_PKG}}}Relationship")}
    for sheet in root.find(f"{{{NS_MAIN}}}sheets"):
        if sheet.attrib.get("name") == sheet_name:
            target = rel_map[sheet.attrib[f"{{{NS_REL_DOC}}}id"]].replace("\\", "/")
            return target.lstrip("/") if target.startswith("/") else target if target.startswith("xl/") else "xl/" + target
    raise ValueError(f"Worksheet {sheet_name!r} not found in workbook")


def _formula_map(path: Path) -> dict[tuple[str, str], str]:
    return rectangular._formula_map(path)


def generate(template: Path, csv_path: Path, output: Path) -> list[str]:
    if not template.is_file():
        raise InputValidationError([f"Template workbook not found: {template}"])
    if template.resolve() == output.resolve():
        raise InputValidationError(["Output path must not be the original template workbook"])
    records, errors = read_records(csv_path)
    if errors:
        raise InputValidationError(errors)
    values, suppressed, errors = build_values(records, template)
    if errors:
        raise InputValidationError(errors)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        workbook = load_workbook(template, data_only=False, read_only=False)
    for sheet_name, cells in values.items():
        for coordinate, value in cells.items():
            workbook[sheet_name][coordinate] = float(value) if isinstance(value, Decimal) else value

    written = {name: set(cells) for name, cells in values.items()}
    # Match the established Phase 2 handling: expand untouched carry-down
    # formulas before intentionally replacing the selected A:B/D:F input cells.
    formula_expansions: dict[str, str] = {}
    rect_sheet = workbook[TAKEOFF_SHEET]
    for row in rect_sheet.iter_rows(min_row=rectangular.FIRST_DATA_ROW, max_row=1004, min_col=1, max_col=6):
        for cell in row:
            if cell.column in (1, 2, 4, 5, 6) and cell.coordinate not in written[TAKEOFF_SHEET] and isinstance(cell.value, str) and cell.value.startswith("="):
                formula_expansions[cell.coordinate] = cell.value

    before_formulas = _formula_map(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(template, "r") as source:
        takeoff_part = _sheet_part(source, TAKEOFF_SHEET)
        spiral_part = _sheet_part(source, SPIRAL_SHEET)
        roots = {
            takeoff_part: ET.fromstring(source.read(takeoff_part)),
            spiral_part: ET.fromstring(source.read(spiral_part)),
        }
        for coordinate, formula in formula_expansions.items():
            rectangular._set_formula(roots[takeoff_part], coordinate, formula)
        for sheet_name, cells in values.items():
            part = takeoff_part if sheet_name == TAKEOFF_SHEET else spiral_part
            for coordinate, value in cells.items():
                rectangular._set_cell_value(roots[part], coordinate, value)
        patched = {part: ET.tostring(root, encoding="utf-8", xml_declaration=True) for part, root in roots.items()}
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for member in source.infolist():
                target.writestr(member, patched.get(member.filename, source.read(member.filename)))

    with zipfile.ZipFile(template, "r") as src, zipfile.ZipFile(output, "r") as dst:
        if set(src.namelist()) != set(dst.namelist()):
            raise RuntimeError("XLSX package member list changed")
        changed_parts = [name for name in sorted(src.namelist()) if src.read(name) != dst.read(name)]
        expected_parts = sorted([takeoff_part, spiral_part])
        if changed_parts != expected_parts:
            raise RuntimeError(f"Unexpected XLSX package changes: {changed_parts}")
        if dst.testzip() is not None:
            raise RuntimeError("Output XLSX ZIP integrity check failed")

    after_formulas = _formula_map(output)
    expected_protected = {key: formula for key, formula in before_formulas.items()
                          if not (key[0] in written and key[1] in written[key[0]])}
    if after_formulas != expected_protected:
        raise RuntimeError("A formula outside the explicitly written input cells changed")
    takeoff_calcs_before = rectangular._takeoff_calculation_formulas(before_formulas)
    takeoff_calcs_after = rectangular._takeoff_calculation_formulas(after_formulas)
    if takeoff_calcs_before != takeoff_calcs_after:
        raise RuntimeError("TAKE OFF K:AX calculation formulas changed")
    if any(key[0] == SPIRAL_SHEET for key in before_formulas if after_formulas.get(key) != before_formulas[key]):
        raise RuntimeError("A SPIRALE formula changed")

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        out_book = load_workbook(output, data_only=False, read_only=False)
        src_book = load_workbook(template, data_only=False, read_only=False)
    for sheet_name, cells in values.items():
        for coordinate, expected in cells.items():
            actual = out_book[sheet_name][coordinate].value
            matches = Decimal(str(actual)) == expected if isinstance(expected, Decimal) else actual == expected
            if not matches:
                raise RuntimeError(f"Output mismatch at {sheet_name}!{coordinate}: {actual!r} != {expected!r}")
        if set(map(str, src_book[sheet_name].merged_cells.ranges)) != set(map(str, out_book[sheet_name].merged_cells.ranges)):
            raise RuntimeError(f"Merged cells changed in {sheet_name}")
        if len(src_book[sheet_name]._images) != len(out_book[sheet_name]._images):
            raise RuntimeError(f"Image count changed in {sheet_name}")
    # Check styles and values in all non-written cells on changed sheets.
    for name in (TAKEOFF_SHEET, SPIRAL_SHEET):
        src_sheet, out_sheet = src_book[name], out_book[name]
        for row in src_sheet.iter_rows():
            for cell in row:
                if name == TAKEOFF_SHEET and cell.coordinate in formula_expansions:
                    continue
                if cell.coordinate not in written[name] and out_sheet[cell.coordinate].value != cell.value:
                    raise RuntimeError(f"Unexpected cell value change at {name}!{cell.coordinate}")
                if out_sheet[cell.coordinate].style_id != cell.style_id:
                    raise RuntimeError(f"Style changed at {name}!{cell.coordinate}")

    written_lines = []
    for name, cells in values.items():
        if cells:
            written_lines.append(name + ": " + ", ".join(f"{coord}={value}" for coord, value in cells.items()))
    rect_count = len(values[TAKEOFF_SHEET])
    spiral_count = len(values[SPIRAL_SHEET])
    replaced_template_formulas = sorted(
        f"{sheet}!{coordinate}"
        for sheet, coordinate in before_formulas
        if sheet in written and coordinate in written[sheet]
    )
    return [
        f"PASS: validated {len(records)} riser source rows; suppressed {suppressed} confirmed duplicate floor listings by matching RiserID and identical riser data.",
        "PASS: rectangular risers use the TAKE OFF section, installation RI, documented metric length, and Phase 2 accessory continuation behavior.",
        "PASS: spiral risers use exact existing diameter keys; feet are aggregated into the matching SPIRALE column B cell; non-matches are rejected without rounding.",
        "PASS: wrote only the following input cells: " + " | ".join(written_lines),
        f"PASS: preserved all TAKE OFF K:AX calculation formulas; {len(takeoff_calcs_before):,} formula entries match exactly.",
        f"PASS: preserved every SPIRALE formula; all {sum(1 for (s, _) in before_formulas if s == SPIRAL_SHEET):,} SPIRALE formulas match exactly.",
        f"PASS: preserved other formula cells; only intentionally populated input cells replace existing carry-down formulas; expanded {len(formula_expansions)} unaffected carry-down formulas.",
        "LIMITATION: original carry-down formula cells replaced by independent riser inputs: "
        + (", ".join(replaced_template_formulas) if replaced_template_formulas else "none")
        + ". Calculation formulas and all formulas outside explicitly written input cells are preserved.",
        f"PASS: output opens with openpyxl and ZIP integrity passes; only {takeoff_part} and {spiral_part} differ in the XLSX package.",
        "PASS: styles and merged cells are unchanged on both edited sheets; all other package parts, including drawings and validation extensions, are byte-identical.",
        "DEFERRED: E.H./E.B. remarks storage and English-unit rectangular riser length conversion remain unmapped and are not written or converted.",
        f"OUTPUT: {output}",
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=Path("source_files/Ventilation Jean Roy excel.xlsx"))
    parser.add_argument("--input", type=Path, default=Path("data/riser_phase5_test_cases.csv"))
    parser.add_argument("--output", type=Path, default=Path("output/Ventilation Jean Roy - riser phase5 output.xlsx"))
    parser.add_argument("--report", type=Path, default=Path("output/riser_phase5_validation_report.txt"))
    args = parser.parse_args()
    try:
        lines = generate(args.template, args.input, args.output)
    except InputValidationError as exc:
        print("VALIDATION FAILED")
        for error in exc.errors:
            print(f"ERROR: {error}")
        return 2
    except Exception as exc:
        print(f"GENERATION FAILED: {exc}", file=sys.stderr)
        return 1
    report = "RISER PHASE 5 VALIDATION REPORT\n" + "\n".join(lines) + "\n"
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(report, encoding="utf-8")
    print(report, end="")
    print(f"REPORT: {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
