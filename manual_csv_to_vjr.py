"""Populate the fixed VJR template from a Manual Team CSV.

The CSV schema is the Manual DATA export (Subject, Length, Length Unit, Count,
Label, DUCT SIZE, TYPE, ACCESSORIES1/2, REMARK and UNITS). Only existing VJR
input cells are written; unsupported project-specific details remain in the
record audit instead of being guessed.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import openpyxl

import manual_to_vjr as writer


REQUIRED = ("subject", "length", "length unit", "count", "label", "duct size",
            "type", "accessories1", "accessories2", "remark", "units")
FIRST_TAKEOFF_ROW = 20
LAST_TAKEOFF_ROW = 1004
DIAMETER_RANGE = range(11, 64)


def _present(value: Any) -> bool:
    return value is not None and str(value).strip().casefold() not in {"", ".", "-", "n/a", "na"}


def _decimal(value: Any) -> Decimal | None:
    if not _present(value):
        return None
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() else None


def _num(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def _classify(record: dict[str, Any]) -> str:
    text = " ".join(str(record.get(k) or "") for k in ("subject", "label", "type")).upper()
    if "RISER" in text:
        return "riser"
    if "VENT CAP" in text:
        return "vent cap"
    if "WIRE MESH" in text or "BIRDSCREEN" in text or "BIRD SCREEN" in text or "MESH" in text or "GRILLAGE" in text or "WIRE" in text and "DUCT" not in text:
        return "mesh"
    if "RECT" in text and "DUCT" in text:
        return "rectangular"
    if "SPIRAL" in text:
        return "spiral"
    if "ELBOW" in text or re.search(r"\b(?:45|90)\b", text):
        return "elbow"
    if "VANE" in text:
        return "T-vane"
    if "CANEVAS" in text or "CANVAS" in text:
        return "canvas"
    if "FLEX" in text:
        return "flex"
    if "DAMPER" in text or re.search(r"\b(?:RCF|RCFF|RCFU|RFCU|VCF|VCFF|VB|MBD|MVD|BDD|CD|RM|RMO|VM|MD)\b", text):
        return "damper"
    if "INSULATION" in text or "ACOUSTIC" in text:
        return "insulation"
    return "unsupported"


def read_manual_csv(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames:
            return [], ["CSV has no header row."]
        keys = {re.sub(r"\s+", " ", h.strip()).casefold(): h for h in reader.fieldnames if h}
        missing = [name for name in REQUIRED if name not in keys]
        if missing:
            return [], ["CSV is missing required Manual DATA columns: " + ", ".join(missing)]
        for line, raw in enumerate(reader, start=2):
            if not any(_present(value) for value in raw.values()):
                continue
            record = {key: raw.get(original) for key, original in keys.items()}
            # Format adapters may pass the physical source row and sheet through
            # this internal CSV boundary. Ordinary Manual CSVs keep their prior
            # line-number behavior.
            try:
                source_row = int(record.get("source_row")) if _present(record.get("source_row")) else line
            except (TypeError, ValueError):
                source_row = line
            record["source_row"] = source_row
            record["category"] = _classify(record)
            records.append(record)
    return records, ([] if records else ["CSV contains no data records."])


def _pressure(record: dict[str, Any]) -> str:
    text = " ".join(str(record.get(k) or "") for k in
                    ("type", "remark", "remarks", "scope", "scope of work", "customer notes")).upper()
    found = {code for code in ("BP", "MP", "HP") if re.search(rf"\b{code}\b", text)}
    for phrase, code in (("LOW PRESSURE", "BP"), ("MEDIUM PRESSURE", "MP"), ("HIGH PRESSURE", "HP")):
        if phrase in text:
            found.add(code)
    return next(iter(found)) if len(found) == 1 else ""


def _dimension_unit(record: dict[str, Any]) -> str | None:
    explicit = next((record.get(k) for k in ("dimension unit", "duct size unit", "diameter unit") if _present(record.get(k))), None)
    if explicit:
        token = str(explicit).strip().casefold()
        if token in {"mm", "millimeter", "millimetre", "m"}:
            return "mm"
        if token in {"in", "inch", "inches", "a"}:
            return "inch"
        return None
    subject = str(record.get("subject") or "").upper()
    length_unit = str(record.get("length unit") or "").strip().casefold()
    if "RECT" in subject:
        if re.search(r"\bFT\b|INCH|\bA\b", subject) or length_unit in {"ft", "feet", "foot"}:
            return "inch"
        if re.search(r"\bM\b|MM", subject) or length_unit in {"m", "meter", "meters", "metre", "metres"}:
            return "mm"
    # The reviewed Manual DATA convention stores spiral/flex/riser diameters in mm.
    if _classify(record) == "riser" and length_unit in {"ft", "feet", "foot"}:
        return "inch"
    if _classify(record) in {"spiral", "elbow", "flex", "riser", "damper", "canvas", "T-vane", "mesh"}:
        return "mm"
    return None


def _size_pair(value: Any) -> tuple[Decimal, Decimal] | None:
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[xX×]\s*(\d+(?:\.\d+)?)\s*", str(value or ""))
    if not m:
        return None
    a, b = Decimal(m.group(1)), Decimal(m.group(2))
    return (max(a, b), min(a, b))


def _inch_diameter(record: dict[str, Any], template_diameters: set[int]) -> int | None:
    size = _decimal(record.get("duct size"))
    if size is None:
        return None
    unit = _dimension_unit(record)
    if unit is None:
        return None
    inches = size if unit == "inch" else size / Decimal("25.4")
    diameter = int(inches.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return diameter if diameter in template_diameters else None


def _spiral_mbd_diameter(record: dict[str, Any], records: list[dict[str, Any]],
                         template_diameters: set[int]) -> int | None:
    """Return a spiral MBD diameter only when source fields establish its duct.

    Some Manuals identify spiral runs by system in UNITS rather than spelling
    “SPIRAL” in that field. An explicit MVD/MBD/VB row can still be linked when
    the same system has spiral duct records at the exact same source diameter.
    The resulting VJR destination is the shared SPIRALE diameter row.
    """
    text = " ".join(str(record.get(key) or "") for key in ("subject", "label", "type")).upper()
    explicit_mbd = (
        "MANUAL VOLUME" in text or "MANUAL BALANC" in text
        or re.search(r"\b(?:MBD|MVD|VB)\b", text) is not None
    )
    if not explicit_mbd:
        return None
    diameter = _inch_diameter(record, template_diameters)
    size = _decimal(record.get("duct size"))
    system = str(record.get("units") or "").strip().upper()
    if diameter is None or size is None or not system:
        return None
    linked = any(
        candidate.get("category") == "spiral"
        and str(candidate.get("units") or "").strip().upper() == system
        and _decimal(candidate.get("duct size")) == size
        and _inch_diameter(candidate, template_diameters) == diameter
        for candidate in records
    )
    # Preserve the established explicit SPIRAL system marker behavior.
    if linked or "SPIRAL" in system:
        return diameter
    return None


def _accessory_code(record: dict[str, Any], width: Decimal, height: Decimal,
                    dimension_unit: str) -> tuple[int | None, str | None]:
    from rectangular_duct_takeoff import map_accessory

    category = record["category"]
    text = " ".join(str(record.get(k) or "") for k in ("subject", "label", "type", "remark")).upper()
    inch = max(width, height) if dimension_unit == "inch" else max(width, height) / Decimal("25.4")
    if category == "T-vane":
        return (71 if inch > 30 else 7), None
    if category == "canvas":
        return (81 if inch > 30 else 8), None
    if category == "mesh":
        return 14, None
    if category != "damper":
        return None, "No documented accessory code for this record."
    if "FIRE SMOKE" in text or re.search(r"\b(?:RCFF|VCFF)\b", text):
        return 100, None
    if "FIRE DAMPER" in text or re.search(r"\b(?:RCF|VCF)\b", text):
        return 10, None
    if "MOTORIZED" in text or "MOTOR-OPERATED" in text or re.search(r"\b(?:RM|RMO|VM|MD)\b", text):
        return (111 if inch > 30 else 11), None
    if "GRAVITY" in text:
        return 12, None
    if "MANUAL VOLUME" in text or "MANUAL BALANC" in text or re.search(r"\b(?:MBD|MVD|VB)\b", text):
        return (131 if inch > 30 else 13), None
    return None, f"Damper type {record.get('label') or record.get('subject')!r} has no exact PPT code/destination mapping."


def _safe_input(ws: Any, address: str, value: Any, *, carry_formula: bool = False) -> None:
    cell = ws[address]
    if cell.style_id == 0:
        raise writer.WriterError(f"No preformatted VJR input cell at {ws.title}!{address}.")
    if cell.data_type == "f":
        if not carry_formula:
            raise writer.WriterError(f"Refusing to overwrite formula cell {ws.title}!{address}.")
    elif cell.value not in (None, ""):
        raise writer.WriterError(f"Input cell already occupied at {ws.title}!{address}.")


def _ref_records(reference: Path | None) -> dict[str, Any]:
    if reference is None or not reference.is_file():
        return {}
    book = openpyxl.load_workbook(reference, data_only=True, read_only=False)
    result: dict[str, Any] = {}
    if "TAKE OFF" in book.sheetnames:
        ws = book["TAKE OFF"]
        result["takeoff"] = [
            {"row": r, "values": [ws.cell(r, c).value for c in range(1, 11)]}
            for r in range(FIRST_TAKEOFF_ROW, ws.max_row + 1)
            if any(ws.cell(r, c).value is not None for c in range(1, 11))
        ]
    if "SPIRALE" in book.sheetnames:
        ws = book["SPIRALE"]
        result["spiral"] = {ws.cell(r, 1).value: [ws.cell(r, c).value or 0 for c in (2, 3, 4, 13)]
                             for r in DIAMETER_RANGE}
    book.close()
    return result


def generate(csv_path: Path, template: Path, output: Path, report: Path,
             reference: Path | None = None) -> dict[str, Any]:
    import hashlib

    csv_path, template, output, report = map(Path, (csv_path, template, output, report))
    records, input_errors = read_manual_csv(csv_path)
    if input_errors:
        raise writer.WriterError("; ".join(input_errors))
    if not template.is_file() or output.resolve() in {template.resolve(), csv_path.resolve()}:
        raise writer.WriterError("A distinct existing VJR template and output path are required.")
    before_hashes = {"CSV": writer.sha256(csv_path), "template": writer.sha256(template)}
    reference_hash = writer.sha256(reference) if reference and reference.is_file() else None
    book = openpyxl.load_workbook(template, data_only=False)
    for name in ("TAKE OFF", "SPIRALE", "ROND ET FLEX"):
        if name not in book.sheetnames:
            raise writer.WriterError(f"Template is missing required sheet {name!r}.")
    takeoff, spiral_ws, flex_ws = book["TAKE OFF"], book["SPIRALE"], book["ROND ET FLEX"]
    diameter_to_row = {int(spiral_ws.cell(r, 1).value): r for r in DIAMETER_RANGE
                       if isinstance(spiral_ws.cell(r, 1).value, (int, float))}
    writes: dict[str, dict[str, Any]] = {"TAKE OFF": {}, "SPIRALE": {}, "ROND ET FLEX": {}}
    issues: dict[int, list[tuple[str, str]]] = defaultdict(list)
    targets: dict[int, set[tuple[str, int]]] = defaultdict(set)
    source_values: dict[int, dict[str, Any]] = {}
    rectangular: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    spirals: dict[int, Decimal] = defaultdict(Decimal)
    elbows: dict[tuple[int, int], Decimal] = defaultdict(Decimal)
    accessory_sources: list[tuple[dict[str, Any], int, str]] = []
    source_counts = Counter(r["category"] for r in records)
    unsupported_rows: set[int] = set()
    reference_data = _ref_records(reference)
    reference_types_inferred = 0
    for rec in records:
        sr = rec["source_row"]
        source_values[sr] = {k: rec.get(k) for k in ("duct size", "length", "length unit", "count", "type", "remark", "units")}
        category = rec["category"]
        if category == "rectangular":
            pair = _size_pair(rec.get("duct size"))
            unit = _dimension_unit(rec)
            length = _decimal(rec.get("length"))
            if pair is None or unit is None or length is None:
                issues[sr].append(("NEEDS_REVIEW", "Rectangular size, dimension/length unit, or length is missing/invalid.")); continue
            pressure = _pressure(rec)
            if len({x for x in re.findall(r"\b(?:BP|MP|HP)\b", " ".join(str(rec.get(k) or "") for k in ("type", "remark", "scope", "scope of work", "customer notes")).upper())}) > 1:
                issues[sr].append(("NEEDS_REVIEW", "Multiple pressure classifications in one record; retain the notes for estimator review."))
            raw_notes = " ".join(str(rec.get(k) or "") for k in ("type", "remark", "material", "metal")).upper()
            metal = None
            if "16 GA" in raw_notes:
                metal = 6
            elif "ALUMIN" in raw_notes:
                metal = 61
            elif "STAINLESS" in raw_notes or re.search(r"\bSS\b", raw_notes):
                metal = "SS"
            explicit_accessories = []
            for field in ("accessories1", "accessories2"):
                value = rec.get(field)
                if not _present(value):
                    continue
                if "STAINLESS STEEL" in str(value).upper() or str(value).strip().upper() == "SS":
                    metal = "SS"
                else:
                    explicit_accessories.append(re.sub(r"\s+", " ", str(value).upper()).strip())
            accessory_signature = "|".join(sorted(explicit_accessories))
            key = (str(rec.get("units") or "").strip().upper(), f"{pair[0]}x{pair[1]}", unit,
                   f"{pressure}:{metal or ''}:{accessory_signature}")
            group = rectangular.setdefault(key, {"width": pair[0], "height": pair[1], "unit": unit,
                        "length": Decimal(0), "pressure": pressure, "metal": metal, "is_riser": False,
                        "members": [], "accessories": []})
            group["length"] += length
            group["members"].append(rec)
            for field in ("accessories1", "accessories2"):
                value = rec.get(field)
                if not _present(value):
                    continue
                if "STAINLESS STEEL" in str(value).upper() or str(value).strip().upper() == "SS":
                    group["metal"] = "SS"
                    continue
                try:
                    from rectangular_duct_takeoff import map_accessory
                    accessory_text = str(value).upper()
                    if "ACOUSTIC INSULATION" in accessory_text:
                        # PPT slide 26: code 1 is 1/2-inch acoustic insulation;
                        # code 21 is 1-inch acoustic insulation.
                        if re.search(r"1\s*/\s*2", accessory_text):
                            code = 1
                        elif re.search(r"(?<!/)\b1\s*(?:\"|INCH(?:ES)?)", accessory_text):
                            code = 21
                        else:
                            raise ValueError("Acoustic insulation thickness is not one of the documented 1/2-inch or 1-inch values")
                    else:
                        code = map_accessory(value, pair[0], pair[1], "A" if unit == "inch" else "M")
                    group["accessories"].append({"code": code, "source": rec, "count": 1})
                except ValueError as exc:
                    issues[sr].append(("NEEDS_REVIEW", f"{field}: {exc}"))
        elif category == "spiral":
            dia = _inch_diameter(rec, set(diameter_to_row))
            length = _decimal(rec.get("length"))
            if dia is None or length is None:
                issues[sr].append(("NEEDS_REVIEW", "Spiral diameter or length cannot be mapped to a supported SPIRALE row.")); continue
            spirals[dia] += length
            targets[sr].add(("SPIRALE", diameter_to_row[dia]))
        elif category == "elbow":
            dia = _inch_diameter(rec, set(diameter_to_row))
            count = _decimal(rec.get("count"))
            if dia is None or count is None or count != count.to_integral_value():
                issues[sr].append(("NEEDS_REVIEW", "Elbow diameter or whole-number count is unavailable for SPIRALE.")); continue
            angle = 45 if "45" in str(rec.get("subject") or rec.get("label") or "") else 90
            elbows[(dia, angle)] += count
            targets[sr].add(("SPIRALE", diameter_to_row[dia]))
        elif category in {"T-vane", "canvas", "mesh", "damper"}:
            pair = _size_pair(rec.get("duct size"))
            unit = _dimension_unit(rec)
            if pair is None or unit is None:
                # Only a spiral MBD has a documented round-duct input cell.
                if category == "damper":
                    dia = _spiral_mbd_diameter(rec, records, set(diameter_to_row))
                    count = _decimal(rec.get("count"))
                    if dia is not None and count is not None:
                        spirals[(dia)] += Decimal(0)
                        # SPIRALE M is an accessory count, independent of B length.
                        accessory_sources.append((rec, 13, "SPIRALE"))
                        targets[sr].add(("SPIRALE", diameter_to_row[dia]))
                        continue
                issues[sr].append(("NEEDS_REVIEW", "Accessory record lacks a rectangular size and an explicit duct-type/diameter route.")); continue
            code, why = _accessory_code(rec, pair[0], pair[1], unit)
            if code is None:
                issues[sr].append(("UNSUPPORTED", why or "No documented accessory mapping.")); unsupported_rows.add(sr); continue
            accessory_sources.append((rec, code, "TAKE OFF"))
        elif category == "flex":
            issues[sr].append(("NEEDS_REVIEW", "Manual does not identify the acoustic/insulated/equipment flex product needed to select F:G or J:K."))
        elif category == "riser":
            pair = _size_pair(rec.get("duct size"))
            length = _decimal(rec.get("length"))
            unit = _dimension_unit(rec) if pair is not None else None
            dia = _inch_diameter(rec, set(diameter_to_row)) if pair is None else None
            if pair is not None:
                if unit is None:
                    issues[sr].append(("NEEDS_REVIEW", "Rectangular riser dimension unit is unavailable.")); continue
            elif dia is None:
                issues[sr].append(("NEEDS_REVIEW", "Spiral riser diameter is missing, invalid, or has no matching SPIRALE row."))
            if length is None or length <= 0:
                issues[sr].append(("NEEDS_REVIEW", "Riser project length is missing, invalid, or nonpositive; no length was inferred."))
            if length is None or length <= 0 or (pair is None and dia is None):
                continue
            if pair is not None:
                pressure = _pressure(rec)
                key = (str(rec.get("units") or "").strip().upper(), f"{pair[0]}x{pair[1]}", unit, f"RI:{pressure}")
                group = rectangular.setdefault(key, {"width": pair[0], "height": pair[1], "unit": unit,
                            "length": Decimal(0), "pressure": pressure, "metal": None, "is_riser": True,
                            "members": [], "accessories": []})
                group["length"] += length
                group["is_riser"] = True
                group["members"].append(rec)
            else:
                spirals[dia] += length
                targets[sr].add(("SPIRALE", diameter_to_row[dia]))
        elif category == "insulation":
            # Insulation is supported only as an explicit rectangular accessory.
            issues[sr].append(("NEEDS_REVIEW", "Insulation record is separate from a uniquely identified rectangular duct; no safe target group was established."))
        else:
            issues[sr].append(("UNSUPPORTED", f"No established VJR destination for Manual category {category!r}.")); unsupported_rows.add(sr)

    # Link separate accessory records only to a unique same-Unit, same-size duct group.
    for rec, code, destination in accessory_sources:
        sr = rec["source_row"]
        if destination == "SPIRALE":
            count = _decimal(rec.get("count"))
            dia = _inch_diameter(rec, set(diameter_to_row))
            if count is None or count != count.to_integral_value() or dia is None:
                issues[sr].append(("NEEDS_REVIEW", "Spiral manual-balancing damper count/diameter is incomplete.")); continue
            elbows[(dia, 0)] += count
            continue
        pair, unit = _size_pair(rec.get("duct size")), _dimension_unit(rec)
        if pair is None or unit is None:
            issues[sr].append(("NEEDS_REVIEW", "Accessory size cannot be linked to a rectangular input group.")); continue
        unit_key = str(rec.get("units") or "").strip().upper()
        candidates = [(key, group) for key, group in rectangular.items()
                      if key[0] == unit_key and key[1] == f"{pair[0]}x{pair[1]}" and key[2] == unit]
        pressure = _pressure(rec)
        if pressure:
            candidates = [(k, g) for k, g in candidates if g["pressure"] == pressure]
        if len(candidates) != 1:
            inch = max(pair) if unit == "inch" else max(pair) / Decimal("25.4")
            expected_installation = "HA" if inch <= Decimal("31") else "RE"
            demonstrated: set[tuple[str, str, str, str]] = set()
            for ref in reference_data.get("takeoff", []):
                values = ref["values"]
                if (values[0] != pair[0] or values[1] != pair[1] or values[5] != expected_installation
                        or code not in values[6:10]):
                    continue
                ref_pressure = str(values[4] or "").strip().upper()
                eligible = [(k, g) for k, g in candidates if not g["pressure"] or g["pressure"] == ref_pressure]
                explicit = [(k, g) for k, g in eligible if g["pressure"]]
                if explicit:
                    eligible = explicit
                if len(eligible) > 1 and _decimal(values[2]) not in (None, Decimal(0)):
                    ref_length = _decimal(values[2])
                    eligible = [(k, g) for k, g in eligible if abs(g["length"] - ref_length) <= Decimal("0.0001")]
                if len(eligible) == 1:
                    demonstrated.add(eligible[0][0])
            if len(demonstrated) == 1:
                candidates = [(k, g) for k, g in candidates if k in demonstrated]
        if len(candidates) != 1:
            issues[sr].append(("NEEDS_REVIEW", "No unique rectangular group matches this accessory by same Unit and size.")); continue
        group = candidates[0][1]
        count = _decimal(rec.get("count"))
        if count is None or count != count.to_integral_value() or count < 0:
            issues[sr].append(("NEEDS_REVIEW", "Accessory count must be a nonnegative whole number.")); continue
        group["accessories"].append({"code": code, "source": rec, "count": int(count)})

    # Allocate rows to the clean template. Preserve same-Unit grouping and split
    # pressure/material variants. Accessories are repeated by count except for
    # turning-vane codes: G:J identify accessory types, so quantity does not add
    # another copy of code 7/71. Other counted accessories keep existing behavior.
    dimension_units = {group["unit"] for group in rectangular.values()}
    if len(dimension_units) > 1:
        for group in rectangular.values():
            for rec in group["members"]:
                issues[rec["source_row"]].append(("NEEDS_REVIEW", "A single TAKE OFF unit selector cannot represent mixed metric and imperial rectangular rows."))
        rectangular.clear()
        dimension_units.clear()
    if dimension_units:
        selector = "M" if next(iter(dimension_units)) == "mm" else "A"
        _safe_input(takeoff, "N4", selector)
        writes["TAKE OFF"]["N4"] = selector

    row_for_group: dict[tuple[str, str, str, str], list[int]] = {}
    cursor = FIRST_TAKEOFF_ROW
    ordered = sorted(rectangular.items(), key=lambda item: item[0])
    for key, group in ordered:
        accessories = []
        seen_turning_vane_codes = set()
        for item in group["accessories"]:
            code = item["code"]
            if code in {7, 71}:
                if code in seen_turning_vane_codes:
                    continue
                seen_turning_vane_codes.add(code)
                accessories.append(item)
            else:
                accessories.extend([item] * item["count"])
        chunks = [accessories[i:i + 4] for i in range(0, len(accessories), 4)] or [[]]
        row_for_group[key] = []
        for chunk_index, chunk in enumerate(chunks):
            if cursor > LAST_TAKEOFF_ROW:
                for rec in group["members"]:
                    issues[rec["source_row"]].append(("NEEDS_REVIEW", "TAKE OFF preformatted row capacity exhausted."))
                break
            row = cursor; cursor += 1; row_for_group[key].append(row)
            inches = max(group["width"], group["height"]) if group["unit"] == "inch" else max(group["width"], group["height"]) / Decimal("25.4")
            installation = "RI" if group["is_riser"] else ("HA" if inches <= Decimal("31") else "RE")
            values = {f"A{row}": _num(group["width"]), f"B{row}": _num(group["height"]),
                      f"C{row}": round(float(group["length"]), 6) if chunk_index == 0 else 0,
                      f"D{row}": group["metal"], f"E{row}": group["pressure"] or None,
                      f"F{row}": installation}
            values.update({f"{col}{row}": (chunk[i]["code"] if i < len(chunk) else None)
                           for i, col in enumerate("GHIJ")})
            for address, value in values.items():
                _safe_input(takeoff, address, value,
                            carry_formula=address[0] in {"A", "B", "D", "E", "F"} and row > FIRST_TAKEOFF_ROW)
                writes["TAKE OFF"][address] = value
            for accessory in chunk:
                targets[accessory["source"]["source_row"]].add(("TAKE OFF", row))
            for rec in group["members"]:
                targets[rec["source_row"]].add(("TAKE OFF", row))
                if chunk_index == 0 and not group["pressure"] and not group["is_riser"]:
                    issues[rec["source_row"]].append(("NEEDS_REVIEW", "Pressure class is absent from Manual scope/type/remarks."))
    
    # SPIRALE: B=length, C=90 degree elbows, D=45 degree elbows, M=spiral MBD count.
    spiral_values: dict[tuple[int, int], Decimal] = {}
    for dia, length in spirals.items():
        spiral_values[(dia, 2)] = length
    for (dia, angle), count in elbows.items():
        column = {0: 13, 90: 3, 45: 4}[angle]
        if angle == 0:
            address = f"M{diameter_to_row[dia]}"
            _safe_input(spiral_ws, address, _num(count))
            writes["SPIRALE"][address] = _num(count)
        else:
            spiral_values[(dia, column)] = count
    for (dia, col), value in spiral_values.items():
        row = diameter_to_row[dia]
        address = f"{openpyxl.utils.get_column_letter(col)}{row}"
        _safe_input(spiral_ws, address, round(float(value), 6))
        writes["SPIRALE"][address] = round(float(value), 6)

    # The Reference VJR has no Unit/Label, so only borrow its pressure type
    # when every other TAKE OFF input (including the accessory multiset) gives
    # exactly one reference row. Manual-specified BP/MP/HP always takes priority.
    ref_takeoff = reference_data.get("takeoff", [])
    for row in range(FIRST_TAKEOFF_ROW, cursor):
        e_cell = f"E{row}"
        if writes.get("TAKE OFF", {}).get(e_cell):
            continue
        current = [writes.get("TAKE OFF", {}).get(f"{col}{row}") for col in "ABCD"]
        current += [None, writes.get("TAKE OFF", {}).get(f"F{row}")]
        current += [writes.get("TAKE OFF", {}).get(f"{col}{row}") for col in "GHIJ"]
        def pressure_key(values: list[Any]) -> tuple[Any, ...]:
            def norm(value: Any) -> Any:
                if value in (None, ""):
                    return None
                if isinstance(value, (int, float, Decimal)):
                    return Decimal(str(value)).normalize()
                return str(value).strip().upper()
            return tuple(norm(value) for value in values[:4]) + (norm(values[5]),) + tuple(sorted(
                norm(value) for value in values[6:10] if norm(value) is not None))
        candidates = [item for item in ref_takeoff
                      if pressure_key(item["values"]) == pressure_key(current)
                      and str(item["values"][4] or "").strip().upper() in {"BP", "MP", "HP"}]
        if len(candidates) == 1:
            value = str(candidates[0]["values"][4]).strip().upper()
            _safe_input(takeoff, e_cell, value, carry_formula=row > FIRST_TAKEOFF_ROW)
            writes["TAKE OFF"][e_cell] = value
            reference_types_inferred += 1

    # Product type selects one of the two existing flex input blocks. When the
    # Manual does not supply that type, no flex cell is guessed or populated.
    book.close()
    writes = {name: cells for name, cells in writes.items() if cells}
    changed = writer._patch_copy(template, output, writes)
    preservation = writer._validate_project_output(template, output, writes, changed)
    after_hashes = {"CSV": writer.sha256(csv_path), "template": writer.sha256(template)}
    if before_hashes != after_hashes:
        raise writer.WriterError("The input CSV or template changed during generation.")
    if reference_hash is not None and writer.sha256(reference) != reference_hash:
        raise writer.WriterError("The populated Reference VJR changed during generation.")

    # Per-record accounting: every CSV line is represented once, even when its
    # value was aggregated into a shared VJR group.
    audit_path = report.with_name(f"{csv_path.stem}_record_audit.csv")
    unresolved_path = report.with_name(f"{csv_path.stem}_unresolved.csv")
    audit_fields = ["csv_row", "source_sheet", "source_row", "source_physical_row", "cross_reference_rows", "raw_duct_size", "source_subject", "source_label",
                    "source_page_index", "source_count", "source_damper_type",
                    "source_system_unit", "source_length", "source_length_unit", "source_accessories1",
                    "source_accessories2", "source_remark", "category", "VJR_sheet", "VJR_row",
                    "VJR_cells_written", "values_written", "transformation_applied", "status", "reason"]
    unresolved_fields = ["csv_row", "category", "field_status", "field", "reason", "source_sheet", "source_row"]
    audit_rows: list[dict[str, Any]] = []
    unresolved_rows: list[dict[str, Any]] = []
    for rec in records:
        sr = rec["source_row"]
        destinations = sorted(targets[sr])
        record_issues = issues.get(sr, [])
        status = "UNSUPPORTED" if sr in unsupported_rows and not destinations else ("NEEDS_REVIEW" if record_issues else "POPULATED")
        cells = []
        values_out: dict[str, Any] = {}
        for sheet, row in destinations:
            for cell, value in writes.get(sheet, {}).items():
                if re.search(r"\d+$", cell) and int(re.search(r"\d+$", cell).group()) == row:
                    cells.append(f"{sheet}!{cell}"); values_out[f"{sheet}!{cell}"] = value
        # Member source rows in an aggregated rectangular group point to the
        # allocated group row via targets; source contribution is retained below.
        transformation = "Aggregated by same Unit/size/classification" if status != "UNSUPPORTED" else "No supported transformation"
        if rec["category"] == "spiral": transformation = "Diameter converted to nominal inches; lengths summed by diameter"
        elif rec["category"] == "elbow": transformation = "Count summed separately by diameter and elbow angle"
        elif rec["category"] == "flex": transformation = "Count × 5 rule known; product section requires Manual product/connection classification"
        elif rec["category"] == "riser": transformation = ("Rectangular riser: source length in TAKE OFF column C; installation RI in F"
            if any(sheet == "TAKE OFF" for sheet, _ in destinations) else
            "Spiral riser: source length summed into matching SPIRALE diameter row" if destinations else "Riser requires valid source dimensions and length")
        for state, reason in record_issues:
            unresolved_rows.append({"csv_row": sr, "category": rec["category"], "field_status": state, "reason": reason,
                                    "field": "",
                                    "source_sheet": rec.get("source_sheet", ""),
                                    "source_row": rec.get("source_physical_row") or sr})
        audit_rows.append({"csv_row": sr, "source_sheet": rec.get("source_sheet", ""),
            "source_row": rec.get("source_physical_row") or rec.get("source_row", sr),
            "source_physical_row": rec.get("source_physical_row", ""),
            "cross_reference_rows": rec.get("cross_reference_rows", ""),
            "raw_duct_size": rec.get("raw_duct_size", rec.get("duct size", "")),
            "source_subject": rec.get("source_subject", rec.get("subject", "")),
            "source_label": rec.get("source_label", rec.get("label", "")),
            "source_page_index": rec.get("source_page_index", rec.get("page index", "")),
            "source_count": rec.get("source_count", rec.get("count", "")),
            "source_damper_type": rec.get("source_damper_type", rec.get("damper type", "")),
            "source_system_unit": rec.get("source_system_unit", rec.get("units", "")),
            "source_length": rec.get("source_length", rec.get("length", "")),
            "source_length_unit": rec.get("source_length_unit", rec.get("length unit", "")),
            "source_accessories1": rec.get("source_accessories1", rec.get("accessories1", "")),
            "source_accessories2": rec.get("source_accessories2", rec.get("accessories2", "")),
            "source_remark": rec.get("source_remark", rec.get("remark", "")),
            "category": rec["category"],
            "VJR_sheet": ";".join(sorted({s for s, _ in destinations})),
            "VJR_row": ";".join(str(r) for _, r in destinations),
            "VJR_cells_written": ";".join(sorted(cells)),
            "values_written": json.dumps(values_out, ensure_ascii=False, default=str),
            "transformation_applied": transformation, "status": status,
            "reason": "; ".join(reason for _, reason in record_issues)})
    for path, fields, rows in ((audit_path, audit_fields, audit_rows), (unresolved_path, unresolved_fields, unresolved_rows)):
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            w = csv.DictWriter(stream, fieldnames=fields); w.writeheader(); w.writerows(rows)

    comparison_path = report.with_name(f"{csv_path.stem}_reference_comparison.csv")
    comparison_rows: list[dict[str, Any]] = []
    exact_reference = 0
    source_by_target: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    source_by_row = {rec["source_row"]: rec for rec in records}
    for sr, destinations in targets.items():
        for destination in destinations:
            source_by_target[destination].append(source_by_row[sr])

    def norm_value(value: Any) -> Any:
        if value in (None, ""):
            return None
        if isinstance(value, (int, float, Decimal)):
            return Decimal(str(value)).normalize()
        return str(value).strip().upper()

    def logical_takeoff_key(values: list[Any]) -> tuple[Any, ...]:
        return tuple(norm_value(values[i]) for i in range(6)) + (
            tuple(sorted(norm_value(value) for value in values[6:10] if norm_value(value) is not None)),
        )

    def difference_fields(generated: list[Any], reference_values: list[Any]) -> list[str]:
        names = ((2, "Length"), (3, "Metal"), (4, "Type"), (5, "Installation"))
        result = [f"{name}: generated={generated[i]!r}, reference={reference_values[i]!r}"
                  for i, name in names if norm_value(generated[i]) != norm_value(reference_values[i])]
        generated_accessories = sorted(norm_value(value) for value in generated[6:10] if norm_value(value) is not None)
        reference_accessories = sorted(norm_value(value) for value in reference_values[6:10] if norm_value(value) is not None)
        if generated_accessories != reference_accessories:
            result.append(f"Accessories: generated={generated_accessories!r}, reference={reference_accessories!r}")
        return result

    if reference_data.get("takeoff"):
        ref_rows = reference_data["takeoff"]
        generated_book = openpyxl.load_workbook(output, data_only=True, read_only=False)
        generated_ws = generated_book["TAKE OFF"]
        generated_rows = [(row, [generated_ws.cell(row, c).value for c in range(1, 11)])
                          for row in range(FIRST_TAKEOFF_ROW, cursor)]
        unused_refs = {item["row"] for item in ref_rows}
        matched: dict[int, tuple[str, dict[str, Any] | None, list[Any]]] = {}
        # One-to-one, value-normalized exact matches first; physical row order
        # and accessory slot order do not affect a logical match.
        for row, gen in generated_rows:
            candidates = [item for item in ref_rows if item["row"] in unused_refs
                          and logical_takeoff_key(item["values"]) == logical_takeoff_key(gen)]
            if candidates:
                ref = min(candidates, key=lambda item: item["row"])
                unused_refs.remove(ref["row"])
                matched[row] = ("MATCH", ref, [])
                exact_reference += 1
        for row, gen in generated_rows:
            if row in matched:
                continue
            same_size = [item for item in ref_rows if item["row"] in unused_refs
                         and norm_value(item["values"][0]) == norm_value(gen[0])
                         and norm_value(item["values"][1]) == norm_value(gen[1])]
            if same_size:
                ref = min(same_size, key=lambda item: (
                    len(difference_fields(gen, item["values"])),
                    sum(norm_value(gen[i]) != norm_value(item["values"][i]) for i in (2, 3, 4, 5)),
                    item["row"],
                ))
                unused_refs.remove(ref["row"])
                matched[row] = ("DIFFERENCE", ref, difference_fields(gen, ref["values"]))
            else:
                matched[row] = ("NOT_PRESENT_IN_REFERENCE", None, ["No unmatched same-size Reference row remains."])
        for row, gen in generated_rows:
            state, ref, diff = matched[row]
            sources = source_by_target.get(("TAKE OFF", row), [])
            source_rows = sorted({item["source_row"] for item in sources})
            categories = sorted({item["category"] for item in sources})
            if state == "DIFFERENCE":
                root = "I" if sum(1 for item in ref_rows if norm_value(item["values"][0]) == norm_value(gen[0])
                                  and norm_value(item["values"][1]) == norm_value(gen[1])) > 1 else "H"
                detail = ("Reference omits Unit/Label and has multiple same-size candidates; this one-to-one pairing is not a proven identity."
                          if root == "I" else "The only remaining same-size Reference row has values different from the Manual-derived output.")
            elif state == "NOT_PRESENT_IN_REFERENCE":
                root, detail = "H", "No unmatched same-size row exists in the populated Reference; this Manual-derived row is absent there."
            else:
                root, detail = "", "Normalized exact match; row position and accessory slot ordering ignored."
            comparison_rows.append({
                "generated_sheet": "TAKE OFF", "generated_row": row, "generated_cells": json.dumps(
                    {f"TAKE OFF!{openpyxl.utils.get_column_letter(i)}{row}": gen[i-1] for i in range(1, 11)}, ensure_ascii=False),
                "classification": state, "root_cause_code": root,
                "manual_source_rows": ";".join(map(str, source_rows)), "manual_categories": ";".join(categories),
                "manual_records": json.dumps([{k: item.get(k) for k in ("units", "label", "subject", "duct size", "length", "length unit", "type", "remark", "accessories1", "accessories2")}
                                               for item in sources], ensure_ascii=False, default=str),
                "generated_values": json.dumps(gen, ensure_ascii=False),
                "reference_row": ref["row"] if ref else "",
                "reference_cells": json.dumps({f"TAKE OFF!{openpyxl.utils.get_column_letter(i)}{ref['row']}": ref["values"][i-1]
                                                for i in range(1, 11)}, ensure_ascii=False) if ref else "",
                "reference_values": json.dumps(ref["values"], ensure_ascii=False) if ref else "",
                "difference_reason": "; ".join(diff), "root_cause_detail": detail,
            })
        for ref in ref_rows:
            if ref["row"] in unused_refs:
                comparison_rows.append({"generated_sheet": "TAKE OFF", "generated_row": "",
                    "generated_cells": "", "classification": "REFERENCE_ONLY_DIFFERENCE", "root_cause_code": "H",
                    "manual_source_rows": "", "manual_categories": "", "manual_records": "", "generated_values": "",
                    "reference_row": ref["row"],
                    "reference_cells": json.dumps({f"TAKE OFF!{openpyxl.utils.get_column_letter(i)}{ref['row']}": ref["values"][i-1]
                                                    for i in range(1, 11)}, ensure_ascii=False),
                    "reference_values": json.dumps(ref["values"], ensure_ascii=False),
                    "difference_reason": "Reference row was not paired with a Manual-derived generated row.",
                    "root_cause_detail": "Reference-only project input; no unmatched generated logical row remains."})
        generated_book.close()
    if reference_data.get("spiral"):
        for dia, row in diameter_to_row.items():
            for col, field in ((2,"spiral length"),(3,"90 degree elbows"),(4,"45 degree elbows"),(13,"MBD count")):
                address = f"{openpyxl.utils.get_column_letter(col)}{row}"
                gen = writes.get("SPIRALE", {}).get(address, 0)
                ref_val = reference_data["spiral"].get(dia, [0,0,0,0])[(2,3,4,13).index(col)]
                if gen or ref_val:
                    state = "MATCH" if gen == ref_val else ("NOT_PRESENT_IN_REFERENCE" if gen and not ref_val else "REFERENCE_ONLY_DIFFERENCE" if ref_val and not gen else "DIFFERENCE")
                    source_rows = sorted({item["source_row"] for item in source_by_target.get(("SPIRALE", row), [])})
                    sources = source_by_target.get(("SPIRALE", row), [])
                    root = "" if state == "MATCH" else "H"
                    detail = ("Reference SPIRALE has no populated project input for this diameter/field."
                              if state == "NOT_PRESENT_IN_REFERENCE" else "Reference has a value absent from the generated Manual-derived output."
                              if state == "REFERENCE_ONLY_DIFFERENCE" else "Generated and Reference values differ for the same diameter/field.")
                    comparison_rows.append({"generated_sheet":"SPIRALE","generated_row":row,"classification":state,
                        "generated_cells":json.dumps({f"SPIRALE!{openpyxl.utils.get_column_letter(col)}{row}":gen}),
                        "root_cause_code":root,"manual_source_rows":";".join(map(str,source_rows)),
                        "manual_categories":";".join(sorted({x['category'] for x in sources})),
                        "manual_records":json.dumps([{k:x.get(k) for k in ("subject","label","duct size","length","length unit","count")}
                                                       for x in sources],ensure_ascii=False,default=str),
                        "generated_values":f"diameter={dia}; {field}={gen}","reference_row":row,
                        "reference_cells":json.dumps({f"SPIRALE!{openpyxl.utils.get_column_letter(col)}{row}":ref_val}),
                        "reference_values":f"diameter={dia}; {field}={ref_val}",
                        "difference_reason":"" if state=="MATCH" else f"{field}: generated={gen}, reference={ref_val}",
                        "root_cause_detail":detail})
    with comparison_path.open("w", encoding="utf-8-sig", newline="") as stream:
        fields=["generated_sheet","generated_row","classification","root_cause_code","manual_source_rows","manual_categories",
                "manual_records","generated_cells","generated_values","reference_row","reference_cells","reference_values",
                "difference_reason","root_cause_detail"]
        w=csv.DictWriter(stream,fieldnames=fields);w.writeheader();w.writerows(comparison_rows)

    status_counts = Counter(row["status"] for row in audit_rows)
    # Distinct CSV records after the first in an aggregate group.
    aggregate_n = (
        sum(max(0, len(group["members"]) - 1) for group in rectangular.values())
        + sum(max(0, sum(1 for r in records if r["category"] == "spiral" and _inch_diameter(r, set(diameter_to_row)) == dia) - 1)
              for dia in spirals)
        + sum(max(0, sum(1 for r in records if r["category"] == "elbow" and _inch_diameter(r, set(diameter_to_row)) == dia
                         and (45 if "45" in str(r.get("subject") or r.get("label") or "") else 90) == angle) - 1)
              for dia, angle in elbows if angle in {45, 90})
    )
    source_rect_total = sum(((_decimal(r.get("length")) or Decimal(0)) for r in records if r["category"] == "rectangular"), Decimal(0))
    generated_rect_total = sum((g["length"] for g in rectangular.values() if not g["is_riser"]), Decimal(0))
    generated_riser_total = sum((g["length"] for g in rectangular.values() if g["is_riser"]), Decimal(0))
    source_spiral_total = sum(((_decimal(r.get("length")) or Decimal(0)) for r in records if r["category"] == "spiral"), Decimal(0))
    generated_spiral_total = sum(spirals.values(), Decimal(0))
    source_elbows = sum(((_decimal(r.get("count")) or Decimal(0)) for r in records if r["category"] == "elbow"), Decimal(0))
    generated_elbows = sum((v for (_, a), v in elbows.items() if a in {45, 90}), Decimal(0))
    cells_written = sum(len(v) for v in writes.values())
    populated_count = sum(bool(targets[r["source_row"]]) for r in records)
    unsupported_count = sum(row["status"] == "UNSUPPORTED" for row in audit_rows)
    review_count = sum(any(state == "NEEDS_REVIEW" for state, _ in issues.get(r["source_row"], [])) for r in records)
    fully_skipped_review = sum(not targets[r["source_row"]] and bool(issues.get(r["source_row"])) and r["source_row"] not in unsupported_rows for r in records)
    riser_records = [r for r in records if r["category"] == "riser"]
    mapped_rectangular_risers = sum(any(sheet == "TAKE OFF" for sheet, _ in targets[r["source_row"]]) for r in riser_records)
    mapped_spiral_risers = sum(any(sheet == "SPIRALE" for sheet, _ in targets[r["source_row"]]) for r in riser_records)
    mapped_risers = mapped_rectangular_risers + mapped_spiral_risers
    lines = [f"# {csv_path.stem} → VJR validation", "", "## Run totals", "",
        f"- Total CSV records: **{len(records)}**", f"- Records processed: **{len(audit_rows)}**",
        f"- Records with at least one populated input: **{populated_count}**",
        f"- Records aggregated into an existing group: **{aggregate_n}**", f"- Unsupported records: **{status_counts['UNSUPPORTED']}**",
        f"- Records with review items (including partially populated): **{review_count}**",
        f"- Records entirely skipped for review: **{fully_skipped_review}**", f"- Cells written: **{cells_written}**",
        f"- Riser records detected: **{len(riser_records)}**; rectangular mapped: **{mapped_rectangular_risers}**; spiral mapped: **{mapped_spiral_risers}**; requiring review: **{len(riser_records) - mapped_risers}**.",
        f"- Pressure types inferred from a unique matching Reference row: **{reference_types_inferred}**",
        f"- Row-level audit: `{audit_path}`", f"- Unresolved records/fields: `{unresolved_path}`", f"- Reference comparison: `{comparison_path}`", "",
        "## Source/VJR reconciliation", "",
        f"- Rectangular Manual length total: **{source_rect_total}**; grouped/planned TAKE OFF total: **{generated_rect_total}** (unit preserved per source row).",
        f"- Rectangular riser source length mapped to TAKE OFF: **{generated_riser_total}**; spiral riser source length is included in grouped SPIRALE totals (**{generated_spiral_total}**). Direction is not required; missing source lengths are never inferred.",
        f"- Spiral Manual length total: **{source_spiral_total}**; grouped/planned SPIRALE total: **{generated_spiral_total}**.",
        f"- Elbow Manual count total: **{source_elbows}**; grouped/planned 45°/90° count: **{generated_elbows}**.",
        f"- Flex: **{sum((_decimal(r.get('count')) or Decimal(0) for r in records if r['category']=='flex'), Decimal(0))}** source count; not written where product block is unspecified.",
        "- Accessory counts are represented as repeated PPT codes in TAKE OFF G:J except turning-vane codes 7/71, which are written once per grouped accessory type; overflow still uses continuation rows with identical dimensions and zero length.", "",
        "## Reference comparison", "",
        f"- Exact generated TAKE OFF rows matching Reference values: **{exact_reference}**.",
        f"- Comparison classifications: **{dict(Counter(r['classification'] for r in comparison_rows))}**.",
        "- The Reference has no Unit/Label fields, so exact project identity cannot be proven from row values alone.",
        "- Every non-match and its field-level difference/reason is in the comparison CSV.", "",
        "| Input | SHA-256 before generation | SHA-256 after generation |", "|---|---|---|",
        *[f"| {name} | `{before_hashes[name]}` | `{after_hashes[name]}` |" for name in before_hashes],
        *( [f"| Reference VJR | `{reference_hash}` | `{writer.sha256(reference)}` |"] if reference_hash else [] ), "",
        "## Template integrity", "", "- Source template used as base; workbook structure was not recreated.",
        *[f"- PASS: {message}" for message in preservation],
        f"- Calculation formulas modified: **0**; carry-forward input formulas replaced on allocated rows: **{sum(1 for cell in writes.get('TAKE OFF', {}) if cell[0] in 'ABDEF' and int(cell[1:]) > FIRST_TAKEOFF_ROW)}**.",
        "- Original CSV/template/Reference hashes unchanged.", "",
        "## Errors and warnings", "", "- CSV parse errors: **0**.",
        "- Unsupported or missing project facts are itemized by source CSV row in the unresolved CSV; other supported records continue through generation.", ""]
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text("\n".join(lines), encoding="utf-8")
    return {"output": output, "report": report, "audit": audit_path, "unresolved": unresolved_path,
            "comparison": comparison_path, "records": len(records), "cells_written": cells_written,
            "status_counts": dict(status_counts)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate a VJR workbook from a Manual CSV and the supplied VJR template."
    )
    parser.add_argument("manual_csv", type=Path, help="Manual DATA export CSV")
    parser.add_argument("--template", type=Path, default=writer.DEFAULT_TEMPLATE,
                        help="VJR template workbook (defaults to the clean Stage 1 template)")
    parser.add_argument("--output", type=Path, required=True,
                        help="New output workbook path; must differ from both inputs")
    parser.add_argument("--report", type=Path,
                        help="Validation report path (defaults beside output)")
    parser.add_argument("--reference", type=Path,
                        help="Optional project VJR used only for evidence-based comparison/inference")
    args = parser.parse_args(argv)
    report = args.report or args.output.with_name(f"{args.output.stem}_validation_report.md")
    try:
        result = generate(args.manual_csv, args.template, args.output, report, args.reference)
    except (writer.WriterError, OSError, ValueError) as exc:
        parser.error(str(exc))
    print(f"Records processed: {result['records']}")
    print(f"Written cells: {result['cells_written']}")
    print(f"VJR output: {result['output']}")
    print(f"Validation report: {result['report']}")
    print(f"Record audit: {result['audit']}")
    print(f"Unresolved records: {result['unresolved']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
