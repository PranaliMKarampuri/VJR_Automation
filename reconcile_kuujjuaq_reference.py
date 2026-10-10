"""Build read-only, record-level Kuujjuaq Manual-to-reference-VJR reports."""
from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import vjr_ppt_rules as ppt

ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
MANUAL = ROOT / "source_files" / "TAKEOFF_CS KATIVIK - CENTRE D'ÉDUCATION AUX ADULTES.xlsx"
VJR = ROOT / "source_files" / "TAKEOFF_CS KATIVIK - CENTRE D'ÉDUCATION AUX ADULTES_VJR.xlsx"
DRYRUN = REPORTS / "kuujjuaq_manual_to_vjr_dry_run.csv"
RECT = REPORTS / "kuujjuaq_rectangular_reconciliation.csv"
EXISTING = REPORTS / "kuujjuaq_vjr_existing_rows_analysis.csv"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def json_cell(value: Any) -> Any:
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    return str(value)


def json_compact(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


def load_manual() -> tuple[list[str], dict[int, dict[str, Any]]]:
    book = load_workbook(MANUAL, data_only=True, read_only=True)
    sheet = book["DATA"]
    headers = [text(x) for x in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    records = {}
    for rownum, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if any(x is not None for x in values):
            records[rownum] = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
    book.close()
    return headers, records


def load_vjr() -> tuple[dict[str, Any], dict[int, dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    formulas = load_workbook(VJR, data_only=False, read_only=False)
    values = load_workbook(VJR, data_only=True, read_only=False)
    ws_f, ws_v = formulas["TAKE OFF"], values["TAKE OFF"]
    metadata = {"unit_selector": text(ws_v["N4"].value).upper(), "sheets": formulas.sheetnames}
    rows: dict[int, dict[str, Any]] = {}
    for rownum in range(20, 1005):
        cells = {f"{col}{rownum}": ws_f[f"{col}{rownum}"].value for col in "ABCDEFGHIJ"}
        val = {f"{col}{rownum}": ws_v[f"{col}{rownum}"].value for col in "ABCDEFGHIJ"}
        if any(ws_f[f"{col}{rownum}"].value not in (None, "")
               and not (isinstance(ws_f[f"{col}{rownum}"].value, str)
                        and ws_f[f"{col}{rownum}"].value.startswith("=")) for col in "ABC"):
            rows[rownum] = {
                "cells_raw": cells,
                "cells_values": val,
                "values": {col: ws_v[f"{col}{rownum}"].value for col in "ABCDEFGHIJ"},
                "formulas": {col: ws_f[f"{col}{rownum}"].value for col in "ABCDEFGHIJ"
                             if ws_f[f"{col}{rownum}"].data_type == "f"},
                "styled_input_cells": [f"{col}{rownum}" for col in "ABCDEFGHIJ"
                                       if ws_f[f"{col}{rownum}"].style_id > 0],
            }
    full = {}
    for name in formulas.sheetnames:
        sf, sv = formulas[name], values[name]
        full[name] = {"formula_sheet": sf, "value_sheet": sv}
    formulas.close()
    values.close()
    return metadata, rows, full


def vjr_literal_cells(row: dict[str, Any]) -> dict[str, Any]:
    result = {}
    for address, raw in row["cells_raw"].items():
        if raw not in (None, "") and not (isinstance(raw, str) and raw.startswith("=")):
            result[address] = json_cell(raw)
    return result


def target_writer_assessment(target_rows: list[int], vjr_rows: dict[int, dict[str, Any]]) -> str:
    reasons = []
    for rownum in target_rows:
        item = vjr_rows.get(rownum)
        if not item:
            reasons.append(f"TAKE OFF row {rownum} not inventoried")
            continue
        raw = item["cells_raw"]
        blocked = [a for a, x in raw.items() if x not in (None, "")]
        if blocked:
            reasons.append(f"TAKE OFF row {rownum} is occupied/formula-bearing in A:J ({','.join(blocked)})")
        else:
            reasons.append(f"TAKE OFF row {rownum} does not correspond to an observed record")
    return "; ".join(reasons) if reasons else "No destination identified."


def make_reports() -> None:
    _, manual_rows = load_manual()
    meta, vjr_rows, sheets = load_vjr()
    dry = {int(x["manual_source_row"]): x for x in read_csv(DRYRUN)}
    rect = {int(x["source row"]): x for x in read_csv(RECT)}
    existing_rows = read_csv(EXISTING)
    existing_keys = {(x["sheet"], int(x["row"])): x for x in existing_rows}

    # The existing reconciliation explicitly identifies these DATA-to-TAKEOFF
    # group rollups. Membership remains a reference fact pending full approval.
    group_members: dict[int, list[int]] = defaultdict(list)
    group_rollup_meta: dict[int, dict[str, str]] = {}
    for source_row, row in rect.items():
        if row["classification"] == "EXACT_GROUP_ROLLUP_MATCH":
            target = int(row["VJR candidate row"])
            group_members[target].append(source_row)
            group_rollup_meta[target] = row

    reference_vjr_rows = set(vjr_rows)
    exact_vjr_rows = set(group_members)
    vjr_only_rows = sorted(reference_vjr_rows - exact_vjr_rows)
    source_status = {row: dry[row]["mapping_status"] for row in manual_rows}
    safe_count = 0
    status_counts = Counter()
    rows_out: list[dict[str, Any]] = []
    manifest_rows: list[dict[str, Any]] = []

    for source_row, rec in manual_rows.items():
        dr = dry.get(source_row, {})
        rr = rect.get(source_row, {})
        group_target = None
        for target, members in group_members.items():
            if source_row in members:
                group_target = target
                break
        if group_target is not None:
            mapping_class = "GROUPED_REFERENCE_MATCH"
            candidates = [group_target]
            rationale = (f"DATA row is included in the existing reconciliation's exact group roll-up to TAKE OFF row "
                         f"{group_target}; this is a grouped size/length reference, not a full-record authorization.")
            method = ("Aggregate the listed Manual DATA member lengths once into the cited TAKEOFF-DUCT/reference "
                      "TAKE OFF group. The exact group size/length roll-up is evidenced; no second write per member "
                      "and no full type/material/accessory transformation is authorized.")
        else:
            candidates = []
            cand_sheet = text(dr.get("vjr_candidate_sheet"))
            cand_rows = text(dr.get("vjr_candidate_row"))
            if cand_sheet and cand_rows:
                candidates = [int(x) for x in re.findall(r"\d+", cand_rows)]
            mapping_class = "CANDIDATE_REFERENCE" if candidates else "UNMATCHED_MANUAL"
            rationale = text(rr.get("reason") or dr.get("reason")) or "No supported reference counterpart identified."
            method = "No transformation is established. " + (text(dr.get("reason")) or text(rr.get("reason")) or "No exact reference mapping.")

        # SAFE is deliberately strict: this writer can only write explicitly
        # confirmed empty styled input cells and cannot overwrite this reference.
        status = text(dr.get("mapping_status")) or "NEEDS_CONFIRMATION"
        if status == "SAFE_TO_AUTOMATE":
            status = "NEEDS_CONFIRMATION"
        status_counts[status] += 1
        target = candidates[0] if len(candidates) == 1 else None
        vjr = vjr_rows.get(target) if target is not None else None
        destination = f"TAKE OFF!{target}" if target is not None else text(dr.get("vjr_destination"))
        if target is not None and vjr:
            destination_cells = json_compact(vjr_literal_cells(vjr))
        elif candidates:
            destination_cells = json_compact({str(t): vjr_literal_cells(vjr_rows[t]) for t in candidates if t in vjr_rows})
        else:
            destination_cells = "{}"

        accessories = {k: rec.get(k) for k in ("ACCESSORIES1", "ACCESSORIES2", "REMARK", "DAMPER TYPE")}
        target_vals = vjr["values"] if vjr else {}
        target_formulas = vjr["formulas"] if vjr else {}
        target_accessories = {f"{col}{target}": target_vals.get(col)
                              for col in "GHIJ" if target is not None and target_vals.get(col) not in (None, "")}
        ppt_checks = []
        if target is not None and vjr:
            f = text(target_vals.get("F")).upper()
            pressure = text(target_vals.get("E")).upper()
            if f in {"HA", "RE", "RI"}:
                ppt_checks.append(f"Reference install={f}; see VJR conflict analysis")
            if pressure and pressure not in {"BP", "MP", "HP"}:
                ppt_checks.append(f"Reference TYPE {pressure!r} is not a PPT BP/MP/HP pressure label")
            for col in "GHIJ":
                val = target_vals.get(col)
                if isinstance(val, (int, float)) and int(val) not in ppt.ACCESSORY_RULES:
                    ppt_checks.append(f"Reference code {val} in {col}{target} is not mapped in PPT slide 26")

        out = {
            "mapping_id": f"M{source_row:03d}", "record_kind": "MANUAL_RECORD",
            "manual_sheet": "DATA", "manual_row": source_row,
            "manual_source_values_json": json_compact({h: json_cell(v) for h, v in rec.items()}),
            "manual_subject": text(rec.get("Subject")), "manual_label": text(rec.get("Label")),
            "manual_system_unit": text(rec.get("UNITS")), "manual_size": text(rec.get("DUCT SIZE")),
            "manual_length": text(rec.get("Length")), "manual_length_unit": text(rec.get("Length Unit")),
            "manual_count": text(rec.get("Count")), "manual_type": text(rec.get("TYPE")),
            "manual_remark_accessories_json": json_compact(accessories),
            "record_family": text(dr.get("rule_category")), "reference_mapping_class": mapping_class,
            "reference_vjr_sheet": "TAKE OFF" if target is not None else text(dr.get("vjr_candidate_sheet")),
            "reference_vjr_destination": f"TAKE OFF!row {target}" if target is not None else text(dr.get("vjr_destination")),
            "reference_vjr_row": str(target) if target is not None else ";".join(str(x) for x in candidates),
            "reference_vjr_cells_json": destination_cells,
            "reference_dimensions": f"{target_vals.get('A','')}x{target_vals.get('B','')}" if target else "",
            "reference_length": text(target_vals.get("C")) if target else "",
            "reference_metal": text(target_vals.get("D")) if target else "",
            "reference_type": text(target_vals.get("E")) if target else "",
            "reference_installation": text(target_vals.get("F")) if target else "",
            "reference_accessory_cells_json": json_compact(target_accessories),
            "reference_formulas_json": json_compact(target_formulas),
            "transformation": method,
            "PPT_assessment": "; ".join(ppt_checks) or ("No row-level PPT conflict established; incomplete source facts remain confirmation-required."),
            "classification": status,
            "writer_input_assessment": target_writer_assessment([target], vjr_rows) if target else "No unique target cells; candidates cannot be written.",
            "evidence": rationale,
            "open_questions": text(dr.get("estimator_confirmation_required")),
        }
        rows_out.append(out)
        manifest_rows.append({
            "manifest_scope": "REFERENCE_CANDIDATE_ONLY", "manual_source_sheet": "DATA",
            "manual_source_row": source_row, "reference_mapping_class": mapping_class,
            "reference_vjr_sheet": out["reference_vjr_sheet"], "reference_vjr_row": out["reference_vjr_row"],
            "observed_vjr_cells_json": destination_cells, "transformation": method,
            "PPT_rule_status": text(dr.get("rule_status")), "classification": status,
            "estimator_decision": "NEEDS_CONFIRMATION" if status != "UNSUPPORTED" else "UNSUPPORTED",
            "writer_write_values_json": "{}", "approval_note": "Reference evidence only; never accepted by production writer; no cells authorized." if status != "UNSUPPORTED" else "Unsupported/manual; no cells authorized.",
        })

    # Explicit VJR-only records are added to the main reconciliation so all 45
    # literal TAKE OFF entries have a record-level explanation.
    vjr_only_out = []
    for target in vjr_only_rows:
        item = vjr_rows[target]
        vals = item["values"]
        rowcsv = existing_keys.get(("TAKE OFF", target), {})
        record = {
            "mapping_id": f"VJR-TAKEOFF-{target}", "record_kind": "VJR_ONLY_RECORD",
            "manual_sheet": "", "manual_row": "", "manual_source_values_json": "{}",
            "manual_subject": "", "manual_label": "", "manual_system_unit": "",
            "manual_size": "", "manual_length": "", "manual_length_unit": "",
            "manual_count": "", "manual_type": "", "manual_remark_accessories_json": "{}",
            "record_family": "rectangular_takeoff", "reference_mapping_class": "VJR_ONLY_UNMATCHED",
            "reference_vjr_sheet": "TAKE OFF", "reference_vjr_row": target,
            "reference_vjr_destination": f"TAKE OFF!row {target}",
            "reference_vjr_cells_json": json_compact(vjr_literal_cells(item)),
            "reference_dimensions": f"{vals.get('A','')}x{vals.get('B','')}",
            "reference_length": text(vals.get("C")), "reference_metal": text(vals.get("D")),
            "reference_type": text(vals.get("E")), "reference_installation": text(vals.get("F")),
            "reference_accessory_cells_json": json_compact({f"{c}{target}": vals.get(c) for c in "GHIJ" if vals.get(c) not in (None, "")}),
            "reference_formulas_json": json_compact(item["formulas"]),
            "transformation": "No unique Manual DATA/group counterpart; retained as a VJR-only reference line.",
            "PPT_assessment": "Inspected separately in rule-conflict report; no Manual source is assigned.",
            "classification": "NEEDS_CONFIRMATION", "writer_input_assessment": "Occupied existing reference row; not a writable destination.",
            "evidence": text(rowcsv.get("reason")) or "No exact Manual grouped counterpart identified in existing reconciliation.",
            "open_questions": "Identify source scope/system/group or confirm this VJR-only line is intentional.",
        }
        rows_out.append(record)
        vjr_only_out.append(record)

    fields = list(rows_out[0])
    with (REPORTS / "kuujjuaq_reference_mapping.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader(); writer.writerows(rows_out)
    manifest_fields = list(manifest_rows[0])
    with (REPORTS / "kuujjuaq_reference_mapping_manifest.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=manifest_fields)
        writer.writeheader(); writer.writerows(manifest_rows)

    # Every unmatched Manual record is preserved in a concise unresolved ledger.
    unresolved = [r for r in rows_out if r["record_kind"] == "MANUAL_RECORD" and r["reference_mapping_class"] != "GROUPED_REFERENCE_MATCH"]
    ufields = ["mapping_id", "manual_sheet", "manual_row", "record_family", "manual_subject", "manual_system_unit",
               "manual_size", "manual_length", "manual_length_unit", "manual_count", "manual_type",
               "manual_remark_accessories_json", "candidate_vjr_sheet", "candidate_vjr_rows", "classification",
               "reason", "missing_decisions"]
    with (REPORTS / "kuujjuaq_reference_unresolved.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=ufields)
        writer.writeheader()
        for r in unresolved:
            writer.writerow({
                "mapping_id": r["mapping_id"], "manual_sheet": r["manual_sheet"], "manual_row": r["manual_row"],
                "record_family": r["record_family"], "manual_subject": r["manual_subject"],
                "manual_system_unit": r["manual_system_unit"], "manual_size": r["manual_size"],
                "manual_length": r["manual_length"], "manual_length_unit": r["manual_length_unit"],
                "manual_count": r["manual_count"], "manual_type": r["manual_type"],
                "manual_remark_accessories_json": r["manual_remark_accessories_json"],
                "candidate_vjr_sheet": r["reference_vjr_sheet"], "candidate_vjr_rows": r["reference_vjr_row"],
                "classification": r["classification"], "reason": r["evidence"], "missing_decisions": r["open_questions"],
            })

    # Compare actual VJR TAKE OFF literal installation and accessory codes to PPT.
    conflicts = []
    exact_boundary = []
    unsupported_codes = []
    selector = meta["unit_selector"]
    for target, item in vjr_rows.items():
        vals = item["values"]
        try:
            w, h = Decimal(str(vals["A"])), Decimal(str(vals["B"]))
        except Exception:
            continue
        install = text(vals["F"]).upper()
        if selector in {"M", "MM"}:
            wi, hi = w / Decimal("25.4"), h / Decimal("25.4")
        elif selector in {"A", "IN", "INCH", "INCHES"}:
            wi, hi = w, h
        else:
            wi = hi = None
        if wi is not None and install in {"HA", "RE"}:
            if wi == 36 or hi == 36:
                exact_boundary.append((target, w, h, install))
                conflicts.append({"kind": "PPT_UNDEFINED_36_IN_BOUNDARY", "target": target,
                                  "cells": f"A{target}:B{target},F{target}", "reference": install,
                                  "rule": "PPT slide 30 defines <36 HA and >36 RE, but exactly 36 is not specified.",
                                  "detail": f"Dimensions {w}x{h} under N4={selector} put a parameter exactly at 36 inches."})
            elif max(wi, hi) < 36 and install != "HA":
                conflicts.append({"kind": "PPT_INSTALLATION_CONFLICT", "target": target,
                                  "cells": f"A{target}:B{target},F{target}", "reference": install,
                                  "rule": "PPT slide 30: all dimensions below 36 inches => HA.",
                                  "detail": f"Reference shows {w}x{h} with N4={selector}, both below 36 inches."})
            elif max(wi, hi) > 36 and install != "RE":
                conflicts.append({"kind": "PPT_INSTALLATION_CONFLICT", "target": target,
                                  "cells": f"A{target}:B{target},F{target}", "reference": install,
                                  "rule": "PPT slide 30: if width or height exceeds 36 inches => RE.",
                                  "detail": f"Reference shows {w}x{h} with N4={selector}, at least one exceeds 36 inches."})
        for col in "GHIJ":
            code_value = vals.get(col)
            if code_value in (None, ""):
                continue
            try:
                code = int(float(code_value))
            except (ValueError, TypeError):
                unsupported_codes.append((target, f"{col}{target}", code_value, "not numeric")); continue
            if code not in ppt.ACCESSORY_RULES:
                unsupported_codes.append((target, f"{col}{target}", code, "no mapping in PPT slide 26"))
    for row, cell, code, why in unsupported_codes:
        conflicts.append({"kind": "PPT_UNMAPPED_REFERENCE_ACCESSORY", "target": row, "cells": cell,
                          "reference": str(code), "rule": "PPT slide 26 enumerates supported accessory codes.",
                          "detail": f"Reference code {code} is not mapped by the PPT accessory table ({why})."})
    for target, members in group_members.items():
        expected_codes = []
        source_notes = []
        for source in members:
            rec = manual_rows[source]
            for field in ("ACCESSORIES1", "ACCESSORIES2", "REMARK"):
                note = text(rec.get(field)).upper()
                if "ACOUSTIC" not in note:
                    continue
                source_notes.append(f"DATA!{source}:{field}={text(rec.get(field))}")
                if "1/2" in note or "½" in note:
                    expected_codes.append(1)
                elif '1"' in note or "1 INCH" in note:
                    expected_codes.append(21)
        # These explicit thickness rules are one-code-per-record/group flag;
        # repeated source notes do not authorize repeated VJR slots.
        expected_codes = sorted(set(expected_codes))
        actual_codes = [int(float(vjr_rows[target]["values"][col])) for col in "GHIJ"
                        if vjr_rows[target]["values"].get(col) not in (None, "")
                        and str(vjr_rows[target]["values"][col]).replace(".0", "").isdigit()]
        for expected in expected_codes:
            if expected not in actual_codes:
                conflicts.append({"kind": "PPT_ACCESSORY_REFERENCE_MISMATCH", "target": target,
                                  "cells": f"G{target}:J{target}", "reference": ",".join(map(str, actual_codes)) or "[none]",
                                  "rule": f"PPT slide 26: source acoustic insulation thickness maps to code {expected}.",
                                  "detail": f"Grouped Manual source note(s) {', '.join(source_notes)} require code {expected}; it is absent from the observed VJR accessory slots."})

    conflict_lines = [
        "# Kuujjuaq reference VJR versus PPT rule conflicts", "",
        "## Rule basis", "",
        "The client VJR is treated as evidence of what the team entered, not as authority for business rules. The controlling rules below were checked against `source_files/VENTILATION_DUCTING.pptx` (slides 26, 28, 30) and the existing PPT transcription `vjr_ppt_rules.py`.", "",
        f"Reference workbook `TAKE OFF!N4` selector: **{selector or '[blank]'}**. The PPT says M selects millimetres/metres and A selects inches/feet (slide 28).", "",
        "## Detected conflicts and undefined cases", "",
        f"- Installation conflicts: **{sum(x['kind']=='PPT_INSTALLATION_CONFLICT' for x in conflicts)}** rows.",
        f"- Exactly-36-inch undefined reference cases: **{sum(x['kind']=='PPT_UNDEFINED_36_IN_BOUNDARY' for x in conflicts)}** rows.",
        f"- Reference accessory cells with no PPT slide 26 mapping: **{sum(x['kind']=='PPT_UNMAPPED_REFERENCE_ACCESSORY' for x in conflicts)}** cells.", "",
        f"- Explicit Manual acoustic notes whose mapped code is absent in the linked reference row: **{sum(x['kind']=='PPT_ACCESSORY_REFERENCE_MISMATCH' for x in conflicts)}** grouped cases.", "",
        "| Kind | VJR row/cell | Existing value | PPT rule | Finding |", "|---|---|---|---|---|",
    ]
    for c in conflicts:
        conflict_lines.append(f"| {c['kind']} | TAKE OFF!{c['cells']} | {c['reference']} | {c['rule']} | {c['detail']} |")
    conflict_lines += ["", "## Other limits observed in the reference", "",
        "- A literal VJR TYPE BP/MP/HP value can be compared with slide 30, but it does not establish the Manual customer-note mapping when Manual TYPE/REMARK is blank or placeholder.",
        "- Blank/formula-driven VJR values are not interpreted as blank input or as a numeric value. Formula cells remain protected; see `kuujjuaq_vjr_blank_formula_analysis.csv`.",
        "- Existing reference rows are occupied. The current writer blocks overwriting literal values/formulas even when those coordinates belong to the normal TAKE OFF input area.",
        "- Spiral/flex/elbow/riser/damper/canvas source rows remain unresolved wherever diameter keys, accessory association, drop lengths, confirmation, or exact target cells are missing. Reference proximity or a similar row is not accepted as a mapping.",
        "- The 32 exact grouped rectangular reference pairs are exact on the previously documented group roll-up dimensions/length and evidence keys; that alone does not establish every source row's complete type/material/accessory transformation or a writable target.", ""]
    (REPORTS / "kuujjuaq_reference_rule_conflicts.md").write_text("\n".join(conflict_lines), encoding="utf-8")

    count_exact_groups = len(group_members)
    grouped_manual = sum(len(v) for v in group_members.values())
    manual_unmatched = len(manual_rows) - grouped_manual
    safe_count = 0
    needs = status_counts["NEEDS_CONFIRMATION"]
    unsupported = status_counts["UNSUPPORTED"]
    examples = sorted(group_members.items(), key=lambda kv: (len(kv[1]), kv[0]))[:5]
    md = [
        "# Kuujjuaq Manual Team to reference VJR reconciliation", "",
        "## Evidence and method", "",
        "The actual Manual Team workbook and actual Kuujjuaq VJR workbook were read without saving either source. The VJR is treated as a practical reference. The PPT remains the authority for business rules. Source rows are reconciled against Manual DATA and the existing aggregate `TAKEOFF-DUCT` review evidence; exact candidate rows are taken only from the existing record-level reconciliation. Same-size similarity alone is never promoted to a match.", "",
        "The main CSV contains every Manual DATA record plus each unmatched populated VJR TAKE OFF line. The unresolved CSV lists Manual records that lack the documented exact group-rollup evidence. The manifest candidate is explicitly `REFERENCE_CANDIDATE_ONLY`; every writer payload is `{}` and every row is non-approved.", "",
        "## Counts", "",
        f"- Manual DATA records: **{len(manual_rows)}** (DATA rows {min(manual_rows)}–{max(manual_rows)}).",
        "- Direct, one-Manual-row to one-VJR-row exact matches: **0**.",
        f"- Exact grouped reference pairs: **{count_exact_groups}** VJR rows covering **{grouped_manual}** Manual DATA rows.",
        f"- Unmatched Manual records: **{manual_unmatched}**.",
        f"- VJR-only populated TAKE OFF records: **{len(vjr_only_rows)}**.",
        f"- SAFE_REFERENCE_MAPPING: **{safe_count}**.",
        f"- NEEDS_CONFIRMATION: **{needs}**.",
        f"- UNSUPPORTED: **{unsupported}**.",
        f"- PPT conflict / undefined reference findings: **{len(conflicts)}**.", "",
        "A grouped reference match means the existing reconciliation identified a Manual `TAKEOFF-DUCT` roll-up and a VJR row with the matching roll-up evidence. It does not mean each DATA row can be written separately. Its member lengths must be aggregated once; duplicating the shared group total would be wrong. All 47 member records remain NEEDS_CONFIRMATION because complete field transformation and current writer destination requirements are not satisfied.", "",
        "## Strongest grouped reference examples", "",
        "| Manual DATA rows | Manual group evidence | VJR destination and populated inputs | Transformation/evidence | Status |", "|---|---|---|---|---|",
    ]
    for target, members in examples:
        grp = group_rollup_meta[target]
        vr = vjr_rows[target]
        vals = vr["values"]
        group_total = sum(Decimal(str(manual_rows[source].get("Length") or 0)) for source in members)
        member_lengths = ", ".join(f"DATA!{source}={manual_rows[source].get('Length')} m" for source in sorted(members))
        code_list = ", ".join(f"{col}{target}={vals[col]}" for col in "GHIJ" if vals.get(col) not in (None, "")) or "none"
        md.append(f"| {', '.join(map(str, sorted(members)))} | Member lengths: {member_lengths}; sum **{group_total} m** | `TAKE OFF!A{target}:J{target}`; {vals.get('A')}×{vals.get('B')}, C={vals.get('C')}, D={vals.get('D')}, E={vals.get('E')}, F={vals.get('F')}; {code_list} | Aggregation of the named DATA rows yields the reference C length and size group; type/material/accessory fields still need complete comparison. | NEEDS_CONFIRMATION |")
    md += ["", "These examples are strongest only as reference mappings. None is SAFE_REFERENCE_MAPPING: each target is already occupied in the reference workbook and the current writer will not overwrite it; some source fields (pressure/material/accessory association) also remain incomplete or conflict with observed values.", "",
        "## Category-level reconciliation", "",
        "| Manual family | DATA records | Reference outcome | Primary transformation/blocker |", "|---|---:|---|---|"]
    family_counts = Counter(text(x.get("rule_category")) for x in dry.values())
    for family, n in sorted(family_counts.items()):
        selected = [r for r in rows_out if r["record_kind"] == "MANUAL_RECORD" and r["record_family"] == family]
        n_group = sum(r["reference_mapping_class"] == "GROUPED_REFERENCE_MATCH" for r in selected)
        n_unsup = sum(r["classification"] == "UNSUPPORTED" for r in selected)
        example_reason = next((r["evidence"] for r in selected if r["reference_mapping_class"] != "GROUPED_REFERENCE_MATCH"), "Grouped reference evidence exists, but is not write authorization.")
        md.append(f"| {family or 'unclassified'} | {n} | {n_group} Manual rows participate in exact grouped reference roll-ups; {n-n_group} unresolved/unmatched | {example_reason[:240]} |")
    md += ["", "## VJR-only rows", "", "These populated VJR records had no exact Manual group counterpart in the existing reconciliation. They are retained as VJR-only rather than assigned to the nearest Manual record.", "",
        "| VJR sheet/row | Dimensions | Length | Type | Installation | Accessory input cells |", "|---|---|---:|---|---|---|"]
    for row in vjr_only_rows:
        v = vjr_rows[row]["values"]
        acc = ", ".join(f"{c}{row}={v[c]}" for c in "GHIJ" if v.get(c) not in (None, "")) or "—"
        md.append(f"| TAKE OFF!{row} | {v.get('A')}×{v.get('B')} | {v.get('C')} | {v.get('E')} | {v.get('F')} | {acc} |")
    md += ["", "## Workbook/input-cell implications", "",
        f"The client workbook uses `N4={meta['unit_selector']}`. The existing input area is TAKE OFF A:J; target coordinates and existing values/formulas are recorded in the CSV. The writer requires explicitly listed target cells to be styled and empty, rejects formulas/occupied cells, and blocks insertion/reuse absent an approved destination. Consequently the reference rows provide observed cells, not valid current write targets.",
        "The output manifest candidate includes `reference_vjr_cells_json` observations but leaves `writer_write_values_json={}`. It uses `manifest_scope=REFERENCE_CANDIDATE_ONLY`, which is not an accepted production scope. Do not convert it into an approved writer manifest without a separate decision/sign-off.", "",
        "The reference comparison found no populated manual-derived spiral length/elbow-count inputs in SPIRALE B/C/D, and no populated flex quantity inputs in ROND ET FLEX. The diameter headings and formula/calculation areas are not treated as matched output records. The Manual spiral, elbow, and flex rows remain in the unresolved ledger with their actual diameter/count/length source values and the PPT transformation/blockers.", "",
        "## Detailed ledgers", "",
        "- `kuujjuaq_reference_mapping.csv`: one line per Manual DATA record plus each unmatched populated TAKE OFF reference row; carries source values, candidate/observed VJR cells, dimensions, length, type, installation, accessories, formulas, transformation, status, and reason.",
        "- `kuujjuaq_reference_unresolved.csv`: unmatched/candidate Manual rows with unresolved reason and confirmation needs.",
        "- `kuujjuaq_reference_mapping_manifest.csv`: machine-readable reference candidate only; no write payload is authorized.",
        "- `kuujjuaq_reference_rule_conflicts.md`: cell-specific PPT discrepancies and undefined cases.", ""]
    (REPORTS / "kuujjuaq_reference_mapping.md").write_text("\n".join(md), encoding="utf-8")


if __name__ == "__main__":
    make_reports()
    print("Generated Kuujjuaq reference mapping reports from the actual Manual and VJR workbooks.")
