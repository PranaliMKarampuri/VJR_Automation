import csv, tempfile
from pathlib import Path
import openpyxl, universal_input, manual_csv_to_vjr
source=Path(r'C:\Users\91739\Downloads\VJR EXCEL FILES\VJR EXCEL FILES\DUCTING TAKEOFF_RTL LOT 1\DUCTING TAKEOFF_RTL LOT 1.xlsx')
template=Path('source_files/Ventilation Jean Roy excel.xlsx')
with tempfile.TemporaryDirectory() as d:
    p=Path(d); info=universal_input.normalize_file(source,p/'normalized.csv')
    rows=list(csv.DictReader((p/'normalized.csv').open(encoding='utf-8-sig')))
    rect=[r for r in rows if 'RECT' in r['Subject'].upper()]
    print('source selection',info['source_selection'],'sheets',info['sheets'],'categories',info['category_counts'])
    print('rect count',len(rect),'first',[(r['Subject'],r['DUCT SIZE'],r['TYPE'],r['REMARK'],r['source_sheet']) for r in rect[:8]])
    manual_csv_to_vjr.generate(p/'normalized.csv',template,p/'out.xlsx',p/'report.md')
    w=openpyxl.load_workbook(p/'out.xlsx',read_only=True,data_only=True)
    ws=w['TAKE OFF']; print('generated E20:E55',[ws.cell(i,5).value for i in range(20,56)])
    w.close()
    wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
    for sn in ['DATA','DATA-DUCT','TAKEOFF-DUCT','DATA-ACCESSORIES']:
        ws=wb[sn]; print('SHEET',sn,'max',ws.max_row,ws.max_column)
        for i,r in enumerate(ws.iter_rows(max_row=5,values_only=True),1):print(i,r)
    wb.close()
