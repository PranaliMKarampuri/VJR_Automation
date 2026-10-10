#!/usr/bin/env python3
"""Build one VJR workbook from the approved Phase 1–5 input datasets.

This orchestrator reuses the Phase 2 rectangular transformation, Phase 3
spiral aggregation, Phase 4 round/flex mapping, and Phase 5 riser mapping. It
starts from the original template and patches only the worksheet XML parts
containing approved user inputs.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
import warnings
import zipfile
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

import rectangular_duct_takeoff as rectangular
import rond_et_flex_takeoff as round_flex
import riser_takeoff as riser
import spiral_duct_takeoff as spiral


ROOT = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = ROOT / "source_files" / "Ventilation Jean Roy excel.xlsx"
DEFAULT_PHASE2_INPUT = ROOT / "data" / "rectangular_duct_phase2_cases.csv"
DEFAULT_PHASE3_INPUT = ROOT / "data" / "spiral_phase3_test_cases.csv"
DEFAULT_PHASE4_INPUT = ROOT / "data" / "rond_et_flex_phase4_test_cases.csv"
DEFAULT_PHASE5_INPUT = ROOT / "data" / "riser_phase5_test_cases.csv"
DEFAULT_OUTPUT = ROOT / "output" / "Ventilation Jean Roy - phases 1-5 integrated.xlsx"
DEFAULT_REPORT = ROOT / "output" / "phase1_5_integration_validation_report.txt"

TAKEOFF = "TAKE OFF"
SPIRAL = "SPIRALE"
ROUND_FLEX = "ROND ET FLEX"
NS_MAIN = rectangular.NS_MAIN
NS_REL_DOC = rectangular.NS_REL_DOC
NS_REL_PKG = rectangular.NS_REL_PKG


class PipelineValidationError(Exception):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("Integration input validation failed")


def _read_phase2_rows(path: Path) -> tuple[list[dict[str, Any]], str]:
    records, errors = rectangular.read_records(path)
    if errors:
        raise PipelineValidationError(["Phase 1–2 rectangular input: " + e for e in errors])
    rows, errors, units = rectangular.transform_records(records)
    if errors:
        raise PipelineValidationError(["Phase 1–2 rectangular rules: " + e for e in errors])
    return rows, units or ""


def _read_phase5_rows(path: Path, template: Path) -> tuple[list[dict[str, Any]], dict[str, Decimal]]:
    records, errors = riser.read_records(path)
    if errors:
        raise PipelineValidationError(["Phase 5 riser input: " + e for e in errors])
    unique, suppressed, errors = riser.deduplicate_floor_records(records)
    if errors:
        raise PipelineValidationError(["Phase 5 riser count-once rule: " + e for e in errors])
    rectangular_sources = [
        {key: record[key] for key in rectangular.REQUIRED_COLUMNS}
        | {"Accessories": record["Accessories"], "Installation": "RI", "Units": "M"}
        for record in unique if record["System"] == "RECTANGULAR"
    ]
    rect_rows: list[dict[str, Any]] = []
    if rectangular_sources:
        rect_rows, errors, units = rectangular.transform_records(rectangular_sources)
        if errors:
            raise PipelineValidationError(["Phase 5 rectangular riser rules: " + e for e in errors])
        if units != "M":
            raise PipelineValidationError(["Phase 5 rectangular risers require metric project units (M)"])
    spiral_sources = [record for record in unique if record["System"] == "SPIRAL"]
    riser_spiral_cells: dict[str, Decimal] = {}
    if spiral_sources:
        mapped, _, errors = riser.build_values(spiral_sources, template)
        if errors:
            raise PipelineValidationError(["Phase 5 spiral riser rules: " + e for e in errors])
        riser_spiral_cells = mapped[SPIRAL]
    return rect_rows, riser_spiral_cells


def _row_cells(records: list[dict[str, Any]], start_row: int) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for index, record in enumerate(records):
        row = start_row + index
        row_values = {
            "A": record["Width"], "B": record["Height"], "C": record["Length"],
            "D": record["MetalCode"], "E": record["Type"], "F": record["Installation"],
        }
        row_values.update({col: record["Accessories"][i] if i < len(record["Accessories"]) else None
                           for i, col in enumerate(rectangular.ACCESSORY_SLOT_COLUMNS)})
        values.update({f"{col}{row}": value for col, value in row_values.items()})
    return values


def build_input_maps(
    template: Path,
    phase2_csv: Path,
    phase3_csv: Path,
    phase4_csv: Path,
    phase5_csv: Path,
) -> tuple[dict[str, dict[str, Any]], dict[str, int], list[str]]:
    """Resolve all approved source records to concrete VJR input-cell maps."""
    phase2_rows, phase2_units = _read_phase2_rows(phase2_csv)
    phase5_rows, phase5_spiral_cells = _read_phase5_rows(phase5_csv, template)
    if phase2_rows and phase5_rows and phase2_units != "M":
        raise PipelineValidationError([
            "Phase 5 rectangular risers are documented in meters; the Phase 1–2 batch uses a different project unit."
        ])
    if phase2_rows and any(str(row["Installation"]).upper() == "RI" for row in phase2_rows):
        # Existing Phase 2 input explicitly contains an RI record; keep it as
        # supplied, while Phase 5's one-count handling applies to RiserID rows.
        pass
    all_rect_rows = phase2_rows + phase5_rows
    if len(all_rect_rows) > 1004 - rectangular.FIRST_DATA_ROW + 1:
        raise PipelineValidationError(["Combined rectangular inputs exceed TAKE OFF rows 20:1004"])
    values: dict[str, dict[str, Any]] = {TAKEOFF: {}, SPIRAL: {}, ROUND_FLEX: {}}
    values[TAKEOFF].update(_row_cells(all_rect_rows, rectangular.FIRST_DATA_ROW))
    if all_rect_rows:
        units_set = {phase2_units} if phase2_rows else {"M"}
        if phase5_rows:
            units_set.add("M")
        if len(units_set) != 1:
            raise PipelineValidationError(["Combined rectangular data has inconsistent project units"])
        values[TAKEOFF]["N4"] = next(iter(units_set))

    # Phase 3 writes spiral feet and elbow counts. Phase 5 riser feet are
    # added to the matching diameter row without replacing elbow counts.
    spiral_records, errors = spiral.read_records(phase3_csv)
    if errors:
        raise PipelineValidationError(["Phase 3 SPIRALE input: " + e for e in errors])
    diameter_to_row = spiral.diameter_rows(template)
    phase3_totals, errors = spiral.aggregate_records(spiral_records, diameter_to_row)
    if errors:
        raise PipelineValidationError(["Phase 3 SPIRALE rules: " + e for e in errors])
    for row, fields in phase3_totals.items():
        for column in spiral.INPUT_COLUMNS:
            values[SPIRAL][f"{column}{row}"] = fields[column]
    for coordinate, feet in phase5_spiral_cells.items():
        phase3_length = values[SPIRAL].get(coordinate, Decimal(0))
        values[SPIRAL][coordinate] = Decimal(str(phase3_length)) + Decimal(str(feet))

    flex_records, errors = round_flex.read_records(phase4_csv)
    if errors:
        raise PipelineValidationError(["Phase 4 ROND ET FLEX input: " + e for e in errors])
    flex_values, errors = round_flex.map_records(flex_records, template)
    if errors:
        raise PipelineValidationError(["Phase 4 ROND ET FLEX rules: " + e for e in errors])
    values[ROUND_FLEX].update(flex_values)

    counts = {
        "phase2_source_rows": len(phase2_rows),
        "phase5_rectangular_rows": len(phase5_rows),
        "phase5_spiral_rows": len(phase5_spiral_cells),
        "phase3_spiral_rows": len(phase3_totals),
        "phase4_cells": len(flex_values),
    }
    notes = [
        "Phase 1 basic A:J row mapping is exercised through the Phase 2 rectangular records and shared row writer.",
        "The legacy Phase 1 sample CSV is not merged: it has no Units field, and the integration does not infer a unit.",
    ]
    return values, counts, notes


def _sheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_map = {r.attrib["Id"]: r.attrib["Target"] for r in rels.findall(f"{{{NS_REL_PKG}}}Relationship")}
    sheets = root.find(f"{{{NS_MAIN}}}sheets")
    for sheet in sheets:
        if sheet.attrib.get("name") == sheet_name:
            target = rel_map[sheet.attrib[f"{{{NS_REL_DOC}}}id"]].replace("\\", "/")
            return target.lstrip("/") if target.startswith("/") else target if target.startswith("xl/") else "xl/" + target
    raise ValueError(f"Worksheet {sheet_name!r} not found in workbook")


def _formula_map(path: Path) -> dict[tuple[str, str], str]:
    return rectangular._formula_map(path)


def _formula_element(root: ET.Element, coordinate: str) -> ET.Element | None:
    for cell in root.findall(f".//{{{NS_MAIN}}}c"):
        if cell.attrib.get("r") == coordinate:
            return cell.find(f"{{{NS_MAIN}}}f")
    return None


def _safe_value(value: Any) -> str:
    if value is None:
        return "<blank>"
    return str(value)


def generate(
    template: Path = DEFAULT_TEMPLATE,
    phase2_csv: Path = DEFAULT_PHASE2_INPUT,
    phase3_csv: Path = DEFAULT_PHASE3_INPUT,
    phase4_csv: Path = DEFAULT_PHASE4_INPUT,
    phase5_csv: Path = DEFAULT_PHASE5_INPUT,
    output: Path = DEFAULT_OUTPUT,
) -> list[str]:
    if not template.is_file():
        raise PipelineValidationError([f"Template not found: {template}"])
    if template.resolve() == output.resolve():
        raise PipelineValidationError(["Output must not overwrite the original VJR template"])
    source_hash_before = hashlib.sha256(template.read_bytes()).hexdigest()
    values, counts, notes = build_input_maps(template, phase2_csv, phase3_csv, phase4_csv, phase5_csv)

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        workbook = load_workbook(template, data_only=False, read_only=False)
    for sheet_name, cells in values.items():
        for coordinate, value in cells.items():
            workbook[sheet_name][coordinate] = float(value) if isinstance(value, Decimal) else value

    written = {name: set(cells) for name, cells in values.items()}
    # Retain the approved Phase 1/2 behavior for independent A:J rows: expand
    # untouched carry-down formulas before replacing formulas in written input
    # cells. Calculation formulas K:AX remain outside the write set.
    formula_expansions: dict[str, str] = {}
    rect_sheet = workbook[TAKEOFF]
    for row in rect_sheet.iter_rows(min_row=rectangular.FIRST_DATA_ROW, max_row=1004, min_col=1, max_col=6):
        for cell in row:
            if (cell.column in (1, 2, 4, 5, 6)
                    and cell.coordinate not in written[TAKEOFF]
                    and isinstance(cell.value, str) and cell.value.startswith("=")):
                formula_expansions[cell.coordinate] = cell.value

    formulas_before = _formula_map(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(template, "r") as source:
        parts = {name: _sheet_part(source, name) for name in (TAKEOFF, SPIRAL, ROUND_FLEX)}
        roots = {parts[name]: ET.fromstring(source.read(parts[name])) for name in parts}
        for coordinate, formula in formula_expansions.items():
            rectangular._set_formula(roots[parts[TAKEOFF]], coordinate, formula)
        for name, cells in values.items():
            root = roots[parts[name]]
            for coordinate, value in cells.items():
                rectangular._set_cell_value(root, coordinate, value)
        patched = {part: ET.tostring(root, encoding="utf-8", xml_declaration=True) for part, root in roots.items()}
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                target.writestr(item, patched.get(item.filename, source.read(item.filename)))

    source_hash_after = hashlib.sha256(template.read_bytes()).hexdigest()
    if source_hash_before != source_hash_after:
        raise RuntimeError("Original VJR template changed during generation")

    with zipfile.ZipFile(template, "r") as source, zipfile.ZipFile(output, "r") as result:
        if set(source.namelist()) != set(result.namelist()):
            raise RuntimeError("Integrated workbook package member list differs from template")
        changed_parts = [name for name in sorted(source.namelist()) if source.read(name) != result.read(name)]
        if changed_parts != sorted(parts.values()):
            raise RuntimeError(f"Unexpected package changes: {changed_parts}")
        if result.testzip() is not None:
            raise RuntimeError("Integrated XLSX ZIP integrity check failed")
        for name, part in parts.items():
            src_root = ET.fromstring(source.read(part))
            out_root = ET.fromstring(result.read(part))
            for tag in ("dataValidations", "extLst"):
                src_node = src_root.find(f"{{{NS_MAIN}}}{tag}")
                out_node = out_root.find(f"{{{NS_MAIN}}}{tag}")
                src_xml = ET.tostring(src_node, encoding="utf-8") if src_node is not None else None
                out_xml = ET.tostring(out_node, encoding="utf-8") if out_node is not None else None
                if src_xml != out_xml:
                    raise RuntimeError(f"{name} {tag} changed")

    formulas_after = _formula_map(output)
    written_cells = {(name, cell) for name, cells in written.items() for cell in cells}
    expected_formula_map = {key: formula for key, formula in formulas_before.items() if key not in written_cells}
    if formulas_after != expected_formula_map:
        raise RuntimeError("A formula outside the explicitly written input cells changed")
    if rectangular._takeoff_calculation_formulas(formulas_before) != rectangular._takeoff_calculation_formulas(formulas_after):
        raise RuntimeError("TAKE OFF K:AX calculation formulas changed")
    for sheet_name in (SPIRAL, ROUND_FLEX):
        if {key: value for key, value in formulas_before.items() if key[0] == sheet_name} != {
            key: value for key, value in formulas_after.items() if key[0] == sheet_name
        }:
            raise RuntimeError(f"Formula preservation failed on {sheet_name}")

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        output_book = load_workbook(output, data_only=False, read_only=False)
        source_book = load_workbook(template, data_only=False, read_only=False)
    if source_book.sheetnames != output_book.sheetnames:
        raise RuntimeError("Worksheet structure changed")
    for sheet_name, cells in values.items():
        src_sheet, out_sheet = source_book[sheet_name], output_book[sheet_name]
        if set(map(str, src_sheet.merged_cells.ranges)) != set(map(str, out_sheet.merged_cells.ranges)):
            raise RuntimeError(f"Merged cells changed in {sheet_name}")
        if len(src_sheet._images) != len(out_sheet._images):
            raise RuntimeError(f"Image count changed in {sheet_name}")
        for coordinate, expected in cells.items():
            actual = out_sheet[coordinate].value
            matches = float(actual) == float(expected) if isinstance(expected, Decimal) else actual == expected
            if not matches:
                raise RuntimeError(f"Output mismatch at {sheet_name}!{coordinate}: {actual!r} != {expected!r}")
        for row in src_sheet.iter_rows():
            for cell in row:
                if sheet_name == TAKEOFF and cell.coordinate in formula_expansions:
                    pass
                elif cell.coordinate not in cells and out_sheet[cell.coordinate].value != cell.value:
                    raise RuntimeError(f"Unexpected cell change at {sheet_name}!{cell.coordinate}")
                if cell.style_id != out_sheet[cell.coordinate].style_id:
                    raise RuntimeError(f"Style changed at {sheet_name}!{cell.coordinate}")

    replaced_formulas = sorted(f"{sheet}!{coordinate}" for sheet, coordinate in formulas_before if (sheet, coordinate) in written_cells)
    changed_cells = []
    for sheet_name in (TAKEOFF, SPIRAL, ROUND_FLEX):
        changed_cells.append(sheet_name + ": " + ", ".join(
            f"{coordinate}={_safe_value(value)}" for coordinate, value in values[sheet_name].items()
        ))
    report_lines = [
        f"PASS: combined rectangular Phase 1/2 rows ({counts['phase2_source_rows']} expanded rows) and Phase 5 rectangular risers ({counts['phase5_rectangular_rows']} expanded rows) into TAKE OFF!A20:J{19 + counts['phase2_source_rows'] + counts['phase5_rectangular_rows']}.",
        f"PASS: combined Phase 3 spiral lengths/elbow counts for {counts['phase3_spiral_rows']} diameter rows with Phase 5 riser feet for {counts['phase5_spiral_rows']} diameter rows.",
        f"PASS: applied Phase 4 round/flex mappings to {counts['phase4_cells']} confirmed ROND ET FLEX input cells.",
        "EXACT INPUT CELLS: " + " | ".join(changed_cells),
        f"PASS: preserved all {len(rectangular._takeoff_calculation_formulas(formulas_before)):,} TAKE OFF K:AX calculation formulas exactly.",
        f"PASS: preserved all {sum(1 for (s, _) in formulas_before if s == SPIRAL):,} SPIRALE formulas and all {sum(1 for (s, _) in formulas_before if s == ROUND_FLEX):,} ROND ET FLEX formulas exactly.",
        "LIMITATION: independent rectangular/riser inputs replace template carry-down formulas at: "
        + (", ".join(replaced_formulas) if replaced_formulas else "none")
        + ". This retains the approved Phase 1–5 behavior; formulas were not recalculated.",
        f"PASS: workbook opens; ZIP integrity, worksheet names, styles, merged cells, and image counts verified. Only {', '.join(sorted(parts.values()))} differ from the template.",
        "PASS: validations and extension XML are unchanged on edited sheets; all other workbook package parts are byte-identical to the template.",
        f"PASS: original template SHA-256 is unchanged during generation ({source_hash_after}).",
        "NOTE: the legacy Phase 1 CSV is not merged because it has no Units column; no unit was inferred. The Phase 2 CSV exercises the Phase 1 A:J input mapping and the approved Phase 2 rules.",
        "NOT CHECKED: formula recalculation. No Excel/LibreOffice calculation engine was used.",
        f"OUTPUT: {output}",
    ]
    return report_lines


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--phase2-input", type=Path, default=DEFAULT_PHASE2_INPUT)
    parser.add_argument("--phase3-input", type=Path, default=DEFAULT_PHASE3_INPUT)
    parser.add_argument("--phase4-input", type=Path, default=DEFAULT_PHASE4_INPUT)
    parser.add_argument("--phase5-input", type=Path, default=DEFAULT_PHASE5_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    try:
        lines = generate(args.template, args.phase2_input, args.phase3_input, args.phase4_input, args.phase5_input, args.output)
    except PipelineValidationError as exc:
        print("VALIDATION FAILED")
        for error in exc.errors:
            print(f"ERROR: {error}")
        return 2
    except Exception as exc:
        print(f"GENERATION FAILED: {exc}", file=sys.stderr)
        return 1
    report = "PHASE 1–5 INTEGRATION VALIDATION REPORT\n" + "\n".join(lines) + "\n"
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(report, encoding="utf-8")
    print(report, end="")
    print(f"REPORT: {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
