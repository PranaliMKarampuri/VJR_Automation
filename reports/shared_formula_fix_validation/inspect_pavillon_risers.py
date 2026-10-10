import json
from pathlib import Path

import openpyxl


source = Path("reports/shared_formula_fix_validation/source_inputs/TAKEOFF_Plans_Pav_C_source.xlsx")
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
print("SHEETS", workbook.sheetnames)
for sheet in workbook.worksheets:
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[0]] if rows else []
    risers = []
    for index, row in enumerate(rows[1:], 2):
        if "RISER" in " ".join(str(value or "") for value in row).upper():
            risers.append({
                headers[column] if column < len(headers) else str(column + 1): value
                for column, value in enumerate(row) if value not in (None, "")
            })
    print("RISER_SAMPLES", sheet.title, "COUNT", len(risers),
          json.dumps(risers[:12], ensure_ascii=False, default=str))
workbook.close()
