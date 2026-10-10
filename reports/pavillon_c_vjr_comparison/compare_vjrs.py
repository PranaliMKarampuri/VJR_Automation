from pathlib import Path
from collections import Counter
import csv, json, math
import openpyxl

ROOT = Path(r"C:\Users\91739\Desktop\VJR_Automation")
MANUAL = Path(r"C:\Users\91739\Downloads\TAKEOFF_Campus MIL Pavillon C\TAKEOFF_Campus MIL Pavillon C\TAKEOFF_Plans Pav. C -révisé add1à8_VJR.xlsx")
GENERATED = ROOT / r"output\jobs\TAKEOFF_Plans_Pav._C_-r_vis_add1_8_2e354dcc6b\TAKEOFF_Plans_Pav._C_-r_vis_add1_8_VJR.xlsx"
OUT = ROOT / r"reports\pavillon_c_vjr_comparison"

def wb(path): return openpyxl.load_workbook(path, read_only=False, data_only=True)
def norm(v):
    if isinstance(v, float):
        if not math.isfinite(v): return str(v)
        return round(v, 8)
    if isinstance(v, str): return v.strip().upper()
    return v
def rowvals(ws, r, cols): return [ws.cell(r,c).value for c in cols]
def populated(vals): return any(v not in (None, "") for v in vals)
def emit(writer, sheet, kind, key, mr, gr, fields, reason):
    writer.append([sheet,kind,str(key),mr or "",gr or "",json.dumps(fields,ensure_ascii=False,default=str),reason])

m,g=wb(MANUAL),wb(GENERATED)
rows=[]
for sh in ["TAKE OFF","SPIRALE","ROND ET FLEX"]:
    if sh not in m.sheetnames or sh not in g.sheetnames: continue
    ms,gs=m[sh],g[sh]
    if sh=="TAKE OFF":
        # Grouping key uses dimensions, length, metal/type/installation. Accessory codes are compared separately.
        cols=[1,2,3,4,5,6]; acc=[7,8,9,10]
        mm={}; gg={}
        for ws,d in [(ms,mm),(gs,gg)]:
            for r in range(20,ws.max_row+1):
                core=rowvals(ws,r,cols); ac=rowvals(ws,r,acc)
                if populated(core) and (core[0] is not None or core[1] is not None):
                    d.setdefault(tuple(map(norm,core)),[]).append((r,ac))
        for key in sorted(set(mm)|set(gg),key=str):
            a,b=mm.get(key,[]).copy(),gg.get(key,[]).copy()
            while a and b:
                mr,ma=a.pop(0); gr,ga=b.pop(0)
                if tuple(map(norm,ma))==tuple(map(norm,ga)): emit(rows,sh,"MATCH",key,mr,gr,{},"Core duct fields and accessory slots agree")
                else: emit(rows,sh,"MAPPING DIFFERENCE",key,mr,gr,{"accessories_manual":ma,"accessories_generated":ga},"Same core duct key; accessory codes differ")
            for mr,ma in a:
                if key[5] == "RI":
                    emit(rows,sh,"UNRESOLVED / INSUFFICIENT SOURCE EVIDENCE",key,mr,None,{"accessories_manual":ma},"Manual RI entry has no generated counterpart; source has zero length on all 197 riser records, so its reference length cannot be validated against source")
                else:
                    emit(rows,sh,"MISSING IN GENERATED",key,mr,None,{"accessories_manual":ma},"Manual duct record has no generated record with same core key")
            for gr,ga in b: emit(rows,sh,"EXTRA IN GENERATED",key,None,gr,{"accessories_generated":ga},"Generated duct record has no manual record with same core key")
        # Riser evidence uses the user's confirmed rule; direction is ignored.
    elif sh=="SPIRALE":
        # A is diameter; B:P are spiral length, fitting counts, dampers and related takeoff inputs.
        cols=list(range(1,17)); mm={};gg={}
        for ws,d in [(ms,mm),(gs,gg)]:
            for r in range(11,ws.max_row+1):
                v=rowvals(ws,r,cols)
                if v[0] is not None and str(v[0]).replace('.','',1).isdigit(): d[norm(v[0])]=(r,v)
        for key in sorted(set(mm)|set(gg),key=str):
            if key not in mm: emit(rows,sh,"EXTRA IN GENERATED",key,None,gg[key][0],{"values":gg[key][1]},"Diameter row present only in generated")
            elif key not in gg: emit(rows,sh,"MISSING IN GENERATED",key,mm[key][0],None,{"values":mm[key][1]},"Diameter row present only in manual")
            else:
                mr,mv=mm[key];gr,gv=gg[key]
                dif={f"col_{i+1}":{"manual":a,"generated":b} for i,(a,b) in enumerate(zip(mv,gv)) if norm(a)!=norm(b)}
                if not dif: emit(rows,sh,"MATCH",key,mr,gr,{},"Diameter and takeoff inputs B:P agree")
                else:
                    cols_changed=[int(x.split('_')[1]) for x in dif]
                    kind="QUANTITY DIFFERENCE" if all(c in (2,3,4,5,6,7,8,9,10,11,12,13,14,15,16) for c in cols_changed) else "MAPPING DIFFERENCE"
                    emit(rows,sh,kind,key,mr,gr,dif,"Diameter row differs in spiral/fitting/damper input cells B:P")
    else:
        # Compare the quantity/type input blocks, not formula-derived costs.
        blocks={"TOTAL ROND":[1,2],"FLEX ACOUSTIQUE":[6,7],"FLEX ISOLE":[10,11],"SEMI RIGIDE ISOLE":[14,15],"SEMI RIGIDE ACOUSTIQUE":[18,19],"SEMI RIGIDE NON ISOLE":[22,23]}
        for name,cols in blocks.items():
            maxr=max(ms.max_row,gs.max_row)
            for r in range(3,maxr+1):
                a=rowvals(ms,r,cols); b=rowvals(gs,r,cols)
                if not populated(a) and not populated(b): continue
                key=(name, tuple(map(norm,a[1:])))
                if tuple(map(norm,a))==tuple(map(norm,b)): emit(rows,sh,"MATCH",key,r,r,{},"Quantity/type input cells agree")
                elif a[0] in (None,0) and b[0] in (None,0):
                    # Avoid reporting empty template rows where only blanks/zeros differ.
                    continue
                else: emit(rows,sh,"QUANTITY DIFFERENCE",key,r,r,{"manual":a,"generated":b},f"{name} quantity/type input differs")

with open(OUT/"row_level_comparison.csv","w",newline="",encoding="utf-8-sig") as f:
    w=csv.writer(f);w.writerow(["sheet","classification","comparison_key","manual_row","generated_row","field_evidence_json","reason"]);w.writerows(rows)
counts=Counter((r[0],r[1]) for r in rows)
summary={f"{s} | {c}":n for (s,c),n in sorted(counts.items())}
summary["manual_takeoff_riser_rows"] = sum(1 for r in range(20,m["TAKE OFF"].max_row+1) if norm(m["TAKE OFF"].cell(r,6).value)=="RI")
summary["generated_takeoff_riser_rows"] = sum(1 for r in range(20,g["TAKE OFF"].max_row+1) if norm(g["TAKE OFF"].cell(r,6).value)=="RI")
summary["manual_takeoff_riser_nonzero_lengths"] = sum(1 for r in range(20,m["TAKE OFF"].max_row+1) if norm(m["TAKE OFF"].cell(r,6).value)=="RI" and m["TAKE OFF"].cell(r,3).value not in (None,0))
summary["manual_takeoff_riser_examples"] = [{"row":r,"size":[m["TAKE OFF"].cell(r,1).value,m["TAKE OFF"].cell(r,2).value],"length":m["TAKE OFF"].cell(r,3).value,"type":m["TAKE OFF"].cell(r,5).value,"install":m["TAKE OFF"].cell(r,6).value} for r in range(20,m["TAKE OFF"].max_row+1) if norm(m["TAKE OFF"].cell(r,6).value)=="RI"][:10]
summary["manual_spiral_diameters_with_input"] = sum(1 for r in range(11,m["SPIRALE"].max_row+1) if m["SPIRALE"].cell(r,1).value is not None and populated(rowvals(m["SPIRALE"],r,range(2,17))))
summary["generated_spiral_diameters_with_input"] = sum(1 for r in range(11,g["SPIRALE"].max_row+1) if g["SPIRALE"].cell(r,1).value is not None and populated(rowvals(g["SPIRALE"],r,range(2,17))))
(OUT/"summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(summary,ensure_ascii=False,indent=2))
