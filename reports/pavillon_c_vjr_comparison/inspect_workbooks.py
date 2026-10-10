from pathlib import Path
import json
import openpyxl

FILES = {
    "manual": Path(r"C:\Users\91739\Downloads\TAKEOFF_Campus MIL Pavillon C\TAKEOFF_Campus MIL Pavillon C\TAKEOFF_Plans Pav. C -révisé add1à8_VJR.xlsx"),
    "generated": Path(r"C:\Users\91739\Desktop\VJR_Automation\output\jobs\TAKEOFF_Plans_Pav._C_-r_vis_add1_8_2e354dcc6b\TAKEOFF_Plans_Pav._C_-r_vis_add1_8_VJR.xlsx"),
}

def val(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v

result = {}
for label, path in FILES.items():
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row_num, row in enumerate(ws.iter_rows(), 1):
            values = [val(c.value) for c in row]
            if any(v is not None for v in values):
                rows.append({"row": row_num, "values": values})
        sheets[ws.title] = {"max_row": ws.max_row, "max_column": ws.max_column, "populated_rows": rows}
    result[label] = {"path": str(path), "sheetnames": wb.sheetnames, "sheets": sheets}
Path("reports/pavillon_c_vjr_comparison/workbook_contents.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps({k: {"path": v["path"], "sheetnames": v["sheetnames"], "dims": {s: [d["max_row"], d["max_column"]] for s,d in v["sheets"].items()}} for k,v in result.items()}, ensure_ascii=False, indent=2))
