"""Read-only inspection of original Pavillon C riser source evidence."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
ARCHIVE = Path(r"C:\Users\91739\Downloads\TAKEOFF_Campus MIL Pavillon C.zip")
SOURCE_COPY = ROOT / "reports" / "shared_formula_fix_validation" / "source_inputs" / "TAKEOFF_Plans_Pav_C_source.xlsx"
CANONICAL_AUDIT = Path(r"C:\Users\91739\Downloads\TAKEOFF_Plans_Pav._C_-r_vis_add1_8_canonical_record_audit.csv")
OUTPUT = ROOT / "reports" / "riser_source_investigation"


def norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def norm_size(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").upper().replace("Ã—", "X"))


def header_map(rows: list[tuple[object, ...]]) -> tuple[int, list[str]]:
    for row_index, row in enumerate(rows[:15]):
        names = [str(value or "").strip() for value in row]
        mapped = {norm_header(value) for value in names if value}
        if {"ductsize", "units"}.issubset(mapped) and ("subject" in mapped or "label" in mapped):
            return row_index, names
    return 0, [str(value or "").strip() for value in rows[0]] if rows else []


def row_dict(headers: list[str], row: tuple[object, ...]) -> dict[str, object]:
    return {headers[index] or f"COLUMN_{index + 1}": value
            for index, value in enumerate(row[:len(headers)]) if value not in (None, "")}


def find_value(data: dict[str, object], aliases: set[str]) -> object:
    return next((value for key, value in data.items() if norm_header(key) in aliases), None)


def is_riser(data: dict[str, object]) -> bool:
    text = " ".join(str(find_value(data, {key}) or "") for key in ("subject", "label", "type"))
    return "RISER" in text.upper()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ARCHIVE) as archive:
        member_name = next(name for name in archive.namelist()
                           if name.casefold().endswith(".xlsx") and "_vjr" not in name.casefold())
        source_bytes = archive.read(member_name)
    archive_hash = hashlib.sha256(source_bytes).hexdigest()
    copy_hash = hashlib.sha256(SOURCE_COPY.read_bytes()).hexdigest()
    assert archive_hash == copy_hash, "Workspace source copy is not byte-identical to the original archive member."

    formulas_book = openpyxl.load_workbook(io.BytesIO(source_bytes), read_only=False, data_only=False)
    values_book = openpyxl.load_workbook(io.BytesIO(source_bytes), read_only=False, data_only=True)
    sheet_info: dict[str, dict[str, object]] = {}
    rows_by_sheet: dict[str, list[dict[str, object]]] = {}
    formula_details: list[dict[str, object]] = []
    formulas_by_sheet_row: dict[tuple[str, int], list[dict[str, object]]] = defaultdict(list)
    merged_details: list[dict[str, str]] = []
    for sheet in formulas_book.worksheets:
        value_sheet = values_book[sheet.title]
        rows = list(sheet.iter_rows(values_only=False))
        tuples = [tuple(cell.value for cell in row) for row in rows]
        header_row, headers = header_map(tuples)
        data_rows = []
        for source_row, row in enumerate(tuples[header_row + 1:], header_row + 2):
            data = row_dict(headers, row)
            if data:
                data["__source_row__"] = source_row
                data_rows.append(data)
                for column, cell in enumerate(rows[source_row - 1], 1):
                    if cell.data_type == "f":
                        formula = {"sheet": sheet.title, "row": source_row,
                            "cell": cell.coordinate, "formula": cell.value,
                            "cached_value": value_sheet.cell(source_row, column).value}
                        formulas_by_sheet_row[(sheet.title, source_row)].append(formula)
                        if is_riser(data):
                            formula_details.append(formula)
        rows_by_sheet[sheet.title] = data_rows
        sheet_info[sheet.title] = {
            "dimensions": sheet.calculate_dimension(), "max_row": sheet.max_row, "max_column": sheet.max_column,
            "header_row": header_row + 1, "headers": headers,
            "measure_columns": [name for name in headers if re.search(
                r"length|drop|height|elevation|quantity|count|total|qty|measure|distance|ductsize|units|page|subject|label|remark|accessor",
                name, re.I)],
            "table_names": list(sheet.tables),
            "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
        }
        for merged in sheet.merged_cells.ranges:
            if "DATA" in sheet.title.upper() and merged.min_row <= sheet.max_row and merged.max_row >= header_row + 2:
                merged_details.append({"sheet": sheet.title, "range": str(merged)})

    # The Manual riser row set is the original DATA sheet, selected by Subject/Label.
    risers = [row for row in rows_by_sheet["DATA"] if is_riser(row)]
    assert len(risers) == 197, f"Expected 197 source DATA risers, found {len(risers)}."
    canonical_rows = list(csv.DictReader(CANONICAL_AUDIT.open(encoding="utf-8-sig", newline="")))
    canonical_by_row = {(row.get("source_sheet"), int(row["source_row"])): row for row in canonical_rows}

    # Index exact unit/size counterparts in candidate schedules. These are evidence of
    # a related row only; values are never treated as a riser length.
    related_sheets = {"TAKEOFF-ACCESSORIES", "DATA-ACCESSORIES", "TAKEOFF-DUCT", "DATA-DUCT",
                      "TAKEOFF-SPIRAL DUCT (FAN)", "TO-ACCESSORIES", "TO-DUCT"}
    result_rows: list[dict[str, object]] = []
    matched_counts: Counter[str] = Counter()
    matched_record_counts: Counter[str] = Counter()
    for source in risers:
        source_row = int(source["__source_row__"])
        subject = find_value(source, {"subject"})
        label = find_value(source, {"label"})
        size = find_value(source, {"ductsize"})
        units = find_value(source, {"units"})
        key = (str(units or "").strip().upper(), norm_size(size))
        related: dict[str, list[dict[str, object]]] = {}
        for sheet_name in related_sheets:
            candidates = []
            for candidate in rows_by_sheet.get(sheet_name, []):
                candidate_key = (str(find_value(candidate, {"units"}) or "").strip().upper(),
                                 norm_size(find_value(candidate, {"ductsize"})))
                if key[0] and key[1] and candidate_key == key:
                    candidate_row = int(candidate["__source_row__"])
                    candidate_values = {k: v for k, v in candidate.items() if not k.startswith("__")}
                    candidate_values["__sheet_row__"] = candidate_row
                    candidate_values["__formula_cells__"] = formulas_by_sheet_row.get((sheet_name, candidate_row), [])
                    candidates.append(candidate_values)
            if candidates:
                related[sheet_name] = candidates
                matched_counts[sheet_name] += len(candidates)
                matched_record_counts[sheet_name] += 1
        canonical = canonical_by_row.get(("DATA", source_row), {})
        all_source_values = {k: v for k, v in source.items() if not k.startswith("__")}
        result_rows.append({
            "record_id": f"DATA!{source_row}", "source_sheet": "DATA", "source_row": source_row,
            "explicit_record_id_in_source": "none; source has no record ID column",
            "subject": subject, "label": label, "page_index": find_value(source, {"pageindex"}),
            "original_size": size, "units_system": units,
            "source_length": find_value(source, {"length", "totallength", "runlength"}),
            "source_length_unit": find_value(source, {"lengthunit", "unitoflength", "lengthunits"}),
            "source_count_quantity": find_value(source, {"count", "quantity", "qty", "total"}),
            "all_original_source_columns_json": json.dumps(all_source_values, ensure_ascii=False, default=str),
            "canonical_source_length": canonical.get("source_length"),
            "canonical_source_length_unit": canonical.get("source_length_unit"),
            "canonical_source_count": canonical.get("source_count"),
            "canonical_size": canonical.get("raw_duct_size"),
            "canonical_source_row": canonical.get("source_physical_row", canonical.get("source_row")),
            "canonical_match": "yes" if canonical else "no",
            "related_exact_system_size_rows_json": json.dumps(related, ensure_ascii=False, default=str),
            "formula_cells_on_record": json.dumps([f for f in formula_details
                if f["sheet"] == "DATA" and f["row"] == source_row], ensure_ascii=False),
            "merged_ranges_covering_record": json.dumps([m for m in merged_details
                if m["sheet"] == "DATA" and int(re.search(r"\d+", m["range"]).group()) <= source_row
                and int(re.search(r":\w+(\d+)", m["range"]).group(1)) >= source_row], ensure_ascii=False),
        })

    out_csv = OUTPUT / "pavillon_c_197_riser_source_rows.csv"
    with out_csv.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(result_rows[0]))
        writer.writeheader()
        writer.writerows(result_rows)

    length_values = Counter(str(row["source_length"]) for row in result_rows)
    length_units = Counter(str(row["source_length_unit"]) for row in result_rows)
    sizes = Counter(str(row["original_size"]) for row in result_rows)
    size_kinds = Counter("rectangular_pair" if re.fullmatch(
        r"\s*\d+(?:\.\d+)?\s*[xXÃ—]\s*\d+(?:\.\d+)?\s*", str(row["original_size"] or ""))
                         else "single_number_or_other" for row in result_rows)
    size_counts = Counter(str(row["original_size"]) for row in result_rows)
    data_accessory_same_rows = sum(bool(row.get("DATA-ACCESSORIES")) for row in
                                   [json.loads(r["related_exact_system_size_rows_json"]) for r in result_rows])
    summary = {
        "original_archive": str(ARCHIVE), "archive_member": member_name,
        "archive_member_sha256": archive_hash, "byte_identical_to_prior_workspace_copy": True,
        "riser_records": len(result_rows), "explicit_source_record_id": False,
        "investigation_record_id_convention": "DATA!<physical row>",
        "source_length_value_counts": dict(length_values), "source_length_unit_counts": dict(length_units),
        "source_count_quantity_value_counts": dict(Counter(str(row["source_count_quantity"]) for row in result_rows)),
        "source_size_shape_counts": dict(size_kinds), "source_size_counts": dict(size_counts),
        "exact_same_system_and_size_related_row_counts": dict(matched_counts),
        "riser_record_counts_with_exact_same_system_size_match": dict(matched_record_counts),
        "riser_records_with_DATA_ACCESSORIES_exact_match": data_accessory_same_rows,
        "formula_cells_on_riser_rows": formula_details,
        "total_formula_cells_in_source_workbook": sum(len(items) for items in formulas_by_sheet_row.values()),
        "data_sheet_merged_ranges": [m for m in merged_details if m["sheet"] == "DATA"],
        "workbook_defined_names": list(formulas_book.defined_names),
        "sheets": sheet_info,
        "canonical_audit_records": len(canonical_rows),
        "canonical_riser_records": sum(row.get("category") == "riser" for row in canonical_rows),
        "canonical_riser_rows_matching_original_data_physical_rows": sum(
            row.get("canonical_match") == "yes" for row in result_rows),
        "single_number_sizes": {size: count for size, count in sizes.items()
                                 if not re.fullmatch(r"\s*\d+(?:\.\d+)?\s*[xXÃ—]\s*\d+(?:\.\d+)?\s*", size)},
        "diameter_35_nominal_candidates": [row["record_id"] for row in result_rows
                                           if str(row["original_size"]).strip() == "881"],
        "row_evidence_csv": str(out_csv),
    }
    (OUTPUT / "pavillon_c_riser_source_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    formulas_book.close()
    values_book.close()
    print(json.dumps(summary, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
